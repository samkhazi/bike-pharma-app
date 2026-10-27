"""Hero e-catalogue (ecatalogue.heromotocorp.biz:8080) - weekly.

Login: public guest user. POST /HeroeCat/getGeneralUserDetails returns an authToken (no password needed).
       The site only answers from Indian IPs.
Data:  fetchModels (listed models) + the hidden/old models in hero_hidden_models.json,
       fetchAggregates (diagrams of a model), fetchBoms (parts of a diagram, with MRP),
       partSearch with supercedence=<prefix> (old -> new part numbers).
Images: diagram  https://ecatalogue.heromotocorp.biz/sol_dealer/dealer/ecat_print/group_image/<groupNo>.plt (a JPEG)
        photo    .../ecat_print/part_image/<PART-NO>.jpg
Change check: a model is rebuilt when it is new or any of its diagrams' part lines / MRPs changed.
"""
import os, re, json, hashlib, logging, itertools
from concurrent.futures import ThreadPoolExecutor
from common import Http, env, data_dir, load_json, save_json, today, Report

log = logging.getLogger('brand_sync.hero')
BASE = 'https://ecatalogue.heromotocorp.biz:8080/HeroeCat/'
IMG = 'https://ecatalogue.heromotocorp.biz/sol_dealer/dealer/ecat_print/'
HERE = os.path.dirname(os.path.abspath(__file__))
CHARS = '0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ'


def nodash(p): return re.sub('[^A-Z0-9]', '', str(p or '').upper())


class Hero:
    name = 'Hero'
    price_label = 'MRP from Hero e-catalogue'

    def __init__(self):
        self.dir = data_dir('hero')
        self.http = Http(env('HERO_ECAT_URL', BASE), delay=float(env('HERO_DELAY', '0.2')))
        self.workers = int(env('HERO_THREADS', '6'))
        self.state_path = os.path.join(self.dir, 'state.json')
        self.state = load_json(self.state_path, {'models': {}, 'groups': {}, 'mrp': {}, 'oldnew': {}})
        # Bike Pharma server: diagrams and photos it already has from the Raw Data upload (written by its Hero build), never downloaded twice
        have = load_json(os.path.join(self.dir, 'hero_have.json'), {})
        self.have_g, self.have_p = set(have.get('groups', [])), set(have.get('photos', []))

    # ---------- login + API ----------
    def login(self):
        r = self.http.post('getGeneralUserDetails')
        try:
            tok = r.json()['authToken']
        except Exception:
            raise SystemExit(f'Hero e-catalogue guest login failed (HTTP {r.status_code}). Server India me hai? Port 8080 khula hai?')
        self.http.s.headers.update({'X-Auth-Token': tok, 'Content-Type': 'application/json'})
        log.info('Hero e-catalogue guest login ok')

    def models(self):
        r = self.http.post('fetchModels', data=json.dumps({'userCode': 'general', 'region': '', 'status': '', 'product': '', 'year': '',
                                                          'model': '', 'isFilter': 'N', 'language': '1', 'modelCC': '', 'sortFlag': ''}))
        out = {}
        for m in r.json().get('models') or []:
            out[m['name']] = {'desc': m.get('desc') or m['name'], 'year': m.get('year'), 'modelCC': m.get('modelCC'), 'listed': True}
        hidden = load_json(os.path.join(HERE, 'hero_hidden_models.json'), {})
        for k, v in hidden.items():
            out.setdefault(k, {'desc': v['desc'], 'year': v.get('year'), 'modelCC': v.get('modelCC'), 'listed': False})
        for k, v in self.state['models'].items():          # models seen earlier stay in the crawl even if the list drops them
            out.setdefault(k, {kk: v.get(kk) for kk in ('desc', 'year', 'modelCC')} | {'listed': False})
        return out

    def aggregates(self, model):
        r = self.http.get('fetchAggregates', params={'variant': '', 'model': model, 'userCode': 'general', 'color': '', 'languageCode': '1'})
        j = r.json()
        groups = []
        for a in j.get('aggregates') or []:
            for g in a.get('groups') or []:
                groups.append({'groupNo': g['groupNo'], 'map': g.get('mapName') or '', 'desc': g.get('groupDesc') or '', 'agg': a.get('name') or g.get('aggregateDesc') or ''})
        return groups, j.get('modelDetails') or {}

    def boms(self, model, group):
        r = self.http.get('fetchBoms', params={'group': group, 'userCode': 'general', 'color': '', 'modelDesc': '', 'modelNo': model,
                                               'userType': 'general', 'languageCode': '1'})
        rows = []
        for b in r.json().get('boms') or []:
            rows.append({'sn': str(b.get('refno') or ''), 'pn': (b.get('component') or '').strip(), 'desc': (b.get('partDesc') or '').strip(),
                         'qty': str(b.get('qty') or ''), 'mrp': b.get('price'), 'moq': b.get('moq'), 'serv': b.get('serviceability')})
        return rows

    def supersedence(self):
        """old -> new pairs; 2-character dashed prefixes cover the full list."""
        pairs = {}
        def one(pre):
            r = self.http.post('partSearch', data=json.dumps({'userCode': 'general', 'partno': '', 'partdesc': '', 'model': '', 'supercedence': pre, 'language': '1'}))
            try: return r.json().get('supercedenceParts') or []
            except ValueError: return []
        pres = [a + b for a, b in itertools.product(CHARS, CHARS)]
        with ThreadPoolExecutor(self.workers) as ex:
            for lst in ex.map(one, pres):
                for s in lst:
                    if s.get('oldpartno') and s.get('partno'):
                        pairs[s['oldpartno']] = {'new': s['partno'], 'old_desc': (s.get('oldpartdesc') or '').strip(), 'new_desc': (s.get('partdesc') or '').strip()}
        return pairs

    # ---------- sync ----------
    def run(self, full=False, report=None):
        rep = report or Report(self.name)
        self.login()
        first = not self.state['models']
        models = self.models()
        for k, m in models.items():
            if k not in self.state['models'] and not first:
                rep.add('Naye model', f"{m['desc']} [{k}]" + ('' if m['listed'] else ' (list me nahi, purana/hidden)'))

        # crawl every model's diagram list, then every diagram's parts (6 at a time)
        def agg(k):
            try: return k, self.aggregates(k)
            except Exception as e: log.warning('aggregates %s: %s', k, e); return k, None
        with ThreadPoolExecutor(self.workers) as ex:
            aggs = dict(ex.map(agg, models))
        jobs = {}
        for k, res in aggs.items():
            if not res: continue
            groups, md = res
            models[k]['groups'] = [g['groupNo'] for g in groups]
            models[k]['status'] = md.get('status') or ''
            models[k]['category'] = md.get('category') or ''
            for g in groups:
                jobs[g['groupNo']] = (k, g)
        def bom(item):
            gno, (k, g) = item
            try: return gno, g, self.boms(k, gno)
            except Exception as e: log.warning('boms %s: %s', gno, e); return gno, g, None
        changed_models, new_mrp = set(), {}
        with ThreadPoolExecutor(self.workers) as ex:
            for gno, g, rows in ex.map(bom, jobs.items()):
                if rows is None: continue
                sig = hashlib.sha1(json.dumps([(r['sn'], r['pn'], r['qty'], r['mrp']) for r in rows]).encode()).hexdigest()
                old = self.state['groups'].get(gno)
                path = self.group_path(gno)
                if full or not old or old['sig'] != sig or not os.path.exists(path):
                    prev = load_json(path, {}).get('rows', [])
                    save_json(path, {'group': g, 'rows': rows, 'fetched': today()})
                    changed_models.add(jobs[gno][0])
                    if old and not first:
                        a, b = {r['pn'] for r in prev}, {r['pn'] for r in rows}
                        if b - a or a - b:
                            rep.add('Parts badle', f"{models[jobs[gno][0]]['desc']} / {g['map']} {g['desc']}: +{len(b - a)} naye, -{len(a - b)} hataye"
                                    + (f" ({', '.join(sorted(b - a)[:6])})" if b - a else ''))
                self.state['groups'][gno] = {'sig': sig, 'model': jobs[gno][0]}
                for r in rows:
                    if r['pn'] and r['mrp'] not in (None, '', '0'):
                        new_mrp[r['pn']] = r['mrp']
                if gno not in self.have_g: self.download(IMG + f'group_image/{gno}.plt', self.diag_path(gno))

        # MRP + new part numbers
        old_mrp = self.state['mrp']
        for pn, p in sorted(new_mrp.items()):
            if pn not in old_mrp:
                if old_mrp: rep.add('Naye part number', f'{pn}  MRP {p}')
            elif str(old_mrp[pn]) != str(p):
                rep.add('MRP badla', f'{pn}: {old_mrp[pn]} -> {p}')
        self.state['mrp'] = {**old_mrp, **new_mrp}
        self.write_mrp_csv()

        # photos of parts not downloaded yet
        todo = [pn for pn in new_mrp if nodash(pn) not in self.have_p and not os.path.exists(self.photo_path(pn)) and not os.path.exists(self.photo_path(pn) + '.missing')]
        with ThreadPoolExecutor(self.workers) as ex:
            list(ex.map(lambda pn: self.download(IMG + f'part_image/{pn}.jpg', self.photo_path(pn)), todo))

        # old -> new
        pairs = self.supersedence()
        for old, d in pairs.items():
            prev = self.state['oldnew'].get(old)
            if self.state['oldnew'] and (not prev or prev['new'] != d['new']):
                rep.add('Purana -> naya', f"{old} -> {d['new']}  {d['new_desc']}")
        self.state['oldnew'].update(pairs)
        save_json(os.path.join(self.dir, 'old_new.json'), self.state['oldnew'])
        self.write_oldnew_xlsx(os.path.join(data_dir('output', self.name), 'Hero_Purana_Naya.xlsx'))

        self.state['models'] = {k: {kk: v.get(kk) for kk in ('desc', 'year', 'modelCC', 'listed', 'groups', 'status', 'category')}
                                for k, v in models.items() if v.get('groups')}
        self.state['last_run'] = today()
        save_json(self.state_path, self.state)
        return rep, sorted(changed_models)

    # ---------- helpers ----------
    def group_path(self, gno): return os.path.join(data_dir('hero', 'groups'), re.sub(r'[^\w.-]', '_', gno) + '.json')
    def diag_path(self, gno): return os.path.join(data_dir('hero', 'diagrams'), re.sub(r'[^\w.-]', '_', gno) + '.jpg')
    def photo_path(self, pn): return os.path.join(data_dir('hero', 'photos'), re.sub(r'[^\w.-]', '_', pn) + '.jpg')

    def download(self, url, path):
        """Saves the file exactly as Hero serves it (original size and quality)."""
        if os.path.exists(path) or os.path.exists(path + '.missing'):
            return
        try:
            r = self.http.s.get(url, timeout=60)
            if r.status_code == 200 and r.content[:3] == b'\xff\xd8\xff':
                open(path, 'wb').write(r.content)
            else:
                open(path + '.missing', 'w').write(str(r.status_code))
        except Exception as e:
            log.warning('image %s: %s', url, e)

    def write_mrp_csv(self):
        import csv
        with open(os.path.join(self.dir, 'hero_ecat_mrp.csv'), 'w', newline='', encoding='utf-8') as f:
            w = csv.writer(f); w.writerow(['part_no', 'mrp', 'part_no_bina_dash'])
            for pn in sorted(self.state['mrp']): w.writerow([pn, self.state['mrp'][pn], nodash(pn)])

    def write_oldnew_xlsx(self, path):
        """Sheet 1: purana -> naya (with the final number after chains). Sheet 2: naya -> saare purane numbers."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        on, mrp = self.state['oldnew'], self.state['mrp']
        def final(p):
            seen = set()
            while p in on and p not in seen:
                seen.add(p); p = on[p]['new']
            return p
        rev = {}
        for old, d in on.items(): rev.setdefault(d['new'], []).append(old)
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Purana -> Naya'
        def head(ws, cols):
            ws.append(cols)
            for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='0B4EA2')
            ws.freeze_panes = 'A2'
        head(ws, ['Purana Part No', 'Bina dash', 'Purana Naam', 'Naya Part No', 'Naya Naam', 'Naya MRP', 'Aakhri (final) Part No', 'Final MRP'])
        for old in sorted(on):
            d = on[old]; f = final(old)
            ws.append([old, nodash(old), d['old_desc'], d['new'], d['new_desc'], mrp.get(d['new']), f, mrp.get(f)])
        ws2 = wb.create_sheet('Naya -> Purane')
        head(ws2, ['Naya Part No', 'Bina dash', 'Naam', 'MRP', 'Kitne purane', 'Saare purane numbers (bina dash bhi)'])
        for new in sorted(rev):
            olds = sorted(rev[new])
            ws2.append([new, nodash(new), on[olds[0]]['new_desc'], mrp.get(new), len(olds), ', '.join(f'{o} ({nodash(o)})' for o in olds)])
        for w in (ws, ws2):
            for col, wd in zip('ABCDEFGH', [16, 14, 32, 16, 32, 10, 18, 10]): w.column_dimensions[col].width = wd
        wb.save(path)

    def prices(self):
        out = {}
        for pn, p in self.state['mrp'].items():
            try: out[pn.upper()] = float(str(p).replace(',', ''))
            except ValueError: pass
        return out

    # ---------- data for the PDF / Excel builders ----------
    @staticmethod
    def family(desc):
        return re.split(r'\s*[\(\-]', desc or '', 1)[0].strip() or 'Other'

    def variants(self):
        for k, m in sorted(self.state['models'].items()):
            plates = []
            for gno in m.get('groups') or []:
                d = load_json(self.group_path(gno), None)
                if not d or not d['rows']: continue
                g = d['group']; img = self.diag_path(gno)
                plates.append({'title': f"{g['map']} - {g['desc']}", 'img': img if os.path.exists(img) else '',
                               'rows': [{'sn': r['sn'], 'pn': r['pn'], 'desc': r['desc'], 'qty': r['qty']} for r in d['rows']]})
            if not plates: continue
            yield {'brand': 'Hero', 'family': self.family(m['desc']), 'model': self.family(m['desc']), 'variant': m['desc'],
                   'sku': k, 'bom': m.get('modelCC') or '', 'mfg': '',   # the description already carries the launch month/year
                   'vertical': m.get('category') or '', 'plates': plates, 'key': k}
