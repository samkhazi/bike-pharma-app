"""Royal Enfield dealer EPC (epc.royalenfield.com, PTC Arbortext Content Delivery), India & SAARC only.

Login: RE_EPC_USER, RE_EPC_PASS from .env (no OTP, no captcha). Same as by hand: LOGIN, user, password, OK.
  The browser profile is kept in data/royal_enfield/browser.
Weekly (all read inside the logged-in page with the EPC's own REST calls, the same ones its screens use):
  products/<INDIA & SAARC>/children     models, then each model's variant codes       -> new models / variants
  effectivityfilters/<variant>/topicRoots  catalogue (root) of new variants
  topics/<root>/childitems (model level)   sections and their parts lists (diagrams)  -> new / removed diagrams
  partslists/<id>/parts                 every parts list: ref no, part no, name, qty, UOM, MOQ, remarks -> changed parts
  media/<id>/files/<name>.svg           diagram, vector SVG exactly as EPC serves it (only new/changed ones and
                                        the diagrams a new/changed model PDF needs)
EPC has no real MRP (its cart shows a dummy $5.99 for every part). MRP comes from Royal Enfield's dealer price list:
  put the newest list (CSV or XLSX with a "Part No" and an "MRP" column) in data/royal_enfield/mrp/. Changes go to
  RE_MRP_WEBHOOK (+ RE_MRP_WEBHOOK_TOKEN) as JSON, like HERO_MRP_WEBHOOK. spares.royalenfield.com is not used.
Only reads: no cart, order or save button is ever touched. Other regions are never opened.
RE's Akamai answers data-centre IPs (this server too) 'Access Denied', so RE_EPC_PROXY sends only this browser through an
Indian home/mobile connection, e.g. http://100.x.y.z:8080 (proxy app on a phone at the shop, over Tailscale) or a paid
Indian residential proxy http://user:pass@host:port. Everything else on the server keeps its own connection.
"""
import os, re, csv, json, gzip, glob, html, time, hashlib, logging, collections
from urllib.parse import unquote
from common import env, data_dir, load_json, save_json, today, Report

log = logging.getLogger('brand_sync.royal_enfield')
BASE = os.environ.get('RE_EPC_BASE', 'https://epc.royalenfield.com')
SEED = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'seed', 'royal_enfield', 'seed.json.gz')
ROW = ['Ref', 'Part No', 'Description', 'Qty', 'UOM', 'MOQ', 'Serviceable', 'Substitute', 'Remarks']

JS = r"""
window.reDealer = () => { try { return JSON.parse(sessionStorage.getItem('SELECTED_DEALER')).id; } catch (e) { return null; } };
window.reX = (path, params, raw) => new Promise((res, rej) => {
  const dealer = reDealer();
  const q = new URLSearchParams({...(params || {}), dealerId: dealer});
  const x = new XMLHttpRequest();
  x.open('GET', '/ACD/servlet/servicecenter/rest/' + path + '?' + q.toString());
  x.setRequestHeader('Accept', raw ? '*/*' : 'application/vnd.ptc.sc+json;version=2');
  x.setRequestHeader('X-Requested-With', 'XMLHttpRequest');
  x.setRequestHeader('Accept-Language', 'en-US');
  x.setRequestHeader('x-ptc-dealer-affiliation', dealer);
  x.setRequestHeader('x-ptc-active-cart-id', '');
  x.setRequestHeader('x-ptc-shopping-parameters', '');
  x.timeout = 180000;
  x.onload = () => {
    if (x.status !== 200) return rej('HTTP ' + x.status);
    if (/^\s*<(!doctype|html)/i.test(x.responseText)) return rej('logged out (login page)');
    if (raw) return res(x.responseText);
    try { res(JSON.parse(x.responseText)); } catch (e) { rej('not json'); }
  };
  x.onerror = () => rej('network error'); x.ontimeout = () => rej('timeout');
  x.send();
});
window.reGet = async (path, params, raw) => {
  for (let k = 0; k < 4; k++) {
    try { return await reX(path, params, raw); }
    catch (e) { if (k === 3 || /logged out|HTTP 4/.test(e)) throw String(e); await new Promise(r => setTimeout(r, 3000 * (k + 1))); }
  }
};
window.reEff = id => `ContextId eq '${id}'`;
window.rePool = async (items, n, fn) => { const out = new Array(items.length); let i = 0;
  await Promise.all(Array.from({length: n}, async () => { while (i < items.length) { const k = i++; try { out[k] = await fn(items[k]); } catch (e) { out[k] = {err: String(e)}; } } }));
  return out; };
window.rePaged = async (path, params, top) => {
  const out = [];
  for (let s = 0; s < 50000; s += top) {
    const j = await reGet(path, {...params, '$skip': s, '$top': top});
    const it = j.items || [];
    out.push(...it);
    if (it.length < top) break;
  }
  return out;
};
// parts list key = last node of the id's path, same as the folder names of the Mac download
window.reKey = id => { try { let b = decodeURIComponent(id).split('::')[1]; b = b.replace(/-/g, '+').replace(/_/g, '/'); while (b.length % 4) b += '=';
  const o = JSON.parse(atob(b).split('~').slice(1).join('~')); return o.op[o.op.length - 1].replace(/[^A-Za-z0-9_#.-]+/g, '_'); } catch (e) { return id; } };
window.reRegions = async () => (await rePaged('products', {}, 200)).map(i => ({n: i.displayName, id: i.id}));
window.reChildren = async id => (await rePaged('products/' + encodeURIComponent(id) + '/children', {effectivity: reEff(id)}, 500))
  .map(i => ({n: i.displayName, id: i.id}));
window.reRoots = async eid => ((await reGet('effectivityfilters/' + encodeURIComponent(eid) + '/topicRoots',
  {'$expand': 'elementcounts', effectivity: reEff(eid), isFindRoots: 'true'})).items || [])
  .map(i => ({n: i.displayName, id: i.id, rid: (i.attributes || {}).PERSISTENTID}));
window.reItems = async (tid, eid) => (await rePaged('topics/' + encodeURIComponent(tid) + '/childitems',
  {'$filter': "docType eq 'PART,kit,partsList,Other'", effectivity: reEff(eid)}, 200))
  .map(i => ({n: i.displayName, id: i.id, th: (i.attributes || {}).THUMBNAIL || ''}));
// one model: variants, roots of variants not known yet, and the sections/parts lists of every root (model level = all variants)
window.reModel = async (mid, known) => {
  const vs = await reChildren(mid);
  const vroots = {}, topic = {};
  for (const v of vs) {
    if (known[v.n]) { vroots[v.n] = known[v.n]; continue; }
    const rs = await reRoots(v.id); vroots[v.n] = rs.map(r => r.rid); rs.forEach(r => topic[r.rid] = r.id);
  }
  const need = [...new Set(Object.values(vroots).flat())].filter(r => !topic[r]);
  for (const rid of need) {
    const vn = Object.keys(vroots).find(k => vroots[k].includes(rid)); const v = vs.find(x => x.n === vn);
    for (const r of await reRoots(v.id)) topic[r.rid] = r.id;
  }
  const tree = [];
  for (const rid of Object.keys(topic).sort()) {
    const secs = await reItems(topic[rid], mid);
    const sub = await rePool(secs, 4, s => /PartsList/.test(s.id) ? [s] : reItems(s.id, mid));
    if (sub.some(x => x && x.err)) throw 'section read failed: ' + sub.find(x => x && x.err).err;
    secs.forEach((s, i) => sub[i].filter(p => /PartsList/.test(p.id)).forEach(p => tree.push([s.n, reKey(p.id), p.n, p.th, p.id])));
  }
  return {variants: vs.map(v => [v.n, v.id]), vroots, topic, tree};
};
// one new variant: only the parts lists that apply to it
window.reVariant = async (vid, topics) => {
  const tree = [];
  for (const tid of topics) {
    const secs = await reItems(tid, vid);
    const sub = await rePool(secs, 4, s => /PartsList/.test(s.id) ? [s] : reItems(s.id, vid));
    secs.forEach((s, i) => (sub[i].err ? [] : sub[i]).filter(p => /PartsList/.test(p.id)).forEach(p => tree.push([s.n, reKey(p.id), p.n, p.th, p.id])));
  }
  return tree;
};
window.reParts = async (plid, eid) => {
  const its = await rePaged('partslists/' + encodeURIComponent(plid) + '/parts',
    {'$expand': 'associatedPart,associatedPartsLists,alternatePart', UOMLanguage: 'en-US', effectivity: reEff(eid)}, 500);
  return its.map(p => { const a = p.attributes || {}, c = p.customMetadata || {};
    return [a.itemNumber || '', a.partNumber || '', a.partsListItemName || '', a.quantity || '', (c.UOM || {}).value || '', (c.MOQ || {}).value || '',
            a.serviceable || '', a.hasSubstitute || '', a.remarks || '']; });
};
window.rePartsMany = async jobs => { const r = await rePool(jobs, 4, j => reParts(j[1], j[2])); const o = {}; jobs.forEach((j, i) => o[j[0]] = r[i]); return o; };
window.reSvg = async (plid, eid, fname) => {
  const b64 = decodeURIComponent(plid).split('::')[1];      // ids come URL-encoded (%3A%3A)
  const base = fname.replace(/\.[^.]+$/, ''), ext = fname.slice(base.length);
  const mid = encodeURIComponent('Element.Media::' + b64 + '::' + base) + ext;
  return await reGet('media/' + mid + '/files/' + encodeURIComponent(fname), {effectivity: reEff(eid), SRCType: 'Media'}, true);
};
window.reSvgMany = async jobs => { const r = await rePool(jobs, 4, j => reSvg(j[1], j[2], j[0])); const o = {}; jobs.forEach((j, i) => o[j[0]] = r[i]); return o; };
0;   // page.evaluate would otherwise call the last function defined above
"""


def num(v):
    try: return float(str(v).replace(',', '').replace('₹', '').strip())
    except (TypeError, ValueError): return None


def safe(s): return re.sub(r'\s+', ' ', re.sub(r'[\\/:*?"<>|\x00-\x1f]+', '-', str(s or ''))).strip().rstrip('.')[:150]


def split_rev(pn):
    """'RAJ00480/B' -> ('RAJ00480', 'B'); part numbers without a revision -> (pn, '')"""
    m = re.match(r'^(.+?)/([A-Z]{1,2})$', pn or '')
    return (m.group(1), m.group(2)) if m else (pn, '')


def revkey(r): return (len(r), r)


def proxy_opt():
    """RE_EPC_PROXY -> Playwright proxy (only the EPC browser uses it); None when not set"""
    v = env('RE_EPC_PROXY', '').strip()
    if not v:
        return None
    from urllib.parse import urlsplit
    u = urlsplit(v if '://' in v else 'http://' + v)
    o = {'server': f'{u.scheme}://{u.hostname}' + (f':{u.port}' if u.port else '')}
    if u.username: o['username'] = unquote(u.username)
    if u.password: o['password'] = unquote(u.password)
    return o


class RoyalEnfield:
    name = 'Royal Enfield'
    pdfs = False                 # builds its own PDFs (vector SVG diagrams kept as vectors)
    price_label = 'MRP from Royal Enfield dealer price list'

    def __init__(self):
        self.dir = data_dir('royal_enfield')
        self.state_path = os.path.join(self.dir, 'state.json')
        self.state = load_json(self.state_path, None)
        if self.state is None:
            self.state = {'models': {}, 'pl': {}, 'mrp': {}, 'mrp_file': '', 'swaps': [], 'pending_pdf': []}
            if os.path.exists(SEED):
                # 24 Oct 2026 Mac download (140 models, 4,894 variants, 6,278 parts lists) + Sam's RE MRP list
                s = json.load(gzip.open(SEED, 'rt', encoding='utf-8'))
                self.state['models'] = s['models']
                self.state['pl'] = {k: {'n': v['n'], 'th': v['th'], 'rows': v['rows']} for k, v in s['pl'].items()}
                self.state['mrp'] = {pn: {'name': v[0], 'mrp': v[1]} for pn, v in s['mrp'].items()}
                self.state['seed_date'] = s['date']

    # ---------- login ----------
    def shot(self, page, name):
        p = os.path.join(data_dir('logs'), f're_{name}_{today()}.png')
        try: page.screenshot(path=p)
        except Exception: pass
        return p

    @staticmethod
    def logged_in(page):
        try:
            if '/ACD/delivery/' not in page.url or page.locator('#username:visible, #password:visible').count():
                return False
            page.evaluate(JS)
            return page.evaluate("async () => !!reDealer() && (await reRegions()).length > 0")
        except Exception:
            return False

    def login(self, page):
        dealer = env('RE_EPC_DEALER', '')
        url = BASE + '/ACD/delivery/index.jsp' + (f'?dealer={dealer}' if dealer else '')
        try:
            page.goto(url, wait_until='load')
        except Exception as e:
            if env('RE_EPC_PROXY') and ('ERR_PROXY' in str(e) or 'ERR_TUNNEL' in str(e)):
                raise SystemExit(f"RE EPC: RE_EPC_PROXY se connect nahi hua ({str(e)[:120]}). Dukaan wala phone/proxy on he? Tailscale chal raha he?")
            raise
        if page.title().strip().lower() == 'access denied':
            raise SystemExit('RE EPC: Royal Enfield (Akamai) ne is IP ko "Access Denied" diya. '
                             + ('RE_EPC_PROXY wala IP bhi block he, dusra proxy lo.' if env('RE_EPC_PROXY') else
                                'Server ka IP data-centre ka he: brand_sync/.env me RE_EPC_PROXY set karo (README dekho).'))
        for _ in range(20):                                    # the app sets the selected dealer after it starts
            page.wait_for_timeout(1500)
            if page.locator('#username').count() or page.evaluate("() => !!sessionStorage.getItem('SELECTED_DEALER')"):
                break
        if self.logged_in(page):
            log.info('RE EPC: still logged in (saved browser profile)'); return
        if not page.locator('#username:visible').count():
            try:
                page.locator('text=LOGIN').first.click(timeout=60000)
                page.wait_for_selector('#username:visible', timeout=30000)
            except Exception:
                body = ' '.join((page.locator('body').inner_text(timeout=5000) or '').split())[:300] if page.locator('body').count() else ''
                raise SystemExit(f"RE EPC: login page nahi khula. URL {page.url} | title '{page.title()}' | page: {body} | screenshot {self.shot(page, 'no_login')}")
        page.fill('#username', env('RE_EPC_USER', required=True))
        page.fill('#password', env('RE_EPC_PASS', required=True))
        page.click('#ok')
        end = time.time() + 120
        while time.time() < end:
            page.wait_for_timeout(2000)
            if page.evaluate("() => !!sessionStorage.getItem('SELECTED_DEALER')") and self.logged_in(page):
                log.info('RE EPC login ok'); return
            if page.locator('#username:visible').count():
                msg = ' '.join(t.strip() for t in page.locator('.alert:visible, .error:visible, .text-danger:visible, #errorMsg:visible').all_inner_texts())[:200]
                if msg or time.time() > end - 100:
                    self.shot(page, 'login_failed')
                    raise SystemExit(f"RE EPC login failed{': ' + msg if msg else ''}. RE_EPC_USER / RE_EPC_PASS check karo.")
        raise SystemExit(f"RE EPC: login ke baad catalogue nahi khula ({self.shot(page, 'after_login')})")

    def call(self, page, expr, arg=None):
        try:
            return page.evaluate(expr, arg)
        except Exception as e:
            if 'logged out' not in str(e) and 'HTTP 4' in str(e):
                raise
            log.warning('RE call failed (%s), reloading', str(e)[:120])
            page.wait_for_timeout(5000)
            self.login(page)
            page.evaluate(JS)
            return page.evaluate(expr, arg)

    # ---------- sync ----------
    def run(self, full=False, report=None):
        rep = report or Report(self.name)
        from playwright.sync_api import sync_playwright
        st = self.state
        self.changes, self.new_parts, self.gone, self.swaps_now = [], [], [], []
        affected, new_variants = set(st.get('pending_pdf') or []), []
        with sync_playwright() as p:
            ctx = p.chromium.launch_persistent_context(data_dir('royal_enfield', 'browser'), args=['--disable-dev-shm-usage', '--js-flags=--max-old-space-size=4096'],
                                                       headless=env('RE_HEADFUL', '') == '', viewport={'width': 1500, 'height': 900},
                                                       proxy=proxy_opt())
            if proxy_opt(): log.info('RE EPC: via RE_EPC_PROXY %s', proxy_opt()['server'])
            page = ctx.pages[0] if ctx.pages else ctx.new_page()
            if env('RE_FRESH_LOGIN'):                        # test: forget the saved login, do the full login
                ctx.clear_cookies(); page.goto(BASE + '/ACD/delivery/index.jsp', wait_until='load'); page.evaluate('() => { localStorage.clear(); sessionStorage.clear(); }')
                log.info('RE EPC: saved login cleared (RE_FRESH_LOGIN)')
            page.set_default_timeout(600000)
            try:
                self.login(page)
                page.evaluate(JS)
                region = env('RE_REGION', 'INDIA & SAARC')
                regs = self.call(page, '() => reRegions()')
                ind = next((r for r in regs if r['n'] == region), None)
                if not ind:
                    raise RuntimeError(f"EPC me '{region}' region nahi mila ({', '.join(r['n'] for r in regs)})")
                models = self.call(page, 'id => reChildren(id)', ind['id'])
                log.info('RE EPC: %d models in %s', len(models), region)
                if len(models) < 0.8 * len(st['models']):
                    raise RuntimeError(f'EPC ne sirf {len(models)} models diye (pichli baar {len(st["models"])}) - kuch nahi badla')
                names = [m['n'] for m in models]
                for n in sorted(set(names) - set(st['models'])):
                    rep.add('Naya model', n); affected.add(n)
                for n in sorted(set(st['models']) - set(names)):
                    rep.add('EPC se hata model', n)
                    st.setdefault('removed_models', {})[n] = today()
                    del st['models'][n]
                live = {}                                      # key -> [plid, effectivity id] for this run
                for i, m in enumerate(models):
                    old = st['models'].get(m['n'], {'variants': {}, 'tree': []})
                    known = {vn: r for vn, r in old['variants'].items() if r}
                    d = self.call(page, '([id, k]) => reModel(id, k)', [m['id'], known])
                    log.info('RE EPC %d/%d %s: %d variants, %d parts lists', i + 1, len(models), m['n'], len(d['variants']), len(d['tree']))
                    vids = dict(d['variants'])
                    if old['variants']:
                        for vn in sorted(set(vids) - set(old['variants'])):
                            new_variants.append([m['n'], vn, vids[vn], [d['topic'][r] for r in d['vroots'].get(vn, []) if r in d['topic']]])
                            rep.add('Naya variant', f"{m['n']}: {vn}" + ('' if d['vroots'].get(vn) else ' (EPC me iska catalogue abhi nahi)'))
                        for vn in sorted(set(old['variants']) - set(vids)):
                            rep.add('EPC se hata variant', f"{m['n']}: {vn}")
                    elif m['n'] in st['models'] or not st['models']:
                        pass
                    else:
                        for vn in sorted(vids):
                            new_variants.append([m['n'], vn, vids[vn], [d['topic'][r] for r in d['vroots'].get(vn, []) if r in d['topic']]])
                    oldt = {(x[0], x[1]) for x in old['tree']}
                    newt = {(x[0], x[1]) for x in d['tree']}
                    if old['tree']:
                        for sec, k in sorted(newt - oldt):
                            rep.add('Naya diagram', f"{m['n']}: {sec} ({k})"); affected.add(m['n'])
                        for sec, k in sorted(oldt - newt):
                            rep.add('Diagram hata', f"{m['n']}: {sec} ({k})"); affected.add(m['n'])
                    for sec, k, n, th, plid in d['tree']:
                        live.setdefault(k, [plid, m['id'], n, th, set()])[4].add(m['n'])
                    st['models'][m['n']] = {'variants': d['vroots'], 'tree': [x[:4] for x in d['tree']]}
                    save_json(self.state_path, st)
                self.check_parts(page, rep, live, affected)
                for mn, vn, vid, topics in new_variants:
                    if topics:
                        t = self.call(page, '([v, t]) => reVariant(v, t)', [vid, topics])
                        for sec, k, n, th, plid in t:
                            live.setdefault(k, [plid, vid, n, th, set()])
                        st.setdefault('vtree', {})[f'{mn}|{vn}'] = [x[:4] for x in t]
                        st.setdefault('pending_vpdf', []).append(f'{mn}|{vn}')
                st['pending_pdf'] = sorted(affected)
                save_json(self.state_path, st)
                self.fetch_svgs(page, rep, live, affected)
                self.build_pdfs(ctx, rep)
            finally:
                ctx.close()
        self.read_mrp(rep)
        st['last_run'] = today()
        st.pop('seed_date', None)
        save_json(self.state_path, st)
        self.write_files(rep)
        return rep, []

    def check_parts(self, page, rep, live, affected):
        """every parts list once (EPC has no reliable 'last changed' date for parts lists, so all ~6,300 are read: ~30 min)"""
        st = self.state
        keys = sorted(live)
        old_pns = {r[1] for v in st['pl'].values() for r in v['rows']}
        self.pname_old = {r[1]: r[2] for v in st['pl'].values() for r in v['rows']}
        names = st.setdefault('names', {})                 # every part number ever seen (for old -> new revisions)
        names.update({pn: n for pn, n in self.pname_old.items() if pn not in names})
        changed = 0
        for i in range(0, len(keys), 120):
            jobs = [[k, live[k][0], live[k][1]] for k in keys[i:i + 120]]
            got = self.call(page, 'j => rePartsMany(j)', jobs)
            for k, rows in got.items():
                if isinstance(rows, dict):
                    log.warning('RE parts %s: %s', k, rows.get('err')); continue
                plid, eid, n, th, models = live[k]
                o = st['pl'].get(k)
                if o is None:
                    if st['pl']:
                        rep.add('Naya parts list', f"{n} ({', '.join(sorted(models)[:3])}) - {len(rows)} parts")
                    affected.update(models)
                elif [r[:9] for r in o['rows']] != [r[:9] for r in rows] or (th and th != o.get('th')):
                    changed += 1
                    affected.update(models)
                    a = {r[1] for r in rows} - {r[1] for r in o['rows']}
                    g = {r[1] for r in o['rows']} - {r[1] for r in rows}
                    msg = f"{', '.join(sorted(models)[:3])}{' +' + str(len(models) - 3) if len(models) > 3 else ''} > {n}"
                    if a or g:
                        msg += f": +{', '.join(sorted(a)[:6])}" * bool(a) + f" -{', '.join(sorted(g)[:6])}" * bool(g)
                    elif th != o.get('th'):
                        msg += ': diagram badla'
                    else:
                        msg += ': qty / ref / naam badla'
                    rep.add('Parts badle', msg)
                    # same ref number, part number replaced -> old -> new
                    ob = collections.defaultdict(set); nb = collections.defaultdict(set)
                    for r in o['rows']: ob[r[0]].add(r[1])
                    for r in rows: nb[r[0]].add(r[1])
                    for ref in ob:
                        if ref in nb and ob[ref] != nb[ref] and len(ob[ref] - nb[ref]) == 1 and len(nb[ref] - ob[ref]) == 1:
                            sw = [next(iter(ob[ref] - nb[ref])), next(iter(nb[ref] - ob[ref])), n, ', '.join(sorted(models)[:5]), today()]
                            self.swaps_now.append(sw)
                            rep.add('Purana -> naya', f'{sw[0]} -> {sw[1]} ({n})')
                st['pl'][k] = {'n': n, 'th': th, 'rows': [r[:9] for r in rows], 'models': sorted(models)}
            save_json(self.state_path, st)
            log.info('RE EPC parts lists %d/%d', min(i + 120, len(keys)), len(keys))
        st['swaps'] = st.get('swaps', []) + self.swaps_now
        new_pns = {r[1] for k in keys if k in st['pl'] for r in st['pl'][k]['rows']}
        if old_pns:
            self.new_parts = sorted(new_pns - old_pns)
            for pn in self.new_parts: st.setdefault('new_parts', {})[pn] = today()
            for pn in self.new_parts[:500]:
                rep.add('Naya part number', f"{pn} {self.pname(pn)}")
            self.gone = sorted(old_pns - new_pns)
            for pn in self.gone: st.setdefault('gone', {})[pn] = [today(), self.pname_old.get(pn, '')]
            for pn in new_pns: st.get('gone', {}).pop(pn, None)
            if self.gone:
                rep.add('EPC se hata part number', f'{len(self.gone)} parts (list Excel me)')
        st['live_pls'] = keys
        names.update({r[1]: r[2] for k in keys if k in st['pl'] for r in st['pl'][k]['rows']})
        for k in set(st['pl']) - set(keys):            # parts lists no India model uses any more
            del st['pl'][k]
        log.info('RE EPC: %d parts lists, %d changed', len(keys), changed)

    def pname(self, pn):
        for v in self.state['pl'].values():
            for r in v['rows']:
                if r[1] == pn: return r[2]
        return ''

    def svg_path(self, th):
        return os.path.join(data_dir('royal_enfield', 'svg'), safe(th))

    def fetch_svgs(self, page, rep, live, affected):
        """diagrams of new/changed parts lists and of every parts list a PDF to be built needs (vector SVG as served)"""
        st = self.state
        need = set()
        for mn in affected:
            need |= {x[1] for x in st['models'].get(mn, {}).get('tree', [])}
        for key in st.get('pending_vpdf', []):
            need |= {x[1] for x in st.get('vtree', {}).get(key, [])}
        jobs = [[live[k][3], live[k][0], live[k][1]] for k in sorted(need) if k in live and live[k][3] and not os.path.exists(self.svg_path(live[k][3]))]
        bad = 0
        for i in range(0, len(jobs), 40):
            got = self.call(page, 'j => reSvgMany(j)', jobs[i:i + 40])
            for th, svg in got.items():
                if isinstance(svg, str) and '<svg' in svg[:3000]:
                    open(self.svg_path(th), 'w', encoding='utf-8').write(svg)
                else:
                    bad += 1
        if jobs:
            log.info('RE EPC: %d diagrams downloaded, %d failed', len(jobs) - bad, bad)
        if bad:
            rep.add('Diagram download nahi hua', f'{bad} diagrams (agle run me dobara)')

    # ---------- PDF ----------
    def page_html(self, title, sub, tree):
        prices = {pn: v['mrp'] for pn, v in self.state['mrp'].items()}
        secs = collections.OrderedDict()
        for sec, k, n, th in sorted(tree, key=lambda x: (x[0], x[2])):
            secs.setdefault(sec, collections.OrderedDict()).setdefault(k, (n, th))
        pages, idx, missing = [], [], False
        for sec, items in secs.items():
            for k, (n, th) in items.items():
                pl = self.state['pl'].get(k)
                if not pl: continue
                t = sec if (n == sec and len(items) == 1) else f'{sec} — {n}'
                pid = f'p{len(pages)}'; idx.append((t, pid, len(pl['rows'])))
                f = self.svg_path(th) if th else ''
                if f and os.path.exists(f):
                    svg = re.sub(r'<\?xml[^>]*>|<!DOCTYPE[^>]*>', '', open(f, encoding='utf-8', errors='ignore').read())
                    dia = f'<div class="dia">{svg}</div>'
                else:
                    dia = '<div class="nodia">Diagram not available in EPC</div>'; missing = missing or bool(th)
                tr = ''.join(f'<tr><td>{html.escape(r[0])}</td><td class="pn">{html.escape(r[1])}</td><td>{html.escape(r[2])}'
                             f'{"<br><i>" + html.escape(r[8]) + "</i>" if r[8] not in ("", "false") else ""}</td><td>{html.escape(r[3])}</td>'
                             f'<td class="r">{"₹ " + format(prices[r[1].upper()], "g") if prices.get(r[1].upper()) else "—"}</td></tr>' for r in pl['rows'])
                pages.append((f'<section class="pg" id="{pid}"><h2>{html.escape(t)}</h2>{dia}</section>' if th else '') +
                             f'<section class="pg tbl"{"" if th else f" id={chr(34)}{pid}{chr(34)}"}><h2>{html.escape(t)} — Parts</h2><table><tr><th>Ref</th><th>Part No</th><th>Description</th><th>Qty</th><th>MRP</th></tr>{tr}</table></section>')
        cover = (f'<section class="pg cover"><div class="brand">ROYAL ENFIELD</div><h1>{html.escape(title)}</h1><div class="sub">India &amp; SAARC<br>{html.escape(sub)}</div>'
                 f'<div class="note">Source: Royal Enfield EPC, synced {today()} (diagrams original vector SVG). MRP: Royal Enfield dealer price list. Parts lists: {len(idx)}.</div></section>')
        toc = '<section class="pg"><h2>Index</h2><table>' + ''.join(f'<tr><td><a href="#{p}">{html.escape(t)}</a></td><td class="r">{n} parts</td></tr>' for t, p, n in idx) + '</table></section>'
        css = ('@page{size:A4;margin:10mm}body{font-family:Helvetica,Arial,sans-serif;font-size:9pt;color:#111}'
               '.pg{page-break-after:always}.cover{text-align:center;padding-top:60mm}.brand{font-size:28pt;font-weight:700;letter-spacing:4px;color:#8b0000}'
               'h1{font-size:24pt;margin:10mm 0}.sub{font-size:12pt;line-height:1.7}.note{margin-top:40mm;font-size:9pt;color:#555}'
               'h2{font-size:12pt;border-bottom:2px solid #8b0000;padding-bottom:2mm;margin:0 0 3mm}'
               '.dia svg{width:100%;height:auto;max-height:250mm}.nodia{padding:40mm;text-align:center;color:#888}'
               'table{width:100%;border-collapse:collapse}td,th{border:0.5pt solid #bbb;padding:1.2mm 1.5mm;vertical-align:top}th{background:#8b0000;color:#fff;text-align:left}'
               '.pn{font-family:Menlo,monospace;white-space:nowrap}.r{text-align:right;white-space:nowrap}a{color:#111;text-decoration:none}')
        return f'<!doctype html><html><head><meta charset="utf-8"><style>{css}</style></head><body>{cover}{toc}{"".join(pages)}</body></html>', len(idx), missing

    def pdf(self, ctx, path, title, sub, tree):
        doc, n, missing = self.page_html(title, sub, tree)
        if not n:
            return False
        os.makedirs(os.path.dirname(path), exist_ok=True)
        tmp = os.path.join(data_dir('royal_enfield', 'tmp'), 'page.html')
        open(tmp, 'w', encoding='utf-8').write(doc)
        pg = ctx.new_page()
        try:
            pg.goto('file://' + tmp, wait_until='load')
            pg.pdf(path=path, format='A4', print_background=True)
        finally:
            pg.close(); os.remove(tmp)
        return not missing

    def build_pdfs(self, ctx, rep):
        st = self.state
        out = data_dir('output', 'Royal Enfield', 'PDF')
        done = 0
        for mn in list(st.get('pending_pdf', [])):
            m = st['models'].get(mn)
            if not m: st['pending_pdf'].remove(mn); continue
            vns = ', '.join(sorted(m['variants']))
            ok = self.pdf(ctx, os.path.join(out, 'Models', safe(mn) + '.pdf'), mn, f'Variants: {vns}' if len(vns) < 900 else f'{len(m["variants"])} variants', m['tree'])
            done += 1
            if ok: st['pending_pdf'].remove(mn)
            save_json(self.state_path, st)
        vdone = 0
        for key in list(st.get('pending_vpdf', [])):
            mn, vn = key.split('|', 1)
            tree = st.get('vtree', {}).get(key) or []
            ok = self.pdf(ctx, os.path.join(out, 'Variants', safe(mn), f'{safe(mn)} - {safe(vn)}.pdf'), f'{mn} - {vn}', f'Variant {vn}', tree)
            vdone += 1
            if ok or not tree: st['pending_vpdf'].remove(key)
            save_json(self.state_path, st)
        if done or vdone:
            rep.add('Files update hui', f'{done} model PDF, {vdone} variant PDF (output/Royal Enfield/PDF)')

    # ---------- MRP (dealer price list) ----------
    def read_list(self, path):
        rows = []
        if path.lower().endswith(('.xlsx', '.xlsm')):
            import openpyxl
            ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
            rows = [[('' if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
        else:
            rows = list(csv.reader(open(path, encoding='utf-8-sig', errors='ignore')))
        for hi, h in enumerate(rows[:20]):
            hl = [x.strip().lower() for x in h]
            pc = next((i for i, x in enumerate(hl) if re.search(r'part\s*(no|number|#)|^material|partno', x)), None)
            mc = next((i for i, x in enumerate(hl) if re.search(r'\bmrp\b', x)), None)
            if mc is None: mc = next((i for i, x in enumerate(hl) if 'price' in x), None)
            nc = next((i for i, x in enumerate(hl) if re.search(r'desc|name', x) and i != pc), None)
            if pc is not None and mc is not None:
                out = {}
                for r in rows[hi + 1:]:
                    if len(r) <= max(pc, mc): continue
                    pn = r[pc].strip().upper()
                    if pn: out[pn] = {'name': r[nc].strip() if nc is not None and nc < len(r) else '', 'mrp': num(r[mc])}
                return out
        raise ValueError('Part No / MRP column nahi mila')

    def read_mrp(self, rep):
        st = self.state
        files = sorted(glob.glob(os.path.join(data_dir('royal_enfield', 'mrp'), '*')), key=os.path.getmtime)
        files = [f for f in files if f.lower().endswith(('.csv', '.xlsx', '.xlsm'))]
        if not files:
            return
        f = files[-1]
        sig = f'{os.path.basename(f)}|{int(os.path.getmtime(f))}'
        if sig == st.get('mrp_file'):
            return
        try:
            new = self.read_list(f)
        except Exception as e:
            rep.add('MRP list padhi nahi gayi', f'{os.path.basename(f)}: {e}'); return
        if len(new) < 100:
            rep.add('MRP list padhi nahi gayi', f'{os.path.basename(f)}: sirf {len(new)} parts'); return
        for pn, v in sorted(new.items()):
            o = st['mrp'].get(pn)
            if o and v['mrp'] is not None and o.get('mrp') is not None and abs(v['mrp'] - o['mrp']) >= 0.5:
                self.changes.append({'part_no': pn, 'old_mrp': o['mrp'], 'new_mrp': v['mrp'], 'name': v['name'] or o.get('name', '')})
        st['mrp_log'] = ([dict(c, date=today()) for c in self.changes] + st.get('mrp_log', []))[:5000]
        for c in self.changes[:300]:
            rep.add('MRP badla', f"{c['part_no']} {c['name']}: Rs {c['old_mrp']:g} -> Rs {c['new_mrp']:g}")
        rep.add('Nayi MRP list li', f'{os.path.basename(f)}: {len(new)} parts, {len(self.changes)} MRP badle')
        merged = dict(st['mrp']); merged.update(new)
        st['mrp'], st['mrp_file'] = merged, sig
        self.push(self.changes)

    def push(self, changes):
        url = env('RE_MRP_WEBHOOK')
        if not url or not changes:
            return
        import requests
        h = {'Content-Type': 'application/json'}
        if env('RE_MRP_WEBHOOK_TOKEN'): h['Authorization'] = 'Bearer ' + env('RE_MRP_WEBHOOK_TOKEN')
        for i in range(0, len(changes), 2000):
            r = requests.post(url, headers=h, timeout=120, data=json.dumps({'source': 'royal_enfield', 'date': today(), 'changes': changes[i:i + 2000]}))
            log.info('webhook %s: HTTP %s', url, r.status_code)

    # ---------- Excel ----------
    def write_files(self, rep):
        import openpyxl
        st = self.state
        out = data_dir('output', 'Royal Enfield')
        with open(os.path.join(out, f'mrp_changes_{today()}.csv'), 'w', newline='', encoding='utf-8') as f:
            w = csv.writer(f); w.writerow(['part_no', 'old_mrp', 'new_mrp', 'name'])
            for c in self.changes: w.writerow([c['part_no'], c['old_mrp'], c['new_mrp'], c['name']])
        live = set(st.get('live_pls') or st['pl'])
        info = {}
        used = collections.defaultdict(set); secs = collections.defaultdict(set)
        for mn, m in st['models'].items():
            for sec, k, n, th in m['tree']:
                for r in (st['pl'].get(k) or {}).get('rows', []):
                    used[r[1]].add(mn); secs[r[1]].add(sec)
        for k in live:
            for r in (st['pl'].get(k) or {}).get('rows', []):
                if r[1] and r[1] not in info: info[r[1]] = r
        mrp = lambda pn: (st['mrp'].get(pn.upper()) or {}).get('mrp')
        new_at = st.get('new_parts', {})
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        def sheet(title, head, rows, widths):
            ws = wb.create_sheet(title); ws.append(head)
            for r in rows: ws.append(r)
            for i, wd in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = wd
            ws.freeze_panes = 'A2'
        sheet('All Parts A-Z', ['Part No', 'Description', 'MRP (Rs)', 'UOM', 'MOQ', 'No. of models', 'Models', 'Sections', 'Remarks', 'Naya kab aaya'],
              [[pn, r[2], mrp(pn), r[4], r[5], len(used[pn]), ', '.join(sorted(used[pn]))[:1500], ', '.join(sorted(secs[pn]))[:500],
                r[8] if r[8] not in ('', 'false') else '', new_at.get(pn, '')] for pn, r in sorted(info.items())],
              [16, 36, 10, 6, 6, 8, 60, 40, 30, 8])
        sheet('MRP Badla', ['Part No', 'Naam', 'Purana MRP', 'Naya MRP', 'Kab'], [[c['part_no'], c['name'], c['old_mrp'], c['new_mrp'], c.get('date', '')] for c in st.get('mrp_log', [])], [16, 36, 11, 11, 11])
        sheet('Naye Parts', ['Part No', 'Description', 'MRP', 'Models', 'Kab aaya'],
              [[pn, info[pn][2], mrp(pn), ', '.join(sorted(used[pn]))[:800], d] for pn, d in sorted(new_at.items(), key=lambda x: (x[1], x[0]), reverse=True) if pn in info], [16, 36, 10, 60, 11])
        sheet('EPC se Hata', ['Part No', 'Description', 'Kab hata'], [[pn, v[1], v[0]] for pn, v in sorted(st.get('gone', {}).items(), key=lambda x: (x[1][0], x[0]), reverse=True)], [16, 36, 11])
        sheet('MRP Nahi Mila', ['Part No', 'Description', 'Models'], [[pn, r[2], ', '.join(sorted(used[pn]))[:800]] for pn, r in sorted(info.items()) if mrp(pn) is None], [16, 36, 60])
        wb.save(os.path.join(out, 'RE_All_Parts.xlsx'))
        # old <-> new: newer revision of the same number, and part numbers replaced in a parts list
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Purana -> Naya (revision)'
        ws.append(['Purana Part No', 'Purana MRP', 'Naya Part No', 'Naya MRP', 'Description'])
        revs = collections.defaultdict(set)
        allnames = {**st.get('names', {}), **{pn: r[2] for pn, r in info.items()}}
        for pn in allnames: b, r = split_rev(pn); revs[b].add(r) if r else None
        for b, rs in sorted(revs.items()):
            if len(rs) < 2: continue
            rs = sorted(rs, key=revkey); newest = f'{b}/{rs[-1]}'
            for r in rs[:-1]:
                ws.append([f'{b}/{r}', mrp(f'{b}/{r}'), newest, mrp(newest), allnames.get(newest, '')])
        ws2 = wb.create_sheet('Diagram me badle')
        ws2.append(['Purana Part No', 'Naya Part No', 'Parts list', 'Models', 'Kab dekha'])
        for s in st.get('swaps', []): ws2.append(s)
        for w_ in (ws, ws2):
            w_.freeze_panes = 'A2'
            for col in 'ABCDE': w_.column_dimensions[col].width = 20
        wb.save(os.path.join(out, 'RE_Purana_Naya.xlsx'))

    def rebuild(self):
        """python sync.py royal_enfield --rebuild: Excel again from saved data (+ a new MRP list if one was put in), no login"""
        rep = Report(self.name)
        self.changes, self.new_parts, self.gone = [], [], []
        self.read_mrp(rep)
        save_json(self.state_path, self.state)
        self.write_files(rep)
        rep.add('Files update hui', 'RE_All_Parts.xlsx, RE_Purana_Naya.xlsx dobara bane')
        return rep

    def prices(self):
        return {pn: v['mrp'] for pn, v in self.state['mrp'].items() if v.get('mrp')}
