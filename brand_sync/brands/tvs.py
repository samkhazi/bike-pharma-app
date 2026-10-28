"""TVS dealer DMS (advantagetvs.in) + the dealer e-catalogue it opens (advantagetvs.com/NewPartsECatalog).

Login: TVS_DEALER_ID, TVS_USER, TVS_PASS from .env, plus an OTP that TVS sends to Sam's phone.
  The OTP is never read or guessed by this tool: Sam types it himself on a small page this tool opens on the server
  (http://<server>:TVS_OTP_PORT/tvs-otp/<code>, the link is logged and emailed), or writes it to data/tvs/otp.txt.
  The browser profile is kept in data/tvs/browser, so if TVS still trusts the last OTP (Sam: valid 24 hours)
  no new OTP is needed.
Data (weekly):
  NDP-MRP Master (/Parts/NewNdpMrp.aspx, "Export Excel")  every TVS part: NDP, MRP, effective from
  e-catalogue API (Bearer = the page's own sessionStorage authToken)
    Catalouge/GetPartsearch?partid=X        series list (CModel) + price list
    Catalouge/GetVehCoordinatesbySeries     assemblies of a series + bike photo
    Catalouge/GetPartDetailsByAssemblyID    parts of an assembly + diagram (IMAGE_LINK)
    CatalougeMaster/getPrKitMaster          kits with their contents
    Catalouge/GetClusterMaster              photo-cluster parts
Only reads: no cart, order or save buttons are ever clicked.
Server update: TVS_MRP_WEBHOOK (+ TVS_MRP_WEBHOOK_TOKEN) gets MRP changes as JSON, like HERO_MRP_WEBHOOK.
"""
import os, io, re, csv, json, gzip, time, html, hashlib, secrets, logging, threading
from common import env, data_dir, load_json, save_json, today, Report

log = logging.getLogger('brand_sync.tvs')
BASE = 'https://advantagetvs.in'

# select the option whose text contains `want` (case-insensitive) in whichever <select> has it
PICK_JS = """want => {
  const w = want.toLowerCase();
  for (const s of document.querySelectorAll('select')) {
    const o = [...s.options].find(o => o.text.trim().toLowerCase() === w) || [...s.options].find(o => o.text.toLowerCase().includes(w));
    if (!o) continue;
    if (s.value === o.value) return 'already ' + (s.id || s.name) + ' = ' + o.text.trim();
    s.value = o.value; s.dispatchEvent(new Event('change', {bubbles: true}));
    return 'changed ' + (s.id || s.name) + ' = ' + o.text.trim();
  }
  return '';
}"""
SEED = env('TVS_SEED') or os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'seed', 'tvs', 'seed.json.gz')

JS = r"""
window.tvGet = async (path, params) => {
  const t = (sessionStorage.getItem('authToken') || '').replace(/^"|"$/g, '');
  const u = '/PartEcommerceAPI/api/' + path + (params ? '?' + new URLSearchParams(params) : '');
  for (let k = 0; k < 3; k++) {
    try {
      const r = await fetch(u, {headers: {Authorization: 'Bearer ' + t, 'Content-Type': 'text/plain'}});
      if (!r.ok) throw r.status;
      return (await r.json()).data;
    } catch (e) { if (k === 2) throw e; await new Promise(r => setTimeout(r, 2000)); }
  }
};
window.tvMany = async (jobs, dealer) => {        // [[series, assembly], ...] -> {series:assembly: data}, 3 at a time
  const out = {}; let i = 0;
  const w = async () => { while (i < jobs.length) { const [s, a] = jobs[i++];
    try { out[s + ':' + a] = await tvGet('Catalouge/GetPartDetailsByAssemblyID', {SERIES: s, ASSEMBLY_ID: a, DealerID: dealer}); }
    catch (e) { out[s + ':' + a] = {error: String(e)}; }
    await new Promise(r => setTimeout(r, 250)); } };
  await Promise.all([w(), w(), w()]);
  return out;
};
0;   // end on a value: Playwright calls a script whose last value is a function
"""


def qty_of(q):
    qd = [x for x in str(q.get('QV_DETAILS') or '').split(',') if x.strip() not in ('', '0')]
    return qd[0] if qd and len(set(qd)) == 1 else ('/'.join(qd) if qd else str(q.get('QV') or ''))


def asm_rows(rec):
    rows = []
    for q in rec.get('partDetails') or []:
        pn = str(q.get('PART_NO') or '').strip().upper()
        if not re.search(r'[A-Z0-9]', pn):
            pn = ''
        rows.append({'sn': str(q.get('REF_NO') or ''), 'pn': pn, 'desc': (q.get('PART_DESC') or '').strip(), 'qty': qty_of(q),
                     'mrp': q.get('MRP'), 'remarks': (q.get('REMARKS') or '').strip()})
    return rows


def sig(rows):
    return hashlib.sha1(json.dumps([[r['sn'], r['pn'], r['qty']] for r in rows]).encode()).hexdigest()[:16]


class OtpPage:
    """Tiny page on the server where Sam types the OTP he got on his phone."""
    def __init__(self, port):
        from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
        from urllib.parse import parse_qs
        self.code = secrets.token_urlsafe(12)
        self.otp, self.resend = None, False
        me = self

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a): pass
            def page(self, msg=''):
                b = (f'<html><head><meta name="viewport" content="width=device-width"></head><body style="font-family:sans-serif;padding:24px">'
                     f'<h3>TVS DMS OTP</h3><p>{html.escape(msg)}</p><form method="post"><input name="otp" inputmode="numeric" autofocus '
                     f'style="font-size:24px;width:180px"> <button style="font-size:20px">Bhejo</button></form><br>'
                     f'<form method="post"><input type="hidden" name="resend" value="1"><button>OTP nahi aaya, dobara bhejo</button></form>'
                     f'</body></html>').encode()
                self.send_response(200); self.send_header('Content-Type', 'text/html; charset=utf-8'); self.end_headers(); self.wfile.write(b)
            def do_GET(self):
                if self.path.rstrip('/') != f'/tvs-otp/{me.code}':
                    self.send_response(404); self.end_headers(); return
                self.page('Phone pe aaya OTP yahan daalo.' if not me.otp else 'OTP mil gaya, login ho raha hai.')
            def do_POST(self):
                if self.path.rstrip('/') != f'/tvs-otp/{me.code}':
                    self.send_response(404); self.end_headers(); return
                q = parse_qs(self.rfile.read(int(self.headers.get('Content-Length') or 0)).decode())
                if q.get('resend'):
                    me.resend = True; self.page('Naya OTP bheja ja raha hai, phone check karo.')
                else:
                    me.otp = re.sub(r'\D', '', (q.get('otp') or [''])[0]) or None
                    self.page('OTP mil gaya, login ho raha hai. Ye page band kar sakte ho.' if me.otp else 'OTP khali tha, dobara daalo.')
        self.srv = ThreadingHTTPServer((env('CODE_BIND', '0.0.0.0'), port), H)   # Bike Pharma server: 127.0.0.1, the website passes /tvs-otp/ on
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()

    def close(self):
        self.srv.shutdown()


class TVS:
    name = 'TVS'
    price_label = 'MRP from TVS DMS (NDP-MRP Master)'

    def __init__(self):
        self.dir = data_dir('tvs')
        self.state_path = os.path.join(self.dir, 'state.json')
        self.state = load_json(self.state_path, None)
        self.seed_parts = {}
        if self.state is None:
            self.state = {'series': {}, 'mrp': {}, 'kits': {}, 'cluster': [], 'all': []}
            if os.path.exists(SEED):
                # 22 Oct 2026 e-catalogue download: the first run downloads everything again (PDFs, diagrams)
                # but reports only what changed in the e-catalogue after 22 Oct
                sd = json.load(gzip.open(SEED, 'rt', encoding='utf-8'))
                self.state['series'], self.state['all'] = sd['series'], sd['all']
                self.seed_parts = sd['parts']

    # ---------- login ----------
    def shot(self, page, name):
        p = os.path.join(data_dir('logs'), f'tvs_{name}_{today()}.png')
        try: page.screenshot(path=p)
        except Exception: pass
        return p

    @staticmethod
    def logged_in(page):
        # strict: a login form on screen means not logged in, whatever other text the page has
        if 'login' in page.url.lower() or page.locator('#txtDealerId, #txtUserID, #txtPassword').count():
            return False
        return page.locator('text=Signout').count() > 0

    def get_otp(self, page):
        f = os.path.join(self.dir, 'otp.txt')
        if os.path.exists(f): os.remove(f)
        port = int(env('TVS_OTP_PORT', '8765'))
        op = OtpPage(port)
        link = f"{env('TVS_OTP_LINK_BASE', f'http://SERVER-IP:{port}').rstrip('/')}/tvs-otp/{op.code}"
        msg = (f'TVS DMS ne aapke phone pe OTP bheja hai. Ye link kholo aur OTP daalo:\n{link}\n\n'
               f'(ya server pe {f} file me OTP likh do). {env("TVS_OTP_WAIT", "180")} minute tak wait karega.')
        log.warning(msg)
        from notify import send
        send(msg, subject='TVS DMS OTP chahiye')
        end = time.time() + 60 * int(env('TVS_OTP_WAIT', '180'))
        try:
            while time.time() < end:
                if op.resend:
                    op.resend = False
                    if page.locator('#btnResendOtp').count():
                        page.click('#btnResendOtp'); self.settle(page); log.info('TVS: resend OTP clicked')
                if os.path.exists(f):
                    op.otp = re.sub(r'\D', '', open(f).read()) or None
                    os.remove(f)
                if op.otp:
                    return op.otp
                page.wait_for_timeout(3000)
        finally:
            op.close()
        raise SystemExit(f'TVS: {env("TVS_OTP_WAIT", "180")} minute me OTP nahi mila, sync nahi chala.')

    @staticmethod
    def settle(page):
        # TVS pages keep polling, so 'networkidle' never comes: wait for load, then a short pause
        try: page.wait_for_load_state('load', timeout=30000)
        except Exception: pass
        page.wait_for_timeout(2000)

    def login(self, page):
        dealer = env('TVS_DEALER_ID', required=True); user = env('TVS_USER', required=True); pw = env('TVS_PASS', required=True)
        # open the real login page first (Home.aspx timed out from the server); 3 tries
        for i in range(3):
            try:
                page.goto(BASE + env('TVS_LOGIN_PAGE', '/Login.aspx'), wait_until='domcontentloaded', timeout=90000)
                break
            except Exception as e:
                log.warning(f'TVS: login page try {i + 1} failed: {e}')
                if i == 2:
                    raise SystemExit(f'TVS: {BASE} server se nahi khul raha (3 baar timeout).')
                page.wait_for_timeout(10000)
        page.wait_for_timeout(3000)
        if self.logged_in(page):
            log.info(f'TVS: still logged in (saved browser profile) at {page.url} | {page.title()}'); return
        log.info(f'TVS: login page {page.url} | {page.title()}')
        page.fill('#txtDealerId', dealer)
        page.click('#btnGo'); self.settle(page)
        try:
            page.wait_for_selector('#txtUserID', state='visible', timeout=60000)
        except Exception:
            ids = page.evaluate("[...document.querySelectorAll('input,select,button')].map(e => e.id || e.name).filter(Boolean).join(' ')")
            raise SystemExit(f"TVS: Dealer ID ke baad user/password box nahi aaya (screenshot: {self.shot(page, 'after_go')}). "
                             f"URL {page.url} | boxes: {ids}")

        def fill():
            # Role, Branch, Language: find the dropdown that has the wanted option, in order (each can reload the page)
            for key, want in (('TVS_ROLE', 'Admin'), ('TVS_BRANCH', 'RANEBENNUR'), ('TVS_LANG', 'ENGLISH')):
                want = env(key) or want
                hit = page.evaluate(PICK_JS, want)
                log.info(f'TVS: {key} -> {hit or "dropdown nahi mila"}')
                if hit and hit.startswith('changed'):
                    self.settle(page)
            page.fill('#txtUserID', user)
            page.fill('#txtPassword', pw)
        fill()
        page.click('#btnOk'); self.settle(page); page.wait_for_timeout(3000)
        if not self.logged_in(page) and page.locator('#txtOTP').count():
            self.shot(page, 'otp_asked')
            otp = self.get_otp(page)
            if not page.locator('#txtPassword').input_value():
                fill()
            page.fill('#txtOTP', otp)
            page.click('#btnOk'); self.settle(page); page.wait_for_timeout(4000)
        if not self.logged_in(page):
            raise SystemExit(f"TVS login failed (screenshot: {self.shot(page, 'login_failed')}). "
                             'TVS_DEALER_ID / TVS_USER / TVS_PASS / OTP check karo; role/branch alag ho to TVS_ROLE / TVS_BRANCH do.')
        log.info('TVS login ok')

    # ---------- NDP-MRP list ----------
    def ndp_mrp(self, page):
        for _ in range(3):
            page.goto(BASE + '/Parts/NewNdpMrp.aspx', wait_until='domcontentloaded'); page.wait_for_timeout(3000)
            if page.locator('#MainContent_btnExptExcel').count():
                break
        else:
            raise RuntimeError(f"NDP-MRP page did not open ({self.shot(page, 'ndp_page')})")
        self.clear_overlay(page)
        with page.expect_download(timeout=30 * 60 * 1000) as dl:
            page.click('#MainContent_btnExptExcel', timeout=60000)
        f = os.path.join(self.dir, f'ndp_mrp_{today()}' + (os.path.splitext(dl.value.suggested_filename)[1] or '.xls'))
        dl.value.save_as(f)
        rows = self.read_table(f)
        out = {}
        for r in rows:
            k = {re.sub(r'[^a-z]', '', (a or '').lower()): v for a, v in r.items()}
            pn = str(k.get('partno') or k.get('partnumber') or '').strip().lstrip("'").strip().upper()   # TVS writes 'G2010170
            if not pn:
                continue
            out[pn] = {'desc': str(k.get('partdesc') or k.get('partdescription') or '').strip(),
                       'ndp': num(k.get('ndp') or k.get('rolndp')), 'mrp': num(k.get('mrp')),
                       'eff': str(k.get('effectivefrom') or k.get('effectivefromdate') or '').strip()[:10], 'region': str(k.get('region') or '').strip()}
        log.info('NDP-MRP: %d parts (%s)', len(out), os.path.basename(f))
        return out

    def clear_overlay(self, page):
        # a Telerik popup (notice window or loading cover) can sit over the page and block clicks
        ov = page.locator('.TelerikModalOverlay')
        for _ in range(30):
            if not ov.count() or not ov.first.is_visible():
                return
            page.wait_for_timeout(2000)
        shot = self.shot(page, 'telerik_popup')
        txt = page.evaluate("(() => [...document.querySelectorAll('.RadWindow, .rwContent, iframe[name^=Rad]')].map(e => (e.innerText || e.src || '').trim().slice(0, 200)).join(' | '))()")
        log.warning(f'TVS: popup over the page ({shot}): {txt}')
        for sel in ('.rwCloseButton', '.RadWindow a[title=Close]'):
            if page.locator(sel).count() and page.locator(sel).first.is_visible():
                page.locator(sel).first.click(); page.wait_for_timeout(2000); log.info('TVS: popup closed')
                break
        if ov.count() and ov.first.is_visible():
            raise RuntimeError(f'TVS: popup page pe atka hua he, band nahi hua (screenshot: {shot}). Text: {txt}')

    @staticmethod
    def read_table(path):
        b = open(path, 'rb').read()
        if b[:2] == b'PK':                                   # xlsx
            import openpyxl
            ws = openpyxl.load_workbook(io.BytesIO(b), read_only=True).worksheets[0]
            it = ws.iter_rows(values_only=True)
            head = [str(h or '').strip() for h in next(it)]
            return [dict(zip(head, r)) for r in it]
        if b[:4] == b'\xd0\xcf\x11\xe0':                     # old binary .xls
            import xlrd
            sh = xlrd.open_workbook(file_contents=b).sheet_by_index(0)
            head = [str(h).strip() for h in sh.row_values(0)]
            return [dict(zip(head, sh.row_values(i))) for i in range(1, sh.nrows)]
        t = b.decode('utf-8', 'replace')
        if '<tr' in t.lower():                               # HTML table saved as .xls (Telerik export)
            cell = lambda s: html.unescape(re.sub(r'<[^>]+>', '', s)).strip()
            trs = re.findall(r'<tr[^>]*>(.*?)</tr>', t, re.S | re.I)
            grid = [[cell(c) for c in re.findall(r'<t[hd][^>]*>(.*?)</t[hd]>', tr, re.S | re.I)] for tr in trs]
            grid = [g for g in grid if any(g)]
            hi = next(i for i, g in enumerate(grid) if any(re.sub(r'\W', '', x).lower() == 'partno' for x in g))
            return [dict(zip(grid[hi], g)) for g in grid[hi + 1:]]
        t = t.lstrip('\ufeff')
        first = t.split('\n', 1)[0]
        return list(csv.DictReader(io.StringIO(t), delimiter='\t' if '\t' in first else ','))   # tab text (TVS export, 15z+) or csv

    # ---------- e-catalogue ----------
    def open_ecat(self, page):
        page.goto(BASE + '/Parts/ECatalouge.aspx', wait_until='domcontentloaded')
        page.wait_for_url(re.compile(r'NewPartsECatalog'), timeout=120000)
        page.wait_for_function("() => !!sessionStorage.getItem('authToken')", timeout=120000)
        page.evaluate(JS)
        m = re.search(r'dealerId=(\d+)', page.url)
        return m.group(1) if m else env('TVS_DEALER_ID')

    def ecat(self, page, dealer, full):
        meta = page.evaluate("() => tvGet('Catalouge/GetPartsearch', {partid:'X',description:'',partdesc:'',partSeries:'',modelID:'',page:1,pageSize:1,frameNumber:''})")
        series = {m['ID']: {'name': (m.get('name') or '').strip(), 'cat': str(m.get('CatID') or '')} for m in meta.get('CModel') or []}
        cats = {str(c.get('ID') or c.get('CatID') or ''): (c.get('name') or c.get('NAME') or '') for c in meta.get('CCategory') or []}
        if len(series) < 0.8 * max(1, len(self.state['series'])):
            raise RuntimeError(f'e-catalogue gave only {len(series)} models (last time {len(self.state["series"])}) - stopping')
        veh, jobs = {}, []
        for sid in series:
            v = page.evaluate('s => tvGet("Catalouge/GetVehCoordinatesbySeries", {series: s})', sid) or {}
            veh[sid] = v
            seen = set()
            for c in v.get('coordinates') or []:
                a = c.get('ASSEMBLY_ID')
                if a and a not in seen:
                    seen.add(a); jobs.append([sid, a, (c.get('ASSEMBLY_NAME') or '').strip(), c.get('FIG_NO') or ''])
            time.sleep(0.3)
        recs = {}
        for i in range(0, len(jobs), 60):
            recs.update(page.evaluate('([j, d]) => tvMany(j, d)', [[x[:2] for x in jobs[i:i + 60]], dealer]))
        bad = [k for k, v in recs.items() if not isinstance(v, dict) or 'error' in v]
        if len(bad) > 0.1 * max(1, len(jobs)):
            raise RuntimeError(f'{len(bad)} of {len(jobs)} assemblies failed - stopping')
        kits = page.evaluate('d => tvGet("CatalougeMaster/getPrKitMaster", {dealerID: d})', dealer) or {}
        cluster = page.evaluate('() => tvGet("Catalouge/GetClusterMaster")') or {}
        return series, cats, veh, jobs, recs, kits, cluster, meta

    def save_image(self, page, url, path):
        if not url or os.path.exists(path):
            return
        try:
            r = page.request.get(url, timeout=60000)
            if r.ok and r.body():
                open(path, 'wb').write(r.body())          # bytes exactly as TVS serves them
            else:
                open(path + '.missing', 'w').write(str(r.status))
        except Exception as e:
            log.warning('image %s: %s', url, e)

    def img_path(self, url):
        ext = os.path.splitext(url.split('?')[0])[1].lower()[:5] or '.png'
        return os.path.join(data_dir('tvs', 'img'), hashlib.sha1(url.encode()).hexdigest()[:16] + ext)

    def asm_path(self, sid, aid):
        return os.path.join(data_dir('tvs', 'asm'), f'{sid}_{aid}.json')

    # ---------- sync ----------
    def run(self, full=False, report=None):
        rep = report or Report(self.name)
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            ctx = p.chromium.launch_persistent_context(data_dir('tvs', 'browser'), args=['--disable-dev-shm-usage', '--js-flags=--max-old-space-size=4096'],   # 15s: big part list crashed the tab (small /dev/shm, JS heap)
                                                       headless=env('TVS_HEADFUL', '') == '',
                                                       accept_downloads=True, viewport={'width': 1500, 'height': 900})
            page = ctx.pages[0] if ctx.pages else ctx.new_page()
            if env('TVS_FRESH_LOGIN'):                       # test: forget the saved login, do the full login + OTP
                ctx.clear_cookies(); log.info('TVS: saved login cleared (TVS_FRESH_LOGIN)')
            page.set_default_timeout(120000)
            try:
                self.login(page)
                mrp = self.ndp_mrp(page)
                if len(mrp) < 0.9 * len(self.state['mrp']):
                    raise RuntimeError(f'NDP-MRP export has only {len(mrp)} parts (last time {len(self.state["mrp"])}) - stopping, nothing changed')
                dealer = self.open_ecat(page)
                series, cats, veh, jobs, recs, kits, cluster, meta = self.ecat(page, dealer, full)
                changed = self.merge_ecat(rep, page, series, cats, veh, jobs, recs, full)
            finally:
                ctx.close()
        self.merge_rest(rep, mrp, kits, cluster)
        self.state['last_run'] = today()
        save_json(self.state_path, self.state)
        return rep, changed

    def merge_ecat(self, rep, page, series, cats, veh, jobs, recs, full):
        old = self.state['series']
        first_run = not old
        changed = set()
        before_all = set(self.state.get('all') or [])
        new_series = {}
        for sid, s in series.items():
            v = veh.get(sid) or {}
            if v.get('IMAGE_LINK'):
                self.save_image(page, v['IMAGE_LINK'], self.img_path(v['IMAGE_LINK']))
            new_series[sid] = {'name': s['name'], 'cat': cats.get(s['cat'], s['cat']), 'photo': v.get('IMAGE_LINK') or '', 'asms': {}}
            if sid not in old and not first_run:
                rep.add('Naya model', f"{s['name']} [{sid}]")
                changed.add(sid)
        for sid in set(old) - set(series):
            rep.add('E-catalogue se hata', f"{old[sid].get('name')} [{sid}]")
        for sid, aid, name, fig in jobs:
            rec = recs.get(f'{sid}:{aid}')
            if not isinstance(rec, dict) or 'error' in rec:
                prev = (old.get(sid) or {}).get('asms', {}).get(aid)
                if prev: new_series[sid]['asms'][aid] = prev          # keep last good copy
                continue
            rows = asm_rows(rec)
            if not rows:
                continue
            sg = sig(rows)
            img = rec.get('IMAGE_LINK') or ''
            new_series[sid]['asms'][aid] = {'name': name, 'fig': fig or rec.get('FIG_NO') or '', 'sig': sg, 'img': img}
            prev = (old.get(sid) or {}).get('asms', {}).get(aid)
            path = self.asm_path(sid, aid)
            prev_rows = load_json(path, {}).get('rows')
            prev_set = {r['pn'] for r in prev_rows if r['pn']} if prev_rows else set(self.seed_parts.get(f'{sid}:{aid}', []))
            if full or not prev or prev.get('sig') != sg or not os.path.exists(path):
                save_json(path, {'rows': rows, 'fetched': today()})
                changed.add(sid)
                if prev and prev.get('sig') != sg and sid in old:
                    cur = {r['pn'] for r in rows if r['pn']}
                    add, gone = sorted(cur - prev_set), sorted(prev_set - cur)
                    if add or gone:
                        rep.add('Parts badle', f"{series[sid]['name']} - {name}: +{len(add)} naye, -{len(gone)} hataye"
                                + (f" (naye: {', '.join(add[:8])})" if add else ''))
                elif not prev and sid in old:
                    rep.add('Naya diagram', f"{series[sid]['name']} - {name}")
            if img:
                self.save_image(page, img, self.img_path(img))
        allp = {r['pn'] for sid, s in new_series.items() for aid in s['asms']
                for r in load_json(self.asm_path(sid, aid), {}).get('rows', []) if r['pn']}
        if before_all:
            for pn in sorted(allp - before_all)[:500]:
                rep.add('Naya part number (e-catalogue)', pn)
        self.state['series'], self.state['all'] = new_series, sorted(allp)
        return sorted(series) if first_run else sorted(changed)

    def compare_old(self, rep, mrp=None):
        """MRP of the NDP-MRP list vs the MRP in the TVS catalogue already on the server (tvs_build.json + tvs_extra.json
        from Sam's TVS_NDP_MRP_List upload). MRP only in the report; the full list goes to a CSV in reports/."""
        mrp = mrp if mrp is not None else self.state.get('mrp') or {}
        d = env('TVS_OLD_DIR') or os.path.join(data_dir(), '..', 'tvs')
        old = {}
        for fn in ('tvs_build.json', 'tvs_extra.json'):
            try:
                for pn, m in (json.load(open(os.path.join(d, fn), encoding='utf-8')).get('mrp') or {}).items():
                    m = m.get('mrp') if isinstance(m, dict) else m
                    k = str(pn).strip().lstrip("'").upper()
                    if k and k not in old and num(m):
                        old[k] = num(m)
            except FileNotFoundError:
                pass
        if not old:
            rep.add('Purani list se milan', f'server pe purani TVS MRP nahi mili ({os.path.abspath(d)})'); return
        rows, new_pn = [], []
        for pn, v in sorted(mrp.items()):
            if v.get('mrp') is None: continue
            if pn not in old: new_pn.append(pn); continue
            if abs(v['mrp'] - old[pn]) >= 0.01:
                rows.append((pn, v['desc'], old[pn], v['mrp'], v['eff']))
        gone = sorted(set(old) - set(mrp))
        f = os.path.join(data_dir('reports'), f'tvs_mrp_vs_old_{today()}.csv')
        with open(f, 'w', newline='', encoding='utf-8') as fh:
            w = csv.writer(fh); w.writerow(['part_no', 'name', 'old_mrp', 'new_mrp', 'effective_from'])
            w.writerows(rows)
        up = sum(1 for r in rows if r[3] > r[2])
        rep.add('Purani list se milan', f'{len(old)} purane MRP vs {len(mrp)} naye: {len(rows)} MRP badle ({up} badhe, {len(rows) - up} ghate), '
                                        f'{len(new_pn)} naye parts, {len(gone)} purani list me the par nayi me nahi. Poori list: {f}')
        for r in rows[:300]:
            rep.add('MRP badla (purani list se)', f'{r[0]} {r[1]}: Rs {r[2]:g} -> Rs {r[3]:g} ({r[4]})')
        log.info('TVS: compared with the old MRP list: %d changed, %d new, %d gone', len(rows), len(new_pn), len(gone))

    def merge_rest(self, rep, mrp, kits, cluster):
        old = self.state['mrp']
        changes = []
        if old:
            for pn, v in sorted(mrp.items()):
                o = old.get(pn)
                if not o:
                    rep.add('Naya part (NDP-MRP list)', f"{pn} {v['desc']} MRP {v['mrp']}")
                elif v['mrp'] is not None and o.get('mrp') is not None and abs(v['mrp'] - o['mrp']) >= 0.01:
                    changes.append({'part_no': pn, 'old_mrp': o['mrp'], 'new_mrp': v['mrp'], 'old_ndp': o.get('ndp'), 'new_ndp': v['ndp'],
                                    'effective_from': v['eff'], 'name': v['desc']})
            for c in changes[:300]:
                rep.add('MRP badla', f"{c['part_no']} {c['name']}: Rs {c['old_mrp']:g} -> Rs {c['new_mrp']:g} ({c['effective_from']})")
            gone = sorted(set(old) - set(mrp))
            if gone:
                rep.add('NDP-MRP list se hata', f"{len(gone)} parts (list Excel me)")
        else:
            gone = []
            rep.add('Pehli baar', f'NDP-MRP list ke {len(mrp)} parts save hue; agle hafte se MRP badlav dikhenge')
            self.compare_old(rep, mrp)
        # kits
        dt = {}
        for d in kits.get('kPartsDtls') or []:
            dt.setdefault(str(d.get('KIT_ID')), []).append([str(d.get('PART_NO') or '').strip(), (d.get('PART_DESCRIPTION') or '').strip(), d.get('QTY')])
        newk = {}
        for k in kits.get('kitParts') or []:
            pn = str(k.get('KIT_PART_NO') or '').strip().upper()
            if pn:
                newk[pn] = {'desc': (k.get('KIT_DESCRIPTION') or '').strip(), 'model': (k.get('MODEL') or '').strip(),
                            'appl': (k.get('MODEL_APPL') or '').strip(), 'ndp': k.get('NDP'), 'mrp': k.get('MRP'),
                            'parts': sorted(dt.get(str(k.get('KIT_ID')), []))}
        if self.state['kits']:
            for pn, k in newk.items():
                o = self.state['kits'].get(pn)
                if not o:
                    rep.add('Naya kit', f"{pn} {k['desc']} ({k['model']}), {len(k['parts'])} parts")
                elif o['parts'] != k['parts']:
                    rep.add('Kit ke andar badla', f"{pn} {k['desc']}")
        cl = sorted({str(c.get('PART_NO') or c.get('PartNo') or '').strip().upper() for c in (cluster.get('CParts') or [])} - {''})
        self.state['mrp'], self.state['kits'], self.state['cluster'] = mrp, newk, cl
        self.write_files(mrp, changes, gone, newk, cl)
        self.push(changes)

    def write_files(self, mrp, changes, gone, kits, cluster):
        import openpyxl
        out = data_dir('output', 'TVS')
        with open(os.path.join(out, 'TVS_NDP_MRP.csv'), 'w', newline='', encoding='utf-8') as f:
            w = csv.writer(f); w.writerow(['part_no', 'description', 'ndp', 'mrp', 'effective_from', 'region'])
            for pn, v in sorted(mrp.items()):
                w.writerow([pn, v['desc'], v['ndp'], v['mrp'], v['eff'], v['region']])
        with open(os.path.join(out, f'mrp_changes_{today()}.csv'), 'w', newline='', encoding='utf-8') as f:
            w = csv.writer(f); w.writerow(['part_no', 'old_mrp', 'new_mrp', 'old_ndp', 'new_ndp', 'effective_from', 'name'])
            for c in changes:
                w.writerow([c['part_no'], c['old_mrp'], c['new_mrp'], c['old_ndp'], c['new_ndp'], c['effective_from'], c['name']])
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        def sheet(title, head, rows, widths):
            ws = wb.create_sheet(title); ws.append(head)
            for r in rows: ws.append(r)
            for i, wd in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = wd
            ws.freeze_panes = 'A2'
        sheet('MRP Badla', ['Part No', 'Naam', 'Purana MRP', 'Naya MRP', 'Purana NDP', 'Naya NDP', 'Kab se'],
              [[c['part_no'], c['name'], c['old_mrp'], c['new_mrp'], c['old_ndp'], c['new_ndp'], c['effective_from']] for c in changes],
              [14, 36, 11, 11, 11, 11, 12])
        sheet('List se Hata', ['Part No'], [[pn] for pn in gone], [14])
        sheet('Kits (andar ke parts)', ['Kit Part No', 'Kit Naam', 'Model', 'Kis me lagta', 'NDP', 'MRP', 'Andar ka Part No', 'Naam', 'Qty'],
              [[pn, k['desc'], k['model'], k['appl'], k['ndp'], k['mrp'], p[0], p[1], p[2]] for pn, k in sorted(kits.items()) for p in (k['parts'] or [['', '', '']])],
              [14, 30, 18, 22, 9, 9, 14, 30, 6])
        sheet('Photo Cluster Parts', ['Part No', 'Naam', 'MRP'], [[pn, (mrp.get(pn) or {}).get('desc', ''), (mrp.get(pn) or {}).get('mrp')] for pn in cluster], [14, 36, 10])
        sheet('Saare Parts NDP-MRP', ['Part No', 'Naam', 'NDP', 'MRP', 'Kab se', 'Region'],
              [[pn, v['desc'], v['ndp'], v['mrp'], v['eff'], v['region']] for pn, v in sorted(mrp.items())], [14, 36, 10, 10, 12, 8])
        wb.save(os.path.join(out, f'TVS_DMS_Report_{today()}.xlsx'))

    def push(self, changes):
        url = env('TVS_MRP_WEBHOOK')
        if not url or not changes:
            return
        import requests
        h = {'Content-Type': 'application/json'}
        if env('TVS_MRP_WEBHOOK_TOKEN'): h['Authorization'] = 'Bearer ' + env('TVS_MRP_WEBHOOK_TOKEN')
        for i in range(0, len(changes), 2000):
            r = requests.post(url, headers=h, timeout=120, data=json.dumps({'source': 'tvs_dms', 'date': today(), 'changes': changes[i:i + 2000]}))
            log.info('webhook %s: HTTP %s', url, r.status_code)

    # ---------- data for the PDF / Excel builders ----------
    def prices(self):
        return {pn: v['mrp'] for pn, v in self.state['mrp'].items() if v.get('mrp') is not None}

    def variants(self):
        for sid, s in sorted(self.state['series'].items(), key=lambda kv: kv[1]['name']):
            plates = []
            for aid, a in s['asms'].items():
                d = load_json(self.asm_path(sid, aid), None)
                if not d or not d['rows']:
                    continue
                img = self.img_path(a['img']) if a.get('img') else ''
                plates.append({'title': a['name'].title(), 'img': img if img and os.path.exists(img) else '',
                               'rows': [{'sn': r['sn'], 'pn': r['pn'], 'desc': r['desc'] + (f" ({r['remarks']})" if r.get('remarks') else ''),
                                         'qty': r['qty']} for r in d['rows']]})
            if not plates:
                continue
            for i, pl in enumerate(plates):
                pl['title'] = f"{i + 1} - {pl['title']}"
            fam = (s.get('cat') or 'TVS').title()
            yield {'brand': 'TVS', 'family': fam, 'model': s['name'], 'variant': s['name'], 'sku': sid, 'bom': '', 'mfg': '',
                   'vertical': fam.upper(), 'plates': plates, 'key': sid}


def num(v):
    try: return float(str(v).replace(',', '').replace('₹', '').strip())
    except (TypeError, ValueError): return None
