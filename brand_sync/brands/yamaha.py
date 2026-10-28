"""Yamaha dealer portal pymidol (www.pymidol.com).

Login: YAMAHA_USER and YAMAHA_PASS from .env (no captcha). Steps are the same as by hand: Company "Yamaha" + User Id,
  Sign In, then password, Sign in. If pymidol asks for the 4 digit OTP sent to the dealer's phone, the tool puts a
  small page on this server (http://<server>:YAMAHA_OTP_PORT/yamaha-otp/<code>, link logged and emailed) where Sam types it.
  The browser profile is kept in data/yamaha/browser, so while pymidol keeps the login alive no login is needed.
Data (weekly, all read inside the logged-in page with the portal's own calls, the same ones its screens use):
  /ydms/spare-report/spr12801   SPR128 Parts MRP list (all ~52,000 parts)          -> new parts, MRP changes
  /ydms/spare-report/spr13601   SPR136 Superseding parts (old -> new)               -> new old<->new pairs
  /ydms/spare-report/spr11801/2 SPR118 Spares kits and what is inside them          -> new / changed kits
  /spare-parts/spm1301          SPM13 Part Information: model parts catalogues (PDF) and PI bulletins
  /spare-parts/pza0301          Circulars (PDF + extra xlsx/pdf)
  /spare-parts/svt0801          SVT08 Service Actions (DSA PDFs)
  /spare-parts/model-value-list model codes
  PDFs are saved byte for byte as pymidol gives them. Only new ones are downloaded.
Only reads: no order, cart, claim or save button is ever touched. YPEC is not used.
Server update: YAMAHA_MRP_WEBHOOK (+ YAMAHA_MRP_WEBHOOK_TOKEN) gets MRP changes as JSON, like HERO_MRP_WEBHOOK.
"""
import os, re, csv, json, gzip, time, html, base64, secrets, logging, threading, collections
from notify import send_code_link
from common import env, data_dir, load_json, save_json, today, Report

log = logging.getLogger('brand_sync.yamaha')
BASE = os.environ.get('YAMAHA_BASE', 'https://www.pymidol.com')
SEED = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'seed', 'yamaha', 'seed.json.gz')

JS = r"""
window.yUser = () => JSON.parse(localStorage.getItem('user') || 'null');
window.yX = (url, body, blob) => new Promise((res, rej) => {
  const r = new XMLHttpRequest();
  r.open('POST', url);
  if (blob) r.responseType = 'blob';
  r.setRequestHeader('Content-Type', 'application/json');
  r.setRequestHeader('Accept', 'application/json, text/plain, */*');
  r.timeout = 300000;
  r.onload = async () => {
    if (blob) {
      const b = r.response;
      if (r.status !== 200 || !b || !b.size || /json|html/.test(b.type)) return res(null);
      const s = await new Promise(ok => { const f = new FileReader(); f.onload = () => ok(f.result.split(',')[1] || ''); f.readAsDataURL(b); });
      return res(s);
    }
    if (r.status !== 200) return rej('HTTP ' + r.status);
    try { res(JSON.parse(r.responseText)); } catch (e) { rej('not json (logged out?)'); }
  };
  r.onerror = () => rej('network error'); r.ontimeout = () => rej('timeout');
  r.send(JSON.stringify(body));
});
window.yCall = async (url, body, blob) => {
  for (let k = 0; k < 3; k++) {
    try { return await yX(url, body, blob); }
    catch (e) { if (k === 2 || /logged out/.test(e)) throw String(e); await new Promise(r => setTimeout(r, 4000)); }
  }
};
window.yModels = async () => (await yCall('/spare-parts/model-value-list/model-ref', {sysOwnerCd: 'IYM', user: yUser()})).data || [];
window.yPaged = async (url, body, size) => {
  const out = [];
  for (let p = 0; p < 500; p++) {
    const j = await yCall(url, {...body, pageNum: p, pageSize: size});
    const d = j.data || j;
    const rows = d.resultList || [];
    out.push(...rows);
    if (!rows.length || out.length >= (d.totalRecords || 0)) return [out, d.totalRecords || 0];
  }
  return [out, -1];
};
window.yMrp = async () => {
  const u = yUser();
  const [rows, total] = await yPaged('/ydms/spare-report/spr12801/retrieve.json', {sysOwnerCd: 'IYM', dealerCd: u.pymidolDealerCd,
    networkType: u.networkType || '000', networkCd: null, partNo: null, sameGstFlag: u.isYbcGstSame || '1', mnfType: '010'}, 2000);
  return [rows.map(r => [r.partNo, r.partNm, r.mrp]), total];
};
window.ySup = async () => { const [rows, total] = await yPaged('/ydms/spare-report/spr13601/retrieve.json', {}, 2000);
  return [rows.map(r => [r.partNo, r.partNm, r.supersedingPartNo, r.supersedingPartNm]), total]; };
window.yKits = async () => {
  const u = yUser();
  const [kits] = await yPaged('/ydms/spare-report/spr11801/retrieve.json', {}, 100);
  const out = [];
  for (const k of kits) {
    const j = await yCall('/ydms/spare-report/spr11802/initial.json', {companyCode: 'IYM', dealerCd: u.pymidolDealerCd, networkCd: ' ', kitCd: k.kitCd});
    out.push({kit: k.kitCd, name: k.kitNm, from: k.effectiveDate, to: k.expiryDate,
              parts: ((j.data || j).resultList || []).map(p => [p.partNo, p.partNm, p.kitQty])});
  }
  return out;
};
window.yPI = async to => (await yCall('/spare-parts/spm1301/retrieve', {sysOwnerCd: 'IYM', modelCd: null, releaseDate: '19900101', releaseDateTo: to, user: yUser()})).data || [];
window.yCirc = async () => (await yCall('/spare-parts/pza0301/initial', {sysOwnerCd: 'IYM', moduleCd: '2', userModel: yUser()})).data || [];
window.yDSA = async () => (await yCall('/spare-parts/svt0801/retrieve', {user: yUser(), serviceDaNo: null, modelCd: null, fromDate: null,
  toDate: null, priority: null, serviceActionType: null, status: null})).data || [];
window.yGetPI = async (newFile, file) => { const j = await yCall('/spare-parts/spm1301/downloadAttachment', {user: yUser(), downloadFile: newFile, displayFile: file});
  return (j.data && j.data.attachment) || null; };
window.yGetDSA = async (newFile, file) => { const j = await yCall('/spare-parts/svt0801/downloadAttachment', {fileDownloadName: newFile, fileDisplayName: file, user: yUser()});
  return (j.data && j.data.attachment) || null; };
window.yGetCirc = async (name, type) => {
  // names with characters like '–' only work encoded; plain names work as they are
  const tries = /[^\x20-\x7e]/.test(name) ? [encodeURIComponent(name), name] : [name, encodeURIComponent(name)];
  for (const n of tries) { const b = await yX('/spare-parts/pza0301/download-file', {fileDisplayName: n, fileType: type}, true).catch(() => null); if (b) return b; }
  return null;
};
0;   // page.evaluate would otherwise call the last function defined above
"""


def num(v):
    try: return float(str(v).replace(',', '').strip())
    except (TypeError, ValueError): return None


def iso(d):
    """'16/10/2026' or '20261016' -> '2026-10-16'"""
    d = str(d or '').strip()
    m = re.match(r'(\d{2})/(\d{2})/(\d{4})$', d)
    if m: return f'{m.group(3)}-{m.group(2)}-{m.group(1)}'
    m = re.match(r'(\d{4})(\d{2})(\d{2})$', d)
    if m: return f'{m.group(1)}-{m.group(2)}-{m.group(3)}'
    return d


def safe(s): return re.sub(r'\s+', ' ', re.sub(r'[\\/:*?"<>|\x00-\x1f]+', '-', str(s or ''))).strip().rstrip('.')[:150]


def is_catalogue(r):
    return bool(re.search(r'cata|cate|catlo', f"{r.get('partInfoListNoForDisplay') or ''} {r.get('details') or ''}", re.I))


class OtpPage:
    """Small page on the server where Sam types the OTP pymidol sends to the dealer's phone. Nothing else is asked."""
    def __init__(self, port):
        from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
        from urllib.parse import parse_qs
        self.code = secrets.token_urlsafe(12)
        self.msg, self.answer = 'Enter the 4 digit OTP sent to the dealer phone.', None
        me = self

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a): pass
            def page(self):
                b = (f'<html><head><meta name="viewport" content="width=device-width"></head><body style="font-family:sans-serif;padding:24px">'
                     f'<h3>Enter code</h3><p>{html.escape(me.msg)}</p>'
                     f'<form method="post"><input name="code" inputmode="numeric" autocomplete="off" autofocus style="font-size:24px;width:180px"> '
                     f'<button style="font-size:20px">Submit</button></form></body></html>').encode()
                self.send_response(200); self.send_header('Content-Type', 'text/html; charset=utf-8')
                self.send_header('Cache-Control', 'no-store'); self.end_headers(); self.wfile.write(b)
            def do_GET(self):
                if self.path.split('?')[0].rstrip('/') != f'/yamaha-otp/{me.code}':
                    self.send_response(404); self.end_headers(); return
                self.page()
            def do_POST(self):
                if self.path.split('?')[0].rstrip('/') != f'/yamaha-otp/{me.code}':
                    self.send_response(404); self.end_headers(); return
                q = parse_qs(self.rfile.read(int(self.headers.get('Content-Length') or 0)).decode())
                me.answer = re.sub(r'\D', '', (q.get('code') or [''])[0]) or None
                me.msg = 'OTP received, logging in. Refresh this page in 30 seconds.' if me.answer else 'The box was empty. Enter the OTP again.'
                self.page()
        self.srv = ThreadingHTTPServer((env('CODE_BIND', '0.0.0.0'), port), H)   # Bike Pharma server: 127.0.0.1, the website passes /yamaha-otp/ on
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()

    def close(self):
        self.srv.shutdown()


class Yamaha:
    name = 'Yamaha'
    pdfs = False                 # pymidol has its own catalogue PDFs: they are saved as they are, nothing is drawn
    price_label = 'MRP from pymidol SPR128'

    def __init__(self):
        self.dir = data_dir('yamaha')
        self.state_path = os.path.join(self.dir, 'state.json')
        self.state = load_json(self.state_path, None)
        if self.state is None:
            self.state = {'mrp': {}, 'sup': [], 'kits': {}, 'models': {}, 'pi': {}, 'circ': {}, 'dsa': {}, 'pending': {}}
            if os.path.exists(SEED):
                # 23 Oct 2026 (the Mac "Yamaha Catalogue" folder): the first run reports and downloads only what came after it
                s = json.load(gzip.open(SEED, 'rt', encoding='utf-8'))
                self.state['mrp'] = {pn: {'name': v[0], 'mrp': v[1]} for pn, v in s['mrp'].items()}
                self.state['sup'] = s['sup']
                self.state['kits'] = {k: {'name': '', 'parts': [[p, '', q] for p, q in v]} for k, v in s['kits'].items()}
                self.state['models'] = {c: '' for c in s['models']}
                self.state['seed'] = {'date': s['date'], 'pi_models': s['pi_models'], 'pi_bulletins': s['pi_bulletins'],
                                      'circulars': s['circulars'], 'dsa': s['dsa']}

    # ---------- login ----------
    def shot(self, page, name):
        p = os.path.join(data_dir('logs'), f'yamaha_{name}_{today()}.png')
        try: page.screenshot(path=p)
        except Exception: pass
        return p

    @staticmethod
    def logged_in(page):
        try:
            if not page.url.startswith(BASE) or '/auth/' in page.url:
                return False
            page.evaluate(JS)
            return page.evaluate("async () => !!yUser() && (await yModels()).length > 0")
        except Exception:
            return False

    def wait_any(self, page, secs=60):
        """-> 'dashboard', 'prelogin', 'password', 'otp', 'vacant', 'microsoft' or ''"""
        end = time.time() + secs
        while time.time() < end:
            u = page.url
            if 'microsoft' in u or 'saml2' in u:
                return 'microsoft'
            if page.locator('input[formcontrolname="otp"]:visible').count():
                return 'otp'
            if page.locator('[formcontrolname="vacantUserId"]:visible').count():
                return 'vacant'
            if page.locator('input[formcontrolname="password"]:visible').count():
                return 'password'
            if page.locator('input[formcontrolname="userId"]:visible').count():
                return 'prelogin'
            if re.search(r'/(pymidolIV|ydms)/', u) and '/auth/' not in u and page.evaluate("() => !!localStorage.getItem('user')"):
                return 'dashboard'
            page.wait_for_timeout(1000)
        return ''

    def toast(self, page):
        loc = page.locator('.toast-message:visible, .toast-error:visible, .text-danger:visible')
        try:
            return ' '.join(t.strip() for t in loc.all_inner_texts() if t.strip())[:200]
        except Exception:
            return ''

    def login(self, page):
        page.goto(BASE + '/', wait_until='load')
        where = self.wait_any(page, 90)
        if where == 'dashboard' and self.logged_in(page):
            log.info('Yamaha: still logged in (saved browser profile)'); return
        user, pw = env('YAMAHA_USER', required=True), env('YAMAHA_PASS', required=True)
        otp = None
        try:
            for step in range(8):
                if where == 'dashboard':
                    page.wait_for_timeout(3000)
                    if self.logged_in(page):
                        log.info('Yamaha login ok'); return
                    raise SystemExit(f"Yamaha: login ke baad bhi pymidol ka data nahi aaya ({self.shot(page, 'after_login')})")
                if where == 'prelogin':
                    sel = page.locator('select[formcontrolname="companyCode"]')
                    if sel.count():
                        labels = sel.locator('option').all_inner_texts()
                        lab = next((l for l in labels if re.search(env('YAMAHA_COMPANY', 'Yamaha'), l, re.I)), None)
                        if lab: sel.select_option(label=lab)
                    page.fill('input[formcontrolname="userId"]', user)
                    page.locator('button', has_text=re.compile(r'^\s*Sign In\s*$', re.I)).first.click()
                elif where == 'password':
                    page.fill('input[formcontrolname="password"]', pw)
                    page.locator('button', has_text=re.compile(r'^\s*Sign in\s*$', re.I)).first.click()
                elif where == 'otp':
                    if otp is None:
                        port = int(env('YAMAHA_OTP_PORT', '8767'))
                        otp = OtpPage(port)
                        link = f"{env('YAMAHA_OTP_LINK_BASE', f'http://SERVER-IP:{port}').rstrip('/')}/yamaha-otp/{otp.code}"
                        wait = int(env('YAMAHA_OTP_WAIT', '180'))
                        log.warning(f'Yamaha needs the OTP sent to the dealer phone. Open this link and enter it:\n{link}\n\nWaiting up to {wait} minutes.')
                        send_code_link(link)
                    end = time.time() + 60 * int(env('YAMAHA_OTP_WAIT', '180'))
                    while not otp.answer and time.time() < end:
                        time.sleep(2)
                    if not otp.answer:
                        raise SystemExit('Yamaha: OTP nahi mila, sync nahi chala.')
                    code, otp.answer = otp.answer, None
                    page.fill('input[formcontrolname="otp"]', code)
                    page.locator('button', has_text=re.compile(r'Submit', re.I)).first.click()
                elif where == 'vacant':
                    raise SystemExit(f"Yamaha: pymidol 'Login As / Assigned Vacant User' puch raha hai ({self.shot(page, 'vacant')}). "
                                     'Ek baar haath se login karke dekho ye kya hai, phir batao.')
                elif where == 'microsoft':
                    raise SystemExit('Yamaha: ye ID Microsoft (Yamaha staff) wala login maang rahi hai. Dealer ID (jaise 184.90) .env me daalo.')
                else:
                    raise SystemExit(f"Yamaha login page nahi pehchana ({page.url}, {self.shot(page, 'login_unknown')})")
                prev = where
                page.wait_for_timeout(2500)
                where = self.wait_any(page, 60)
                if where == prev:
                    msg = self.toast(page) or 'koi jawab nahi'
                    self.shot(page, 'login_try')
                    if prev == 'otp':
                        if otp: otp.msg = f'OTP did not work ({msg}). Enter it again.'
                        continue
                    raise SystemExit(f'Yamaha login failed: {msg}. YAMAHA_USER / YAMAHA_PASS check karo.')
                if otp and where == 'dashboard':
                    otp.msg = 'Logged in. You can close this page.'
        finally:
            if otp: otp.close()
        raise SystemExit('Yamaha login nahi hua (bahut steps).')

    def call(self, page, expr, arg=None):
        try:
            return page.evaluate(expr, arg)
        except Exception as e:
            if re.search(r'HTTP 4\d\d', str(e)):            # pymidol says the file is not there: reloading won't help
                raise
            log.warning('Yamaha call failed (%s), reloading page', str(e)[:120])
            page.wait_for_timeout(5000)
            page.goto(BASE + '/', wait_until='load'); self.wait_any(page, 60)
            if not self.logged_in(page):
                raise RuntimeError('Yamaha login beech me khatam ho gaya')
            return page.evaluate(expr, arg)

    # ---------- sync ----------
    def run(self, full=False, report=None):
        rep = report or Report(self.name)
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            ctx = p.chromium.launch_persistent_context(data_dir('yamaha', 'browser'), args=['--disable-dev-shm-usage', '--js-flags=--max-old-space-size=4096'], headless=env('YAMAHA_HEADFUL', '') == '',
                                                       viewport={'width': 1500, 'height': 900})
            page = ctx.pages[0] if ctx.pages else ctx.new_page()
            if env('YAMAHA_FRESH_LOGIN'):                    # test: forget the saved login, do the full login
                ctx.clear_cookies(); page.goto(BASE + '/', wait_until='load'); page.evaluate('() => { localStorage.clear(); sessionStorage.clear(); }')
                log.info('Yamaha: saved login cleared (YAMAHA_FRESH_LOGIN)')
            page.set_default_timeout(300000)
            try:
                self.login(page)
                page.evaluate(JS)
                models = self.call(page, '() => yModels()')
                mrp, mtotal = self.call(page, '() => yMrp()')
                sup, stotal = self.call(page, '() => ySup()')
                log.info('Yamaha: %d models, %d/%d MRP rows, %d/%d superseding rows', len(models), len(mrp), mtotal, len(sup), stotal)
                if len(mrp) < 0.9 * len(self.state['mrp']) or (mtotal > 0 and len(mrp) < mtotal):
                    raise RuntimeError(f'Yamaha MRP list me sirf {len(mrp)} parts aaye (pichli baar {len(self.state["mrp"])}) - kuch nahi badla')
                kits = self.call(page, '() => yKits()')
                pi = self.call(page, 'd => yPI(d)', today().replace('-', ''))
                circ = self.call(page, '() => yCirc()')
                dsa = self.call(page, '() => yDSA()')
                self.merge_parts(rep, models, mrp, sup, kits)
                todo = self.merge_docs(rep, pi, circ, dsa, full)
                save_json(self.state_path, self.state)
                self.download(page, rep, todo)
            finally:
                ctx.close()
        self.state['last_run'] = today()
        self.state.pop('seed', None)
        save_json(self.state_path, self.state)
        self.write_files(rep)
        return rep, []

    def merge_parts(self, rep, models, mrp, sup, kits):
        st = self.state
        first = not st['mrp']
        # model codes
        newm = {(m.get('modelCd') or '').strip(): (m.get('modelNm') or '').strip() for m in models if m.get('modelCd')}
        if st['models']:
            for c in sorted(set(newm) - set(st['models'])):
                rep.add('Naya model code', f'{c} {newm[c]}')
        st['models'] = {**{c: n for c, n in st['models'].items() if n}, **newm}
        # MRP
        new = {}
        for pn, nm, m in mrp:
            pn = (pn or '').strip().upper()
            if pn: new[pn] = {'name': (nm or '').strip(), 'mrp': num(m)}
        self.changes = []
        if not first:
            for pn in sorted(set(new) - set(st['mrp']))[:1000]:
                rep.add('Naya part number', f"{pn} {new[pn]['name']} MRP {new[pn]['mrp']}")
            for pn, v in sorted(new.items()):
                o = st['mrp'].get(pn)
                if o and v['mrp'] is not None and o.get('mrp') is not None and abs(v['mrp'] - o['mrp']) >= 0.5:
                    self.changes.append({'part_no': pn, 'old_mrp': o['mrp'], 'new_mrp': v['mrp'], 'name': v['name']})
            for c in self.changes[:300]:
                rep.add('MRP badla', f"{c['part_no']} {c['name']}: Rs {c['old_mrp']:g} -> Rs {c['new_mrp']:g}")
            if len(self.changes) > 300:
                rep.add('MRP badla', f'... aur {len(self.changes) - 300} (sab Excel me)')
        self.new_parts = sorted(set(new) - set(st['mrp'])) if not first else []
        self.gone = sorted(set(st['mrp']) - set(new)) if not first else []
        if self.gone:
            rep.add('MRP list se hata', f'{len(self.gone)} parts (list Excel me)')
        st['mrp'] = new
        # superseding
        pairs = sorted({((a or '').strip().upper(), (b or '').strip().upper()) for a, _, b, _ in sup if a and b})
        old = {tuple(x) for x in st['sup']}
        self.new_sup = [p for p in pairs if p not in old] if old else []
        for a, b in self.new_sup[:300]:
            rep.add('Naya purana -> naya', f"{a} -> {b} {(new.get(b) or {}).get('name', '')}")
        self.sup_names = {(a or '').strip().upper(): (an or '').strip() for a, an, b, bn in sup}
        self.sup_names.update({(b or '').strip().upper(): (bn or '').strip() for a, an, b, bn in sup})
        st['sup'] = [list(p) for p in pairs]
        # kits
        nk = {k['kit']: {'name': k['name'], 'from': k['from'], 'to': k['to'], 'parts': k['parts']} for k in kits}
        if st['kits']:
            for k in sorted(set(nk) - set(st['kits'])):
                rep.add('Naya kit', f"{k} {nk[k]['name']} ({len(nk[k]['parts'])} parts)")
            for k in sorted(set(nk) & set(st['kits'])):
                a = sorted((p[0], num(p[2])) for p in nk[k]['parts'])
                b = sorted((p[0], num(p[2])) for p in st['kits'][k]['parts'])
                if a != b:
                    rep.add('Kit ke andar badla', f"{k} {nk[k]['name']}")
            for k in sorted(set(st['kits']) - set(nk)):
                rep.add('Kit hata', k)
        st['kits'] = nk
        self.push(self.changes)

    def merge_docs(self, rep, pi, circ, dsa, full):
        """new Part Information (catalogues + bulletins), circulars and service actions -> list of files to download"""
        st = self.state
        seed = st.get('seed')
        todo = dict(st.get('pending') or {})
        def known(kind, key, row_date, seed_hit):
            if key in st[kind]: return True
            if seed is not None:
                # first run: what pymidol had on the seed day is already in the Mac folder
                return seed_hit or row_date < seed['date']
            return not st.get('last_run') and not full     # very first run without a seed: just remember, download nothing
        sm, sb = (set(seed['pi_models']), set(seed['pi_bulletins'])) if seed else (set(), set())
        for r in pi:
            k = r.get('partInfoListNo')
            if not k: continue
            d = iso(r.get('partInfoReleaseDate'))
            cat = is_catalogue(r)
            row = {'model': r.get('model') or '', 'name': (r.get('modelNm') or '').strip(), 'date': d, 'cat': cat,
                   'title': (r.get('details') or r.get('partInfoListNoForDisplay') or '').strip(), 'file': r.get('file') or '', 'new': r.get('newFile') or ''}
            if not known('pi', k, d, k in sb or f"{row['model']}|{d}" in sm):
                todo['pi|' + k] = row
                rep.add('Naya model catalogue (PDF)' if cat else 'Naya Part Information bulletin', f"{row['model']} {row['name']} {row['title']} ({d})")
            st['pi'][k] = row
        sc = set(seed['circulars']) if seed else set()
        for r in circ:
            k = f"{r.get('description') or ''}|{r.get('publishDate') or ''}"
            d = iso(r.get('publishDate'))
            row = {'title': (r.get('description') or '').strip(), 'date': d, 'file': r.get('circularFile') or '', 'new': r.get('newCircularFile') or '',
                   'info': r.get('infoFile') or '', 'newinfo': r.get('newInfoFile') or ''}
            if not known('circ', k, d, k in sc):
                todo['circ|' + k] = row
                rep.add('Naya circular', f"{row['title']} ({r.get('publishDate')})")
            st['circ'][k] = row
        sd = set(seed['dsa']) if seed else set()
        for r in dsa:
            k = str(r.get('serviceDaNo') or '').strip()
            if not k: continue
            d = iso(r.get('releaseDate'))
            files = [[(r.get(f'file{i}') or '').strip(), (r.get(f'newFile{i}') or '').strip()] for i in (1, 2, 3)]
            row = {'no': r.get('serviceActionNo'), 'type': r.get('serviceActionType'), 'priority': r.get('priority'),
                   'models': r.get('modelCd') or '', 'problem': (r.get('problemDesc') or '').strip(), 'date': d,
                   'files': [f for f in files if f[0] and f[1]]}
            if not known('dsa', k, d, k in sd):
                todo['dsa|' + k] = row
                rep.add('Naya Service Action (DSA)', f"DSA {k}: {row['problem']} ({row['models'][:60]}) {d}")
            st['dsa'][k] = row
        st['pending'] = todo
        return todo

    def out(self, *parts):
        d = os.path.join(data_dir('output', 'Yamaha'), *parts)
        os.makedirs(d, exist_ok=True)
        return d

    def save(self, folder, name, b64):
        b = base64.b64decode(b64)
        p = os.path.join(self.out('PDF', folder), name)
        open(p, 'wb').write(b)                               # bytes exactly as pymidol gives them
        return p

    def download(self, page, rep, todo):
        st = self.state
        n, bad = 0, []
        for key, row in sorted(todo.items(), key=lambda kv: kv[1].get('date') or ''):
            kind = key.split('|', 1)[0]
            try:
                if kind == 'pi':
                    b = self.call(page, '([n, f]) => yGetPI(n, f)', [row['new'], row['file']])
                    ext = os.path.splitext(row['file'])[1] or '.pdf'
                    if row['cat']:
                        name = safe(f"{row['model']} {row['name']} ({row['date']})") + ext
                        folder = 'Models'
                    else:
                        name = safe(f"{row['date']} {row['title'] or os.path.splitext(row['file'])[0]}") + ext
                        folder = 'Part Information Bulletins'
                    ok = bool(b) and self.save(folder, name, b)
                elif kind == 'circ':
                    b = self.call(page, '([n, t]) => yGetCirc(n, t)', [row['new'], 'circular']) if row['new'] else None
                    ok = bool(b) and self.save('Circulars (pymidol)', safe(row['file'] or row['title'] + '.pdf'), b)
                    if row['newinfo']:
                        bi = self.call(page, '([n, t]) => yGetCirc(n, t)', [row['newinfo'], 'info'])
                        if bi: self.save('Circulars (pymidol)', safe('Extra - ' + row['info']), bi)
                else:
                    left = []
                    for disp, new in row['files']:
                        try:
                            b = self.call(page, '([n, f]) => yGetDSA(n, f)', [new, disp])
                        except RuntimeError:
                            raise
                        except Exception as e:
                            log.warning('Yamaha DSA %s: %s', disp, e); b = None
                        if b: self.save('Service Actions (DSA)', safe(disp), b)
                        else: left.append([disp, new])
                    ok = not left
                    row['files'] = left                         # next run tries only the files still missing
            except RuntimeError:
                raise
            except Exception as e:
                log.warning('Yamaha download %s: %s', key, e); ok = False
            if ok:
                n += 1; todo.pop(key, None)
            else:
                row['tries'] = row.get('tries', 0) + 1
                what = {'pi': f"Part Information {row.get('model', '')} {row.get('title', '')}",
                        'circ': f"Circular {row.get('title', '')}",
                        'dsa': f"DSA {key.split('|', 1)[1]}: {', '.join(f[0] for f in row.get('files', []))}"}[kind]
                if row['tries'] >= 3:
                    todo.pop(key, None)
                    bad.append(f'{what} (3 baar try kiya, pymidol file nahi de raha)')
                else:
                    bad.append(f'{what} (agle run me dobara)')
            st['pending'] = todo
            save_json(self.state_path, st)
            time.sleep(float(env('YAMAHA_DELAY', '0.3')))
        if n: rep.add('Files download hui', f'{n} (output/Yamaha/PDF)')
        for b in bad: rep.add('Download nahi hua', b)

    # ---------- files ----------
    def write_files(self, rep):
        import openpyxl
        st = self.state
        out = self.out()
        with open(os.path.join(out, f'mrp_changes_{today()}.csv'), 'w', newline='', encoding='utf-8') as f:
            w = csv.writer(f); w.writerow(['part_no', 'old_mrp', 'new_mrp', 'name'])
            for c in self.changes: w.writerow([c['part_no'], c['old_mrp'], c['new_mrp'], c['name']])
        nxt = collections.defaultdict(list); prv = collections.defaultdict(list)
        for a, b in st['sup']:
            nxt[a].append(b); prv[b].append(a)
        in_kit = collections.defaultdict(list)
        for k, v in st['kits'].items():
            for p in v['parts']: in_kit[p[0]].append(k)
        name = lambda pn: (st['mrp'].get(pn) or {}).get('name') or self.sup_names.get(pn, '')
        mrp = lambda pn: (st['mrp'].get(pn) or {}).get('mrp')
        def final(pn, seen=()):
            if pn not in nxt or pn in seen: return pn
            return final(nxt[pn][-1], seen + (pn,))
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        def sheet(title, head, rows, widths):
            ws = wb.create_sheet(title); ws.append(head)
            for r in rows: ws.append(r)
            for i, wd in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = wd
            ws.freeze_panes = 'A2'
        new_set = set(self.new_parts)
        sheet('All Parts A-Z', ['Part No', 'Naam', 'MRP (Rs)', 'Naya part no.', 'Purane part no.', 'Kit me', 'Is hafte naya'],
              [[pn, v['name'], v['mrp'], ', '.join(nxt.get(pn, [])), ', '.join(prv.get(pn, [])), ', '.join(in_kit.get(pn, [])),
                'Haan' if pn in new_set else ''] for pn, v in sorted(st['mrp'].items())],
              [16, 34, 10, 16, 22, 16, 8])
        sheet('MRP Badla', ['Part No', 'Naam', 'Purana MRP', 'Naya MRP'],
              [[c['part_no'], c['name'], c['old_mrp'], c['new_mrp']] for c in self.changes], [16, 34, 11, 11])
        sheet('Naye Parts', ['Part No', 'Naam', 'MRP'], [[pn, name(pn), mrp(pn)] for pn in self.new_parts], [16, 34, 10])
        sheet('List se Hata', ['Part No'], [[pn] for pn in self.gone], [16])
        sheet('Kits', ['Kit', 'Kit naam', 'Kab se', 'Kab tak', 'Part No', 'Naam', 'Qty', 'MRP'],
              [[k, v['name'], v.get('from'), v.get('to'), p[0], p[1] or name(p[0]), p[2], mrp(p[0])]
               for k, v in sorted(st['kits'].items()) for p in v['parts']], [14, 26, 11, 11, 16, 30, 6, 10])
        sheet('Model Codes', ['Model Code', 'Model'], sorted(st['models'].items()), [12, 30])
        wb.save(os.path.join(out, 'Yamaha_All_Parts.xlsx'))
        # old <-> new, both ways
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Purana -> Naya'
        ws.append(['Purana Part No', 'Purana Naam', 'Purana MRP', 'Naya Part No', 'Naya Naam', 'Naya MRP', 'Aakhri (final) Part No', 'Final MRP', 'Is hafte naya'])
        newp = {tuple(p) for p in self.new_sup}
        for a, b in st['sup']:
            f = final(b)
            ws.append([a, name(a), mrp(a), b, name(b), mrp(b), f, mrp(f), 'Haan' if (a, b) in newp else ''])
        ws2 = wb.create_sheet('Naya -> Purane')
        ws2.append(['Naya Part No', 'Naam', 'MRP', 'Purane Part No'])
        for b, olds in sorted(prv.items()):
            ws2.append([b, name(b), mrp(b), ', '.join(sorted(olds))])
        for w_ in (ws, ws2):
            w_.freeze_panes = 'A2'
            for col in 'ABCDEFGHI': w_.column_dimensions[col].width = 18
        wb.save(os.path.join(out, 'Yamaha_Purana_Naya.xlsx'))
        # what pymidol has: Part Information, circulars, service actions
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        sheet('Model Catalogues', ['Model Code', 'Model', 'Release Date', 'PI No', 'pymidol file'],
              sorted([[r['model'], r['name'], r['date'], k, r['file']] for k, r in st['pi'].items() if r['cat']], key=lambda x: (x[0], x[2])),
              [12, 30, 12, 18, 30])
        sheet('PI Bulletins', ['PI No', 'Bulletin', 'Release Date', 'Model Code', 'Model', 'pymidol file'],
              sorted([[k, r['title'], r['date'], r['model'], r['name'], r['file']] for k, r in st['pi'].items() if not r['cat']], key=lambda x: x[2], reverse=True),
              [18, 50, 12, 12, 24, 30])
        sheet('Circulars', ['Circular', 'Publish Date', 'PDF', 'Extra file'],
              sorted([[r['title'], r['date'], r['file'], r['info']] for r in st['circ'].values()], key=lambda x: x[1], reverse=True), [60, 12, 50, 40])
        sheet('Service Actions', ['DSA No', 'Service Action No', 'Type', 'Priority', 'Model Codes', 'Problem', 'Release Date', 'Files'],
              sorted([[k, r['no'], r['type'], r['priority'], r['models'], r['problem'], r['date'], ' | '.join(f[0] for f in r['files'])]
                      for k, r in st['dsa'].items()], key=lambda x: x[6], reverse=True), [8, 8, 11, 8, 30, 60, 12, 50])
        wb.save(os.path.join(out, 'Yamaha_Documents_Index.xlsx'))

    def push(self, changes):
        url = env('YAMAHA_MRP_WEBHOOK')
        if not url or not changes:
            return
        import requests
        h = {'Content-Type': 'application/json'}
        if env('YAMAHA_MRP_WEBHOOK_TOKEN'): h['Authorization'] = 'Bearer ' + env('YAMAHA_MRP_WEBHOOK_TOKEN')
        for i in range(0, len(changes), 2000):
            r = requests.post(url, headers=h, timeout=120, data=json.dumps({'source': 'yamaha_pymidol', 'date': today(), 'changes': changes[i:i + 2000]}))
            log.info('webhook %s: HTTP %s', url, r.status_code)

    def rebuild(self):
        """python sync.py yamaha --rebuild: Excel files again from saved data, no login"""
        rep = Report(self.name)
        self.changes, self.new_parts, self.gone, self.new_sup, self.sup_names = [], [], [], [], {}
        self.write_files(rep)
        rep.add('Files update hui', 'Yamaha_All_Parts.xlsx, Yamaha_Purana_Naya.xlsx, Yamaha_Documents_Index.xlsx dobara bane')
        return rep

    def prices(self):
        return {pn: v['mrp'] for pn, v in self.state['mrp'].items() if v.get('mrp')}
