"""Every brand's catalogue for the website's Find Part, from Sam's package All Catalogues/BikeSurgeon_Catalog_Import (3 Oct 2026).

For each brand two files go to public/cat/ (tracked in git, staff only on the server):
  <slug>.json     the app's parts list: {brand, src, bk, models, cur, cc, vars, vmod, asm, desc, d, fw, mrp, phf, phz}
                  d row = [book, -1, group, asm, ref, part no., desc, qty, remarks, not serviceable, 0, [variants]]
  <slug>-dg.json  {f: ['pf:<path>:<offset>:<length>'], p: {variant: {group: f}}, dir, files: [[zip, size]], pick}
                  each diagram is read straight out of the brand's own zip on the server (stored zips: the picture's bytes sit
                  at one offset), so it is the original jpg/png byte for byte (Sam: "diagram aur image ki quality thoda bhi kam
                  nhi hona he"). The zips go to the server unchanged from Admin > More > Catalogue files (dir + files).
Part photos (Hero, Royal Enfield) are read the same way: phf {part no.: [zip, offset, length, w, h]}, phz [zip paths].
MRP only (never dealer / NDP prices).

usage: python3 -I brands_build.py <out dir> [hero re yamaha honda tvs bajaj suzuki]
Reads (read only) the files staged from Sam's Mac (src.py) and the Hero / RE model rules of the prototype (hr_pkg_build.py);
Suzuki: suzuki-pages.json (sz_parse.py, from Sam's folder All Catalogues > Suzuki)."""
import collections, itertools, json, os, re, sys, zipfile
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from src import IDX, IMP, rows_zip, rows_csv, zidx

OUT = sys.argv[1]
WHICH = sys.argv[2:] or ['hero', 're', 'yamaha', 'honda', 'tvs', 'bajaj', 'suzuki']
PKG = 'All Catalogues > BikeSurgeon_Catalog_Import > '

# Hero / Royal Enfield model names, editions and emission norms: the same rules as the prototype (Sam saw these)
_src = open('/mnt/project-files/designs/prototype-build/bj/hr_pkg_build.py').read()
HR = {'__name__': 'hr_rules', '__file__': 'hr_pkg_build.py'}
exec(compile(_src[:_src.index('# ---------- the packed file')].replace("WHAT = sys.argv[1] if len(sys.argv) > 1 else 'all'", "WHAT = 'x'"), 'hr_pkg_build.py', 'exec'), HR)
nat = HR['nat']


def jdump(path, o):
    tmp = path + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(o, f, ensure_ascii=False, separators=(',', ':'))
    os.replace(tmp, path)


class Zips:
    """the brand's picture zips as they are on Sam's computer: entry -> 'pf:<dir>/<zip>:<offset>:<length>' (stored entries only)"""
    def __init__(self, dirn, names, key=lambda n: n):
        self.dir, self.names, self.at, self.size, self.used = dirn, names, {}, {}, set()
        self.dims = {}
        for z in names:
            ix, size = zidx(z)
            self.size[z] = size
            dp = os.path.join(IDX, z + '.dims.json')
            dims = json.load(open(dp)) if os.path.exists(dp) else {}
            for n, (meth, off, csize, usize) in ix.items():
                k = key(n)
                if k is None:
                    continue
                assert k not in self.at or self.at[k][0] != z, ('same name twice in one zip', z, n)
                if k in self.at:
                    self.at[k] = None   # in two zips: ambiguous, never used
                    continue
                self.at[k] = (z, meth, off, csize, usize, dims.get(n))

    def link(self, k):
        e = self.at.get(k)
        if not e:
            return None
        z, meth, off, csize, usize, _ = e
        assert meth == 0 and csize == usize and off + usize <= self.size[z], (z, k)
        self.used.add(z)
        return 'pf:%s/%s:%d:%d' % (self.dir, z, off, usize)

    def files(self):
        return [[z, self.size[z]] for z in self.names if z in self.used]


def merge_mrp(pairs):
    """[(part no., price)] -> {part no.: Rs}; a part no. listed with two different prices gets none"""
    c = collections.defaultdict(set)
    for pn, p in pairs:
        try:
            v = float(str(p).replace(',', '').strip() or 0)
        except ValueError:
            continue
        if pn and v > 0:
            c[pn].add(v)
    two = [pn for pn, s in c.items() if len(s) > 1]
    out = {pn: (int(next(iter(s))) if next(iter(s)) == int(next(iter(s))) else next(iter(s))) for pn, s in c.items() if len(s) == 1}
    return out, two


def pack(brand, slug, src, order, cur, rows, fw, bk, cc, zp, mrp=None, photos=None, pick='', by_group=False):
    """rows: dicts m (model), v (edition label), y (year or 0), nb (book), g (group), asm, ref, pn, desc, qty, rem, ns, dg (picture key)
    -> <slug>.json + <slug>-dg.json. Lines that differ only by edition are one line listing the editions. by_group: lines in figure
    order (then ref) rather than model by model, so every model's sections come in its book's order even where models share lines
    (or a sort key of its own for the (line, editions) items)."""
    have = {x['m'] for x in rows}
    ms = [m for m in order if m in have] + sorted(have - set(order))
    cur_ms = [m for m in ms if m in cur]
    ms = cur_ms + [m for m in ms if m not in cur]
    vy, vm = {}, {}
    for x in rows:
        vy[x['v']] = max(vy.get(x['v'], 0), x.get('y') or 0)
        assert vm.setdefault(x['v'], x['m']) == x['m'], ('edition in two models', x['v'])
    vl = sorted(vm, key=lambda v: (ms.index(vm[v]), -vy[v], nat(v)))
    vi = {v: i for i, v in enumerate(vl)}
    lines, asm, desc = collections.OrderedDict(), {}, {}
    for x in sorted(rows, key=lambda x: (ms.index(x['m']), -(x.get('y') or 0), nat(x['v']), nat(x['g']), nat(x['ref']), x['pn'])):
        k = (x['nb'], x['g'], x['asm'], x['ref'], x['pn'], x['desc'], x['qty'], x['rem'], 1 if x['ns'] else 0)
        lines.setdefault(k, set()).add(vi[x['v']])
    items = list(lines.items())
    if by_group:
        items.sort(key=by_group if callable(by_group) else lambda kv: (kv[0][0], nat(kv[0][1]), nat(kv[0][3])))
    d = [[nb, -1, g, asm.setdefault(a, len(asm)), ref, pn, desc.setdefault(ds, len(desc)), qty, rem, ns, 0, sorted(vs)]
         for (nb, g, a, ref, pn, ds, qty, rem, ns), vs in items]
    inv = lambda m: [k for k, _ in sorted(m.items(), key=lambda t: t[1])]
    o = {'brand': brand, 'src': [src], 'bk': bk, 'models': ms, 'cur': len(cur_ms), 'cc': {m: cc[m] for m in ms if cc.get(m)},
         'sheets': [], 'figs': [], 'vars': vl, 'vmod': {v: vm[v] for v in vl}, 'asm': inv(asm), 'desc': inv(desc), 'd': d, 'fw': fw}
    if mrp:
        pns = {r[5] for r in d}
        o['mrp'] = {pn: p for pn, p in mrp.items() if pn in pns}
    # diagrams: per edition and group, the picture most of its lines name
    need = collections.defaultdict(collections.Counter)
    for x in rows:
        if x.get('dg'):
            need[(vi[x['v']], x['g'])][x['dg']] += 1
    f, fi, p, lost, two = [], {}, collections.defaultdict(dict), set(), 0
    for (v, g), c in sorted(need.items()):
        two += len(c) > 1
        u = zp.link(c.most_common(1)[0][0])
        if not u:
            lost.add(c.most_common(1)[0][0])
            continue
        if u not in fi:
            fi[u] = len(f)
            f.append(u)
        p[str(v)][g] = fi[u]
    # part photos
    if photos:
        phz, phf = [], {}
        for pn, k in photos.items():
            u = zp.link(k)
            if not u:
                continue
            _, path, off, ln = u.split(':')
            if path not in phz:
                phz.append(path)
            wh = zp.at[k][5] or [250, 250]
            phf[pn] = [phz.index(path), int(off), int(ln)] + list(wh)
        o['phf'], o['phz'] = phf, phz
    dg = {'f': f, 'p': p, 'dir': zp.dir, 'files': zp.files(), 'spare': [], 'pick': pick}
    for name, obj in ((slug + '.json', o), (slug + '-dg.json', dg)):
        jdump(os.path.join(OUT, name), obj)
    groups = {(r[11][0], r[2]) for r in d}
    print('%s: %d models (%d current), %d editions, %d lines, %d part no., MRP for %d | diagrams %d (edition+group pages %d, %d with 2+ pictures, %d not in the zips)%s | files %s | %.1f MB + %.1f MB'
          % (brand, len(ms), len(cur_ms), len(vl), len(d), len({r[5] for r in d if r[5]}), len(o.get('mrp', {})), len(f), len(need), two, len(lost),
             ' | photos %d' % len(o['phf']) if photos else '', ', '.join('%s %.2f GB' % (z, s / 1e9) for z, s in dg['files']),
             os.path.getsize(os.path.join(OUT, slug + '.json')) / 1e6, os.path.getsize(os.path.join(OUT, slug + '-dg.json')) / 1e6))
    if lost:
        print('   pictures not found:', sorted(lost)[:5])
    return o, dg


def clean(s):
    return re.sub(r'\s+', ' ', (s or '').replace('\xa0', ' ')).strip()


def title(s, keep=()):
    """'APACHE RTR 160 2V' -> 'Apache RTR 160 2V' (letters-only words of 3+ letters capitalised, the rest as they are)"""
    out = []
    for w in clean(s).split(' '):
        out.append(w if w in keep or not w.isalpha() or len(w) <= 2 or (w.isupper() and len(w) <= 3 and w not in ('NEW', 'PRO')) else w.capitalize())
    return ' '.join(out)


# ---------------------------------------------------------------- Hero
def hero():
    src = list(rows_zip('hero'))
    lab, code = collections.defaultdict(set), {}
    for r in src:
        lab[HR['hero_label'](r['model'])].add(r['model'])
        code[r['model']] = r['model_code']
    label = {raw: l if len(raws) == 1 else '%s [%s]' % (l, [t for t in code[raw].split('_') if re.fullmatch(r'[A-Z][A-Z0-9]{3}', t)][-1])
             for l, raws in lab.items() for raw in raws}
    assert len(set(label.values())) == len(label)
    rows, cur, mrp, ph, cc = [], set(), [], {}, {}
    for r in src:
        c = int(re.sub(r'\D', '', r['cc']) or 0)
        m = HR['hero_base'](r['model'], r['model_code'], c)
        if r['status'] == 'REGULAR':
            cur.add(m)
        if c:
            cc.setdefault(m, HR['HERO_CC'].get(m, c))
        pn = clean(r['part_no'])
        mrp.append((pn, r['mrp']))
        rem = [x for x in (clean(r['remarks']),) if x]
        if clean(r['alternate']):
            rem.append('Alternate no. ' + clean(r['alternate']))
        if clean(r['superseded_by']):
            rem.append('Superseded by ' + clean(r['superseded_by']))
        y = int(r['model_year'])
        if r['photo_file'] and pn:
            ph.setdefault(pn, r['photo_file'])
        rows.append(dict(m=m, v=label[r['model']], y=y, nb=HR['NB'][HR['hero_norm'](r['model'], y, r['emission_norm'])], g=clean(r['assembly_no']),
                         asm=clean(r['assembly']), ref=clean(r['ref_no']), pn=pn, desc=clean(r['part_name']), qty=clean(r['qty']), rem='; '.join(rem),
                         ns=r['serviceable'].strip() == 'N', dg=r['diagram_file'] or None))
    unknown = {x['m'] for x in rows} - set(HR['HERO_ORDER'])
    assert not unknown, unknown
    mrp, two = merge_mrp(mrp)
    print('Hero MRP: %d part no., %d with two prices left out' % (len(mrp), len(two)))
    zp = Zips('catalogues/hero', ['HERO_Upload_Package_IMAGES_part1.zip', 'HERO_Upload_Package_IMAGES_part2.zip'], key=lambda n: n.split('/', 1)[1])
    bk = [[0, n[1], n[2]] for n in HR['NORMS']]
    return pack('Hero', 'hero', "Hero MotoCorp e-Parts Catalogue (India), Sam's Hero package, Oct 2026", HR['HERO_ORDER'], cur, rows, 'Fig', bk, cc, zp,
                mrp=mrp, photos=ph, pick=PKG + 'Hero')


# ---------------------------------------------------------------- Royal Enfield
def royal_enfield_pkg():
    """the RE package (EPC data, 3 Oct 2026): replaced by royal_enfield() below in update 9h"""
    rows, mrp, ph = [], [], {}
    for r in rows_zip('re'):
        nb = HR['NB'][r['emission_norm'] or 'BS4']   # only the Bullet Trials books (2019) name no norm: they are BS4
        dgm = re.search(r'_(IL\d+)_', r['diagram_file'])
        m = HR['re_base'](r['model'])
        y = int(r['model_year'])
        pn = clean(r['part_no'])
        mrp.append((pn, r['mrp']))
        rem = [x for x in (clean(r['remarks']),) if x]
        if clean(r['alternate_part_nos']):
            rem.append('Alternate no. ' + clean(r['alternate_part_nos']))
        if clean(r['superseded_by']):
            rem.append('Superseded by ' + clean(r['superseded_by']))
        if r['photo_file'] and pn:
            ph.setdefault(pn, r['photo_file'])
        rows.append(dict(m=m, v=HR['re_label'](r['model'], y, nb), y=y, nb=nb, g='IL%d' % int(dgm.group(1)[2:]) if dgm else '', asm=clean(r['assembly']),
                         ref=clean(r['ref_no']), pn=pn, desc=clean(r['part_name']), qty=clean(r['qty']), rem='; '.join(dict.fromkeys(rem)),
                         ns=r['serviceable'].strip() == 'false', dg=r['diagram_file'] or None))
    unknown = {x['m'] for x in rows} - set(HR['RE_ORDER'])
    assert not unknown, unknown
    mrp, two = merge_mrp(mrp)
    print('Royal Enfield MRP: %d part no., %d with two prices left out' % (len(mrp), len(two)))
    cc = {m: int(re.search(r'\d{3}', m).group(0)) for m in HR['RE_ORDER']}
    zp = Zips('catalogues/royal-enfield', ['RE_Upload_Package_IMAGES_part1.zip', 'RE_Upload_Package_IMAGES_part2.zip'], key=lambda n: n.split('/', 1)[1])
    bk = [[0, n[1], n[2]] for n in HR['NORMS']]
    return pack('Royal Enfield', 'royal-enfield', "Royal Enfield official parts catalogue (EPC, India) with the distributor MRP list (Oct 2026), Sam's RE package",
                HR['RE_ORDER'], set(HR['RE_ORDER'][:HR['RE_CUR']]), rows, 'Diagram', bk, cc, zp, mrp=mrp, photos=ph, pick=PKG + 'Royal_Enfield')


# ---------------------------------------------------------------- Yamaha (the package's own compact list, yamaha-epc.json)
YM_REN = {"Fascino-S": "Fascino S", "Fascino 125 (LCX125)": "Fascino 125 Drum", "FAZER25": "Fazer 25", "FZ-16": "FZ16", "FZ-Sfi": "FZ-S FI",
          "FZ-Sfi Hybrid": "FZ-S FI Hybrid", "FZ- Blue Flex": "FZ Blue Flex", "R15-S": "R15S", "R15 Ver 4.0": "R15 V4", "MT-15 Ver 2.0": "MT-15 V2",
          "R3 Abs": "R3 ABS", "Rayzr Sr Disc Tft": "Ray ZR SR Disc TFT", "Szrr": "SZ-RR", "Szs": "SZ-S", "Yamaha XSR": "XSR", "Yc-Z": "YC-Z",
          "Yzf-R1": "R1", "Yzf-R1X": "R1X", "Yzf-R2": "R2", "Yzf-R2 M": "R2M", "YBR-125": "YBR125DX", "YC150": "FZ-X"}


def yamaha():
    zf = zipfile.ZipFile(os.path.join(IDX, 'YAMAHA_Upload_Package_DATA.zip'))
    e = json.loads(zf.read('YAMAHA_Upload_Package/yamaha-epc.json'))
    fn = [int(m.group(1)) if (m := re.search(r'_(\d+)\.\w+$', p)) else -1 for p in e['imgs']]
    rows = []
    vm = {v: YM_REN.get(m, m) for v, m in e['vmod'].items()}
    for r in e['d']:
        img = e['imgs'][r[10] - 1] if r[10] and fn[r[10] - 1] == int(r[2]) else None   # a picture of another fig (3 books) is no picture
        for v in r[11]:
            lab = e['vars'][v]
            y = max([int(t) for t in re.findall(r'\b(?:19|20)\d\d\b', lab)] or [0])
            rows.append(dict(m=vm[lab], v=lab, y=y, nb=r[0], g=str(r[2]), asm=e['asm'][r[3]], ref=str(r[4]), pn=r[5], desc=e['desc'][r[6]], qty=str(r[7]),
                             rem=r[8], ns=bool(r[9]), dg=img))
    order = list(dict.fromkeys(YM_REN.get(m, m) for m in e['models']))
    cc = {YM_REN.get(m, m): c for m, c in e['cc'].items()}
    names = ['YAMAHA_Upload_Package_DIAGRAMS_part%d.zip' % i for i in range(1, 7)]
    zp = Zips('catalogues/yamaha', names, key=lambda n: n.split('/', 1)[1])
    return pack('Yamaha', 'yamaha', "Yamaha parts catalogues (YMSLI) with MRP from Yamaha's dealer price list, Sam's Yamaha package, Oct 2026",
                order, set(order), rows, 'Fig', e['bk'], cc, zp, mrp=e.get('mrp'), pick=PKG + 'Yamaha')


# ---------------------------------------------------------------- Honda
HONDA_NAME = {'CB Shine/ Shine 125': 'CB Shine / Shine 125', 'Activa125': 'Activa 125', 'HNESS': "H'ness CB350", 'GRAZIA': 'Grazia', 'XBLADE': 'X-Blade',
              'Shine100 DX': 'Shine 100 DX', 'CLIQ': 'Cliq', 'Activa e': 'Activa e:', '500X': 'CB500X', 'Activa i': 'Activa i', 'Navi': 'Navi'}
HONDA_ORDER = ['Activa', 'Activa 125', 'Shine 100', 'Shine 100 DX', 'CB Shine / Shine 125', 'SP125', 'SP160', 'Dio', 'Dio 125', 'Unicorn',
               'CB Unicorn 160', 'Livo', 'CD 110 Dream', 'Hornet 2.0', 'CB125 Hornet', 'X-Blade', 'CB200X', "H'ness CB350", 'CB350', 'CB300F', 'CB300R',
               'Activa e:', 'QC1', 'CB500X', 'CB650R', 'CBR650R', 'Africa Twin', 'Grazia', 'Navi', 'Cliq', 'Activa i', 'Aviator', 'Dream Yuga',
               'Dream Neo', 'CB Shine SP', 'CB Hornet 160R', 'CB Twister', 'CB Trigger', 'CB Unicorn', 'CB Unicorn Dazzler', 'CBF Stunner',
               'CBR150R', 'CBR250R', 'CBR650F', 'Eterno']
HONDA_OLD = {'Activa i', 'Aviator', 'Dream Yuga', 'Dream Neo', 'CB Shine SP', 'CB Hornet 160R', 'CB Twister', 'CB Trigger', 'CB Unicorn', 'CB Unicorn Dazzler',
             'CBF Stunner', 'CBR150R', 'CBR250R', 'CBR650F', 'Eterno', 'Navi', 'Cliq', 'Grazia', 'X-Blade'}
HONDA_CC = {'Activa': 110, 'Activa 125': 125, 'Shine 100': 100, 'Shine 100 DX': 100, 'CB Shine / Shine 125': 125, 'SP125': 125, 'SP160': 160, 'Dio': 110,
            'Dio 125': 125, 'Unicorn': 160, 'CB Unicorn 160': 160, 'Livo': 110, 'CD 110 Dream': 110, 'Hornet 2.0': 184, 'CB125 Hornet': 125, 'X-Blade': 160,
            'CB200X': 184, "H'ness CB350": 350, 'CB350': 350, 'CB300F': 293, 'CB300R': 286, 'CB500X': 471, 'CB650R': 649, 'CBR650R': 649,
            'Africa Twin': 1084, 'Grazia': 125, 'Navi': 110, 'Cliq': 110, 'Activa i': 110, 'Aviator': 110, 'Dream Yuga': 110, 'Dream Neo': 110,
            'CB Shine SP': 125, 'CB Hornet 160R': 163, 'CB Twister': 110, 'CB Trigger': 150, 'CB Unicorn': 150, 'CB Unicorn Dazzler': 150,
            'CBF Stunner': 125, 'CBR150R': 150, 'CBR250R': 249, 'CBR650F': 649, 'Eterno': 150}
# Honda's model-year letter, the last letter of the model code (SCV110S = 2025, CBF125P = 2023; letters I O Q U are not used).
# A code that ends in a digit or an unused letter gets no year (CBF125L3 is a 2020 Shine, not 2003).
HONDA_YEAR = {c: 2010 + i for i, c in enumerate('ABCDEFGHJKLMNPRST')}


def honda_year(code):
    ys = [HONDA_YEAR[p[-1]] for p in code.split('/') if p and p[-1] in HONDA_YEAR]
    return sorted(set(ys))


def honda():
    rows, mrp, raw = [], [], set()
    for r in rows_csv('Honda/honda_parts.csv'):
        raw.add(r['Model'])
        m = HONDA_NAME.get(r['Model'], clean(r['Model']))
        if m == 'CB Unicorn' and False:
            pass
        code = clean(r['Variant'])
        ys = honda_year(code)
        yl = ('%d' % ys[0] if len(ys) == 1 else '%d-%d' % (ys[0], ys[-1])) if ys else ''
        v = '%s · %s%s' % (m, yl + ' · ' if yl else '', code)
        pn = clean(r['Part No'])
        mrp.append((pn, r['MRP']))
        g = clean(r['Fig No'])
        rows.append(dict(m=m, v=v, y=ys[-1] if ys else 0, nb=0, g=g, asm=clean(r['Figure']), ref=clean(r['Ref']), pn=pn, desc=clean(r['Description']),
                         qty=clean(r['Qty']), rem='', ns=False, dg=r['diagram_file'] or None))
    unknown = {x['m'] for x in rows} - set(HONDA_ORDER)
    assert not unknown, unknown
    mrp, two = merge_mrp(mrp)
    print('Honda MRP: %d part no., %d with two prices left out' % (len(mrp), len(two)))
    zp = Zips('catalogues/honda', ['honda_diagrams_HD.zip'])
    return pack('Honda', 'honda', "Honda 2W parts catalogue (HMSI, Bike Pharma retailer login), Sam's Honda package, Oct 2026", HONDA_ORDER,
                set(HONDA_ORDER) - HONDA_OLD, rows, 'Fig', [[0, [], '']], HONDA_CC, zp, mrp=mrp, pick=PKG + 'Honda')


# ---------------------------------------------------------------- TVS
TVS_NAME = {'APACHE RR 310': 'Apache RR 310', 'APACHE RTR 160 2V': 'Apache RTR 160 2V', 'APACHE RTR 180 2V': 'Apache RTR 180 2V',
            'APACHE RTR 200 4V': 'Apache RTR 200 4V', 'APACHE RTR 310': 'Apache RTR 310', 'APACHE RTX': 'Apache RTX', 'Apache RTR 160 4V': 'Apache RTR 160 4V',
            'IQUBE': 'iQube', 'JUPITER 110': 'Jupiter 110', 'JUPITER 110 - 2024': 'Jupiter 110 (2024)', 'JUPITER 125': 'Jupiter 125', 'NTORQ 125': 'Ntorq 125',
            'NTORQ 150': 'Ntorq 150', 'ORBITER': 'Orbiter', 'PHOENIX 125': 'Phoenix 125', 'RADEON': 'Radeon', 'RAIDER 125': 'Raider 125',
            'RONIN 225': 'Ronin 225', 'SCOOTY PEP PLUS': 'Scooty Pep Plus', 'SCOOTY ZEST': 'Scooty Zest', 'SPORT': 'Sport', 'STAR CITY': 'Star City',
            'STAR CITY PLUS': 'Star City Plus', 'VICTOR NEW': 'Victor (new)', 'WEGO 110': 'Wego 110', 'XL100': 'XL100', 'XL SUPER': 'XL Super'}
TVS_ORDER = ['Jupiter 110 (2024)', 'Jupiter 110', 'Jupiter 125', 'XL100', 'Apache RTR 160 4V', 'Apache RTR 160 2V', 'Raider 125', 'Ntorq 125', 'Ntorq 150',
             'Radeon', 'Star City Plus', 'Sport', 'Apache RTR 200 4V', 'Apache RTR 180 2V', 'Apache RTR 310', 'Apache RR 310', 'Apache RTX', 'Ronin 225',
             'iQube', 'Orbiter', 'Scooty Pep Plus', 'Scooty Zest', 'Victor (new)', 'Wego 110', 'Phoenix 125', 'Star City', 'XL Super',
             'Flame 125', 'Jive', 'Max 4R', 'Victor', 'Scooty Pep', 'Scooty Streak', 'Scooty', 'Scooty Teenz', 'Oils & lubricants']
TVS_OLD = {'Victor (new)', 'Wego 110', 'Phoenix 125', 'Star City', 'XL Super', 'Apache RTR 180 2V', 'Scooty Zest',
           'Flame 125', 'Jive', 'Max 4R', 'Victor', 'Scooty Pep', 'Scooty Streak', 'Scooty', 'Scooty Teenz', 'Oils & lubricants'}
TVS_CC = {'Jupiter 110 (2024)': 113, 'Jupiter 110': 110, 'Jupiter 125': 124, 'XL100': 99, 'Apache RTR 160 4V': 160, 'Apache RTR 160 2V': 160,
          'Raider 125': 125, 'Ntorq 125': 125, 'Ntorq 150': 150, 'Radeon': 110, 'Star City Plus': 110, 'Sport': 110, 'Apache RTR 200 4V': 198,
          'Apache RTR 180 2V': 180, 'Apache RTR 310': 312, 'Apache RR 310': 312, 'Apache RTX': 299, 'Ronin 225': 225, 'Scooty Pep Plus': 90,
          'Scooty Zest': 110, 'Victor (new)': 110, 'Wego 110': 110, 'Phoenix 125': 125, 'Star City': 110, 'XL Super': 70,
          'Flame 125': 125, 'Jive': 110, 'Max 4R': 110, 'Victor': 110}
TVS_NB = {'': 0, 'BS-VI': 1, 'BS-IV': 2, 'BS-III': 3}
# painted parts (tvs_painted_parts.csv): its Model Family -> our model; APACHE / JUPITER / STAR-CITY versions go to the model of
# theirs whose part numbers the parts book shares most (a tie: every tied model gets the version)
TVS_PAINT_FAM = {'VICTOR': 'Victor', 'VICTOR 2016': 'Victor (new)', 'SPORT': 'Sport', 'RADEON': 'Radeon', 'FLAME': 'Flame 125', 'JIVE': 'Jive',
                 'MAX 4R': 'Max 4R', 'PHOENIX': 'Phoenix 125', '0010012': 'Raider 125', 'RAIDER125.': 'Raider 125', 'RONIN 225': 'Ronin 225',
                 'XL100': 'XL100', 'XL': 'XL Super', 'XL SUPER HD': 'XL Super', 'PEP': 'Scooty Pep', 'PEP +': 'Scooty Pep Plus',
                 'STREAK': 'Scooty Streak', 'ZEST': 'Scooty Zest', 'IQUBE': 'iQube', 'ORBITER': 'Orbiter', '2': 'Wego 110', 'WEGO': 'Wego 110',
                 'JUPITER125': 'Jupiter 125', 'NTORQ150': 'Ntorq 150', 'NTorq': 'Ntorq 125'}
TVS_PAINT_PICK = {'APACHE': ('Apache RTR 160 2V', 'Apache RTR 180 2V', 'Apache RTR 160 4V', 'Apache RTR 200 4V', 'Apache RR 310', 'Apache RTR 310', 'Apache RTX'),
                  'JUPITER': ('Jupiter 110', 'Jupiter 110 (2024)'), 'STAR-CITY': ('Star City Plus', 'Star City')}
TVS_PAINT_NONE = {'STAR-CITY': ['Star City'],          # versions sharing no part number with the parts books (older ones)
                  '000010000100000764': ['Apache RTR 160 4V']}   # inferred: its neighbours 762 and 766 are the 160 4V
# kits (tvs_kit_parts.csv): its Model -> our model
TVS_KIT = {'Star City': 'Star City', 'Apache RTR 160': 'Apache RTR 160 2V', 'Star City Plus': 'Star City Plus', 'Flame': 'Flame 125',
           'XL Super & HD': 'XL Super', 'XL 100': 'XL100', 'Pep+': 'Scooty Pep Plus', 'Apache RTR 180 ABS': 'Apache RTR 180 2V',
           'New Victor': 'Victor (new)', 'Radeon': 'Radeon', 'Pheonix 125': 'Phoenix 125', 'Apache RTR 200 4V': 'Apache RTR 200 4V',
           'Apache RR 310': 'Apache RR 310', 'Apache RTR 160 4V': 'Apache RTR 160 4V', 'Streak': 'Scooty Streak', 'NTorq 125': 'Ntorq 125',
           'Wego': 'Wego 110', 'Jupiter': 'Jupiter 110', 'Zest': 'Scooty Zest', 'Sport': 'Sport'}


def tvs_extra(src, rows, mrp, photos):
    """TVS painted parts (book 4: a section per colour, its own parts picture), kits (book 5: the kit no. with its MRP, then what is in it)
    and oils & lubricants (book 6, with their photos) -> rows, mrp pairs, photos"""
    pm = collections.defaultdict(set)
    for r in src:
        pm[clean(r['Part No'])].add(TVS_NAME[r['Model']])
    pp = list(rows_csv('TVS/tvs_painted_parts.csv'))
    vers = collections.defaultdict(list)
    for r in pp:
        vers[r['Model ID']].append(r)
    to = {}
    for mid, rs in vers.items():
        fam = clean(rs[0]['Model Family'])
        if fam in TVS_PAINT_FAM:
            to[mid] = [TVS_PAINT_FAM[fam]]
            continue
        c = collections.Counter()
        for r in rs:
            pn = clean(r['Part No'])
            c.update(m for m in (pm.get(pn) or pm.get(pn[:-1]) or pm.get(pn[:-2]) or ()) if m in TVS_PAINT_PICK[fam])
        top = max(c.values(), default=0)
        to[mid] = sorted(m for m, n in c.items() if n == top) if top else (TVS_PAINT_NONE.get(mid) or TVS_PAINT_NONE[fam])
    per = collections.defaultdict(list)
    for mid in sorted(vers):
        for m in to[mid]:
            per[m].append(mid)
    n_paint = 0
    for m, mids in per.items():
        g = 0
        for k, mid in enumerate(mids, 1):
            v = m + ' · painted parts' + (' %d' % k if len(mids) > 1 else '')
            cols = collections.OrderedDict()
            for r in vers[mid]:
                cols.setdefault(clean(r['Colour ID']), []).append(r)
            for cid, rs in cols.items():
                g += 1
                col = clean(rs[0]['Colour']).upper()
                for r in rs:
                    pn = clean(r['Part No'])
                    mrp.append((pn, r['MRP (incl tax)']))
                    rem = [x for x in (clean(r['Remarks']), clean(r['Emission Norm'])) if x]
                    rows.append(dict(m=m, v=v, y=0, nb=4, g='P%d' % g, asm=col + ' · PAINTED PARTS' + (' %d' % k if len(mids) > 1 else ''),
                                     ref=clean(r['Ref No']), pn=pn, desc=clean(r['Part Name']), qty=clean(r['Qty']), rem='; '.join(rem), ns=False,
                                     dg=clean(r['Parts Diagram']) or None))
                    n_paint += 1
    kits = collections.OrderedDict()
    for r in rows_csv('TVS/tvs_kit_parts.csv'):
        kits.setdefault((TVS_KIT[clean(r['Model'])], clean(r['Kit Part No'])), []).append(r)
    kn = collections.Counter()
    for (m, kit), rs in kits.items():
        kn[m] += 1
        n = kn[m]
        mrp.append((kit, rs[0]['Kit MRP (incl tax)']))
        common = dict(m=m, v=m + ' · kits', y=0, nb=5, g='KIT', asm='KITS AND WHAT IS IN THEM', ns=False, dg=None)
        rows.append(dict(common, ref='K%d' % n, pn=kit, desc=clean(rs[0]['Kit Description']), qty='1', rem='Kit for ' + clean(rs[0]['Model Applicable'])))
        for j, r in enumerate(rs, 1):
            rows.append(dict(common, ref='K%d.%d' % (n, j), pn=clean(r['Part No']), desc=clean(r['Part Description']), qty=clean(r['Qty']),
                             rem='; '.join(x for x in ('In kit ' + kit, clean(r['Emission Norm'])) if x.strip())))
    oils = list(rows_csv('TVS/tvs_oils_lubricants.csv'))
    for i, r in enumerate(oils, 1):
        pn = clean(r['Part No'])
        mrp.append((pn, r['MRP (incl tax)']))
        ml = clean(r['Capacity (ml)'])
        rows.append(dict(m='Oils & lubricants', v='Oils & lubricants', y=0, nb=6, g='OIL', asm='TVS OILS & LUBRICANTS', ref=str(i), pn=pn,
                         desc=clean(r['Description']), qty='1', rem=clean(r['Range']) + (' · %s ml' % ml if ml not in ('', '0') else ''), ns=False, dg=None))
        if clean(r['Image']):
            photos[pn] = clean(r['Image'])
    print('TVS extras: painted parts %d lines (%d versions, %d colours) | kits %d (%d lines) | oils %d (%d photos)'
          % (n_paint, len(vers), len({r['Colour ID'] for r in pp}), len(kits), sum(len(x) for x in kits.values()), len(oils), len(photos)))


def tvs():
    src = list(rows_csv('TVS/tvs_parts.csv'))
    allv = collections.defaultdict(set)
    split = lambda q: [t.strip() for t in q.split(',') if t.strip()]
    for r in src:
        allv[r['Model']].update(split(r['Qualified Variant']))
    rows, mrp = [], []
    for r in src:
        m = TVS_NAME[r['Model']]
        vs = split(r['Qualified Variant']) or sorted(allv[r['Model']]) or ['']
        pn = clean(r['Part No'])
        mrp.append((pn, r['MRP (incl tax)']))
        rem = [x for x in (clean(r['Remarks']),) if x]
        if clean(r['Alternate Part']):
            rem.append('Alternate no. ' + clean(r['Alternate Part']))
        for v in vs:
            rows.append(dict(m=m, v=m + (' · ' + title(v, keep=('CVTI', 'EFI', 'ABS', 'SBS', 'DLX', 'ES', 'KS', 'S/ST')) if v else ''), y=0,
                             nb=TVS_NB.get(clean(r['Emission Norm']), 0), g=clean(r['Fig No']), asm=clean(r['Assembly Name']), ref=clean(r['Ref No']),
                             pn=pn, desc=clean(r['Part Name']), qty=clean(r['Qty']), rem='; '.join(rem), ns=False,
                             dg=r['Figure Image'] or None))   # the zip entry itself (images/<category>/<file>.png)
    photos = {}
    tvs_extra(src, rows, mrp, photos)
    mrp, two = merge_mrp(mrp)
    print('TVS MRP: %d part no., %d with two prices left out' % (len(mrp), len(two)))
    zp = Zips('catalogues/tvs', ['tvs_images.zip'])
    bk = [[0, [], ''], [0, ['BS6'], ''], [0, ['BS4'], ''], [0, ['BS3'], ''], [0, [], 'Painted parts in each colour, from TVS’s painted parts list'],
          [0, [], 'TVS kits: the kit part no. with its MRP, then the parts in it'], [0, [], 'TVS oils and lubricants']]
    o, dg = pack('TVS', 'tvs', "TVS Motor dealer e-catalogue (TVS-approved), MRP incl. tax, Sam's TVS package, Oct 2026", TVS_ORDER,
                 set(TVS_ORDER) - TVS_OLD, rows, 'Fig', bk, TVS_CC, zp, mrp=mrp, photos=photos, pick=PKG + 'TVS',
                 by_group=lambda kv: (kv[0][0] >= 4, kv[0][0] if kv[0][0] >= 4 else 0, nat(kv[0][1]), kv[0][0], nat(kv[0][3])))   # figure by figure, its BS6 / BS4 / BS3 lines next to it; painted parts, kits, oils after
    return tvs_books(o, dg)          # then TVS's own PDF books from Sam's folder All Catalogues > TVS ALL


# ---------------------------------------------------------------- TVS books (Sam's folder All Catalogues > TVS ALL, 11 Oct 2026)
# TVS's own parts catalogues and colour part lists (PDF, made from TVS's Excel masters). tvs_parse.py reads their tables (only the
# text a page shows) -> tvs-pages.json. A book whose part numbers are (nearly) all in the dealer e-catalogue already is not added: its
# file is listed as spare. Each other book is added whole after the e-catalogue (which stays exactly as it was), as an edition of its
# model named by what the book prints (title, version, date): '<model> · <book>'. A figure's diagram is its page of the book, drawn from
# the original PDF that Admin uploads unchanged: 'pdf:catalogues/tvs/<book>:0:<size>:<page>:0,0,1,1'. The quantity columns of a book
# (DRUM | DISC, KICK START | ELEC. START ...) are read as the book prints them: a part not in every column of its model says which
# ("For DISC"), different quantities say each. The colour lists add the painted parts the painted parts list does not have: each colour
# its own section, one edition per version of the bike (as the painted parts list is laid out).
TVS_DIR = 'catalogues/tvs'
TVS_SRC = "TVS parts catalogues and colour part lists (PDF), Sam's folder All Catalogues > TVS ALL, Oct 2026"
# parts book -> (its quantity columns as printed, [(model, the columns that are that model, the edition: what the book prints)])
TVS_BOOKS = {
    'SCOOTY 2S & TEENZ 2S.pdf': (['KICK START', 'ELEC. START', 'NHL KS', 'NHL ES', 'UPG ES', 'TEENZ ES'],
                                 [('Scooty', [0, 1, 2, 3, 4], 'Scooty/NHL/UPG/Teenz parts catalogue (15.05.2010)'),
                                  ('Scooty Teenz', [5], 'Scooty/NHL/UPG/Teenz parts catalogue (15.05.2010)')]),
    'FLAME CATALOGUE PART-2.pdf': (['DRUM', 'DISC'], [('Flame 125', [0, 1], 'parts catalogue part 2: figures 14-35 (27.03.2008)')]),
    'STREAK.pdf': (['STREAK'], [('Scooty Streak', [0], 'parts catalogue rev. 01 (27.09.2013)')]),
    'JIVE PARTS C.pdf': (['JIVE'], [('Jive', [0], 'parts catalogue rev. 01 (03.01.14)')]),
    'APACHE RR 310.pdf': (['RR 310'], [('Apache RR 310', [0], 'BSVI parts catalogue (02 10 20)')]),
    'PEP & PEP+ PARTS CATALOG 15.02.2013.pdf': (['SCOOTY PEP', 'SCOOTY PEP+'], [('Scooty Pep', [0], 'Scooty Pep & Pep+ parts catalogue'),
                                                                                 ('Scooty Pep Plus', [1], 'Scooty Pep & Pep+ parts catalogue')]),
    'STAR SITY PC.pdf': (['STAR BASE', 'STAR DLX', 'STAR CITY', 'STAR CITY ALLOY', 'STAR CITY REF.', 'STAR SPORT DLX', 'STAR SPORT CVTI',
                          'STAR CITY 110 CVTI'],
                         [('Star City', [0, 1, 2, 3, 4], 'Star / Star City BSIV parts catalogue rev. 03 (19.02.2017)'),
                          ('Sport', [5, 6], 'Star Sport BSIV parts catalogue rev. 03 (19.02.2017)'),
                          ('Star City Plus', [7], 'Star City 110 CVTI BSIV parts catalogue rev. 03 (19.02.2017)')]),
    'APACHE 160 4V PC ,..pdf': (['CARB. VER.', 'FI VER.', '165RP'], [('Apache RTR 160 4V', [0, 1, 2], 'Ride Mode parts catalogue (Carb. / FI / 165RP)')]),
    'RAIDER 125.pdf': (['RAIDER'], [('Raider 125', [0], 'parts catalogue')]),
    'WEGO.pdf': (['DRUM', 'DISC'], [('Wego 110', [0, 1], 'parts catalogue')]),
    'APACHE 200 4V PC.pdf': (['DISC', 'EFI DISC'], [('Apache RTR 200 4V', [0, 1], 'BSVI & Ride Mode parts catalogue (20 01 21)')]),
    'NTROQ 125.pdf': (['NTORQ DRUM/DISC'], [('Ntorq 125', [0], 'BSVI parts catalogue')]),
    'o qube.pdf': (['DISC', '2022/ELE.S'], [('iQube', [0, 1], '2022 / ELE. S parts catalogue')]),
}
TVS_REM_KIT = {'o qube.pdf'}            # its remarks column is headed KIT PART NO.
# colour lists: the model each version's title names (the versions with no painted part new to the painted parts list are left out)
TVS_COLOUR = {
    'APACHE COLOR PART MATRIX FIN.pdf': ('colour part matrix', [('RTR 160', 'Apache RTR 160 2V'), ('RTR 180', 'Apache RTR 180 2V')]),
    'APCHE COLOUR 2V.pdf': ('colour parts', [('RTR 160 2V', 'Apache RTR 160 2V')]),
    'APACHE 1604VCOLOUR.pdf': ('colour parts', [('RTR 160 4V', 'Apache RTR 160 4V')]),
    'APACHE 200 4V COLOUR.pdf': ('colour parts', [('RTR 200 4V', 'Apache RTR 200 4V')]),
}
# the other files of the folder: their parts are in the e-catalogue already (or in a book above), or they have no text to read
TVS_SPARE = ['APACHE 160 PC.pdf', 'APACHE 18 ABS PC.pdf', 'APACHE 180 PC.pdf', 'APACHE 200 4V PC (2).pdf', 'APACHE RR310 PARTS CATALOGUE.pdf',
             'APACHEPDFNEW_ 26042013.pdf', 'Apache 160 2v raid moad.pdf', 'JUPITER 110 PC.pdf', 'JUPITER 110.pdf', 'JUPITER 125 PC.pdf',
             'JUPITER 2024 NEW.pdf', 'NT XP.pdf', 'NTORQ150 BASE PREMIUM PARTS CATALOGUE.pdf', 'PEP + PC.pdf', 'Phoenix- Parts Catalogue.pdf',
             'REDEON 110.pdf', 'REDON PC.pdf', 'RONIN PC.pdf', 'SCOOTY PEP+.pdf', 'SPORT PC,.pdf', 'STAR CITY 110+ BS 4 BS6.pdf', 'VICTOR NEW.pdf',
             'XL 100 PC.pdf', 'XLSuper_Catalogue.pdf', 'n.pdf', 'nt xt ver.pdf', 'raider khin ver.pdf', 'zest.pdf',
             'JUPITER 125 COLOUR.pdf', 'SPORT COLOUR.pdf', 'ZL 100 COLOUR.pdf', 'SPORT COLOUR PC.pdf']
# what tvs_parse.py cannot know: a remark printed in two lines of a tall cell is read as the end of the line above
TVS_REM_SPLIT = (re.compile(r'^(.*?)\s*\b(REF\.?\s*COLOU?R|Replaces|S\.\s?No\.?\s*-?)$', re.I),
                 re.compile(r'^(MATRIX|MATRIC|MAT\.?|PAGE(?:\s*NO\.?)?|[A-Z]{1,2}\d{6,}[A-Z0-9]*|\d[\d,]*)$'))
TVS_FIX = {('STREAK.pdf', 38, 'SIDE PANEL TOP RH'): {'rem': ''}}     # a smudge of tiny marks in its remarks cell, nothing readable
TVS_DASH = {'', '-', '--', '---', '----', '–', '—', '___', '____', '_____', '…'}
TVS_NOQ = {'', '-', '--', '---', '–', '—', 'NA', 'N.A', 'N.A.', '0'}
TVS_SYM = re.compile(r'^[*#$&@]{1,3}$')


def tvs_key(s):
    return re.sub(r'[^A-Z0-9*#$&@]', '', s.upper())


def tvs_defs(notes):
    """a page's footnotes -> {mark: what it means}: '** - REF. KIT & KIT CONTENTS N.I - NOT ILLUSTRATED', '* MARKED ITEM NOT SERVICED'"""
    out = {}
    for t in notes:
        starts = [m for m in re.finditer(r"(?:^|(?<=[\s,]))(\*{1,3}|#|\$|&|@|'BLOCKED'|\"BLOCKED\"|[A-Z](?:\.\s?[A-Z]){1,2}\.?|[A-Z]{2,4}\.?)"
                                         r"\s*'?\s*(-|:-|:)\s*", t)]
        m0 = re.match(r"^(\*{1,3}|#|\$|&|@)\s*'?\s*", t)          # a mark first on its line needs no dash: '* MARKED ITEM NOT SERVICED'
        if m0 and not (starts and starts[0].start() == 0):
            starts.insert(0, m0)
        for k, m in enumerate(starts):
            end = starts[k + 1].start() if k + 1 < len(starts) else len(t)
            val = t[m.end():end].strip(' ,;.')
            if val:
                out.setdefault(tvs_key(m.group(1)), val)
    return out


def tvs_qty(q):
    """a quantity cell -> None (not for this column) or (quantity, footnote mark)"""
    q = q.strip()
    if q in TVS_NOQ:
        return None
    m = re.fullmatch(r'(\d{1,3}(?:\.\d{1,2})?)\s*([*#$&@]{1,3})?', q)
    if m:
        return (None if m.group(1) == '0' else (m.group(1), m.group(2) or ''))
    if re.fullmatch(r'A\.?\s?R\.?', q):
        return ('AR', '')
    if TVS_SYM.match(q):
        return ('', q)
    return (q, '')


class TvsBook:
    """one book's footnotes: the page's own, else the book's (a mark that means one thing all through the book)"""
    def __init__(self, notes):
        self.page = {int(p): tvs_defs(ns) for p, ns in notes.items()}
        allv = collections.defaultdict(collections.Counter)
        for d in self.page.values():
            for k, v in d.items():
                allv[k][v] += 1
        norm = lambda v: re.sub(r'\b(PL|PLEASE)\b|[^A-Z]', '', v.upper().replace(' AND ', ' & ').replace('REFER', 'REF'))
        self.book = {k: c.most_common(1)[0][0] for k, c in allv.items() if len({norm(v) for v in c}) == 1}

    def means(self, page, mark):
        k = tvs_key(mark)
        return self.page.get(page, {}).get(k) or self.book.get(k)

    def rem(self, page, rem):
        """the remarks column as the book means it -> (text, not serviced)"""
        rem = clean(rem)
        if rem in ('BSVI', 'BSIV', 'BS VI', 'BS IV'):
            return rem.replace(' ', ''), False
        ns, items = False, []
        for it in [x.strip() for x in rem.split(',') if x.strip()]:
            toks = it.split()
            while toks and tvs_key(toks[0]) in ('NS', 'BLOCKED', "BLOCKED'"):
                ns = True
                toks = toks[1:]
            if toks and len(toks) > 1 and tvs_key(toks[0]) == 'N' and tvs_key(toks[1]) == 'S':    # 'N S'
                ns, toks = True, toks[2:]
            it = ' '.join(toks)
            if not it:
                continue
            m = re.fullmatch(r'(ALT|RE)\.?\s*-?\s*([A-Z0-9]{6,})', it)
            if m:
                items.append(('Alternate no. ' if m.group(1) == 'ALT' else 'Replaces ') + m.group(2))
                continue
            pm = re.fullmatch(r'\(\s*(\S+)\s*\)', it)
            if pm and self.means(page, pm.group(1)):
                it = pm.group(1)
            if self.means(page, it) and tvs_key(it) not in ('NS', 'BLOCKED'):
                items.append('%s (%s)' % (it, self.means(page, it)))
                continue
            parts = it.split()
            if len(parts) > 1 and all(self.means(page, p) for p in parts):
                items.append(' '.join('%s (%s)' % (p, self.means(page, p)) for p in parts))
                continue
            items.append(it)
        return ', '.join(dict.fromkeys(items)), ns


def tvs_section_of(lines):
    """the book's figures: page -> group (the FIG.NO.). A page whose FIG.NO. another title holds and whose title is the page before's
    goes on with the page before's figure (TVS printed the next figure's number on it); the same figure printed with two spellings of
    its title stays one figure"""
    pages = collections.OrderedDict()
    for l in lines:
        pages.setdefault(l[0], (l[1], l[2]))
    norm = lambda t: re.sub(r'[^A-Z0-9]', '', t.upper().replace('ASSEMBLY', 'ASSY').replace('BRACKET', 'BRKT').replace('CARBURETOR', 'CARBURETTOR'))
    titles = collections.defaultdict(set)
    for fig, t in pages.values():
        titles[fig].add(norm(t))
    out, prev, moved = {}, None, []
    for p, (fig, t) in pages.items():
        if prev and prev[0] != fig and norm(t) == norm(prev[1]) and len(titles[fig]) > 1:
            moved.append((p, fig, prev[0]))
            fig = prev[0]
        out[p] = fig
        prev = (fig, t)
    return out, moved


def tvs_books(o, dg):
    """the TVS books after the e-catalogue: rows, editions, models, diagrams and the files Admin uploads"""
    B = json.load(open(os.path.join(HERE, 'tvs-pages.json')))['books']
    assert set(TVS_BOOKS) | set(TVS_COLOUR) | set(TVS_SPARE) == set(B), set(B) ^ (set(TVS_BOOKS) | set(TVS_COLOUR) | set(TVS_SPARE))
    painted = {r[5] for r in o['d'] if r[0] == 4}
    have_pn = {r[5] for r in o['d']}
    nv0, nb0 = len(o['vars']), len(o['bk'])
    o['src'].append(TVS_SRC)
    src = len(o['src']) - 1
    vars_, vmod, rows, links, stats = [], {}, [], {}, collections.Counter()
    book_of = {}

    def edition(m, label):
        v = m + ' · ' + label
        if v not in vmod:
            vars_.append(v)
            vmod[v] = m
        return nv0 + vars_.index(v)

    def add(nb, g, page, asm, ref, pn, desc, qty, rem, ns, eds, book):
        pn = '' if pn in TVS_DASH or re.fullmatch(r'[-–—_.…\s]*', pn) else pn
        if not pn and not clean(desc).strip('-–—_.… '):
            stats['lines with no part no. and no description left out'] += 1       # a tool kit's '---' lines: nothing to show
            return
        rows.append((nb, g, asm, ref, pn, clean(desc), qty, rem, 1 if ns else 0, eds))
        for v in eds:
            links.setdefault((v, g), (book, page))

    def colour_rows(nb, b, f, blocks, eds_of, gn):
        for (page, title, mtitle), cs in blocks:
            ver, _, col = title.rpartition(' - ')
            if not ver:
                ver, col = mtitle, title
            eds = eds_of(page, title, mtitle)
            if not eds:
                continue
            gn += 1
            for c in cs:
                add(nb, 'C%d' % gn, page, clean(col) + ' · ' + clean(ver) + b.get('rd', {}).get(page, ''), c[3], c[4], c[5], '', '', False, eds, f)
                stats['colour lines'] += 1
        return gn

    bk_new = []
    for f, (labels, models) in TVS_BOOKS.items():
        b = B[f]
        nb = nb0 + len(bk_new)
        bk_new.append([src, [], ''])
        book_of[nb] = f
        eds = [edition(m, lab) for m, cols, lab in models]
        tb = TvsBook(b['notes'])
        sec, moved = tvs_section_of(b['lines'])
        for p, a, z in moved:
            print('   %s p%d: printed FIG.NO. %s, its title is figure %s going on' % (f, p, a, z))
        # the page's own label of a single quantity column (RAIDER DRUM / RAIDER DISC ...) where it is not the book's usual one
        plab = {int(p): (c[0] if c else '') for p, c in b['cols'].items()}
        usual = collections.Counter(plab.values()).most_common(1)[0][0] if plab else ''
        L = [list(l) for l in b['lines']]
        for l in L:
            assert len(l[6]) == len(labels), (f, l)
        # a part number alone on a line (no ref, no description) under a line that has none: that line's number
        keep = []
        for l in L:
            if not l[3] and not l[5] and l[4] and keep and keep[-1][0] == l[0] and keep[-1][4] in TVS_DASH:
                keep[-1][4] = l[4]
                stats['part no. joined to the line above'] += 1
                continue
            keep.append(l)
        L = keep
        # a remark in two lines of a tall cell, its first line read with the line above
        for i in range(len(L) - 1):
            a, z = L[i], L[i + 1]
            m1, m2 = TVS_REM_SPLIT[0].match(a[7]), TVS_REM_SPLIT[1].match(z[7])
            if a[0] == z[0] and m1 and m2:
                a[7], z[7] = m1.group(1), m1.group(2) + ' ' + z[7]
                stats['remarks in two lines put together'] += 1
        for l in L:
            fx = TVS_FIX.get((f, l[0], l[5]))
            if fx:
                l[7] = fx['rem']
                stats['fixed by hand'] += 1
        # a ref the page draws (not as text): the one between its neighbours
        for i, l in enumerate(L):
            if l[3]:
                continue
            fig = sec[l[0]]
            prv = next((x[3] for x in reversed(L[:i]) if sec[x[0]] == fig), None)
            nxt = next((x[3] for x in L[i + 1:] if sec[x[0]] == fig), None)
            if prv and nxt and prv.isdigit() and nxt.isdigit() and int(nxt) == int(prv) + 2:
                l[3] = str(int(prv) + 1)
            elif prv is None and nxt == '2':
                l[3] = '1'
            else:
                continue
            stats['ref read from its neighbours'] += 1
        first = {}
        for l in L:
            first.setdefault(sec[l[0]], l)
        last = {}                                       # the line above in the figure (for ditto marks)
        for l in L:
            page, fig, ref, pn, desc, qty, rem = l[0], sec[l[0]], l[3], l[4], l[5], list(l[6]), l[7]
            asm = clean(first[fig][2])
            prev = last.get(fig)
            for i, q in enumerate(qty):
                if q.strip() in ('"', '”', '“') and prev:
                    qty[i] = prev[0][i]
            dm = re.match(r'^["”“](?![A-Za-z])\s*(.*)$', rem)
            base = rem
            if dm:                                      # a ditto mark: the remark of the line above (and what follows the mark)
                base = prev[1] if prev else ''
                rem = (base + (' ' + dm.group(1) if dm.group(1) else '')).strip(' ,')
            last[fig] = (qty, base)                     # (the next ditto means the same, not this line's own marks too)
            cells = [tvs_qty(q) for q in qty]
            text, ns = tb.rem(page, rem)
            if f in TVS_REM_KIT and re.fullmatch(r'[A-Z0-9]{7,12}(\s+[A-Z0-9]{7,12})*', text) and any(c.isdigit() for c in text):
                text = 'Kit part no. ' + text
            pre = []
            if len(labels) == 1 and plab.get(page) and plab[page] != usual:
                pre.append('For ' + plab[page])
            anyq = any(c is not None for c in cells)
            for (m, cols, lab), ed in zip(models, eds):
                on = [i for i in cols if cells[i] is not None]
                if not on and anyq:
                    continue
                parts = list(pre)
                if on and len(on) < len(cols):
                    parts.append('For ' + ', '.join(labels[i] for i in on))
                vals = [cells[i][0] for i in on]
                q = vals[0] if vals and len(set(vals)) == 1 else ''
                if len(set(vals)) > 1:
                    parts.append('Qty ' + ', '.join('%s %s' % (labels[i], cells[i][0] or '-') for i in on))
                elif q == 'AR':
                    q = ''
                    parts.append('Qty as required (AR)')
                marks = collections.OrderedDict()
                for i in on:
                    if cells[i][1]:
                        marks.setdefault(cells[i][1], []).append(labels[i])
                for mk, where in marks.items():
                    what = tb.means(page, mk)
                    parts.append(('%s: ' % ', '.join(where) if len(where) < len(on) and len(on) > 1 else '') + mk + (' (%s)' % what if what else ''))
                if text and text not in parts:
                    parts.append(text)
                note = '; '.join(parts)
                if note in ('BSVI', 'BSIV'):            # the app hides a bare BSVI / BSIV (it shows the e-catalogue's norm instead)
                    note = 'For ' + note
                add(nb, fig, first[fig][0], asm, ref, pn, desc, q, note, ns, [ed], f)
                stats['book lines'] += 1
        # its colour pages, kit list and grading list
        blocks = collections.OrderedDict()
        for c in b['colour']:
            blocks.setdefault((c[0], c[1], c[2]), []).append(c)
        colour_rows(nb, b, f, list(blocks.items()), lambda *a: [eds[0]] if len(eds) == 1 else [], 0)
        if b['colour']:
            assert len(eds) == 1, f
        kn = 0
        kits = collections.OrderedDict()
        for k in b['kits']:
            kits.setdefault((k[0], k[1], k[2], k[3], k[4]), []).append(k)
        kpages = sorted({k[0] for k in b['kits']})
        for (page, sno, kpn, kdesc, model), ks in kits.items():
            kn += 1
            g = 'KIT%d' % (kpages.index(page) + 1)
            snos = [s for p, s, *_ in kits if p == page]
            asm = 'KIT & KIT CONTENTS · KITS %s-%s' % (snos[0], snos[-1])
            kpn = '' if kpn in TVS_DASH else kpn
            add(nb, g, page, asm, 'K' + sno, kpn, kdesc, '1', 'Kit for ' + model if model else 'Kit', False, eds, f)
            for k in ks:
                add(nb, g, page, asm, 'K%s.%s' % (sno, k[5]), k[6], k[7], k[8], ('In kit ' + kpn) if kpn else ('In ' + kdesc), False, eds, f)
                stats['kit lines'] += 1
        for page, sno, bpn, desc, qty, ser, grade, gpn in b['grades']:
            bpn = '' if bpn in TVS_DASH else bpn
            add(nb, 'GRADE', page, 'GRADING PART NUMBERS', '%s.%s' % (sno, ser), gpn, '%s GRADE %s' % (desc, grade.upper()), qty,
                'Grade-wise part no.' + (' of ' + bpn if bpn else ''), False, eds, f)
            stats['grading lines'] += 1
    for f, (what, fams) in TVS_COLOUR.items():
        b = B[f]
        nb = nb0 + len(bk_new)
        bk_new.append([src, [], 'Painted parts in each colour, from TVS’s colour part list'])
        book_of[nb] = f
        rd, cur = {}, ''
        for p in range(1, b['pages'] + 1):          # a version's release date, printed on its first page
            cur = b['notes'].get(str(p), [cur])[0] if str(p) in b['notes'] else cur
            rd[p] = cur
        blocks = collections.OrderedDict()
        for c in b['colour']:
            blocks.setdefault((c[0], c[1], c[2]), []).append(c)
        new = [(k, cs) for k, cs in blocks.items() if any(c[4] not in painted for c in cs)]
        first_rd = {}
        for (page, title, mtitle), cs in new:
            ver = title.rpartition(' - ')[0] or mtitle
            first_rd.setdefault(ver, rd.get(page, ''))
        b2 = dict(b, rd={p: (' (%s)' % d if d else '') for p, d in rd.items()})

        def eds_of(page, title, mtitle, fams=fams, what=what):
            ver = title.rpartition(' - ')[0] or mtitle
            m = next((m for k, m in fams if k in ver), None)
            assert m, (f, ver)
            d = first_rd.get(ver, '')
            return [edition(m, '%s: %s%s' % (what, re.sub(r'^TVS\s+', '', ver), ' (%s)' % d.replace('RELEASE DATE :', 'released').strip() if d else ''))]
        colour_rows(nb, b2, f, new, eds_of, 0)
        stats['colour versions left out (no new painted part)'] += len({k[2] for k in blocks} - {k[2] for k, _ in new})
    # one line for the editions that print it the same (the e-catalogue's way)
    lines = collections.OrderedDict()
    for nb, g, asm, ref, pn, desc, qty, rem, ns, eds in rows:
        lines.setdefault((nb, g, asm, ref, pn, desc, qty, rem, ns), set()).update(eds)
    asm_i = {a: i for i, a in enumerate(o['asm'])}
    desc_i = {d: i for i, d in enumerate(o['desc'])}

    def idx(table, ix, s):
        if s not in ix:
            ix[s] = len(table)
            table.append(s)
        return ix[s]
    for (nb, g, asm, ref, pn, desc, qty, rem, ns), eds in lines.items():
        o['d'].append([nb, -1, g, idx(o['asm'], asm_i, asm), ref, pn, idx(o['desc'], desc_i, desc), qty, rem, ns, 0, sorted(eds)])
    o['bk'] += bk_new
    o['vars'] += vars_
    o['vmod'].update(vmod)
    have = set(o['vmod'].values())
    ms = [m for m in TVS_ORDER if m in have] + sorted(have - set(TVS_ORDER))
    cur = set(TVS_ORDER) - TVS_OLD
    o['models'] = [m for m in ms if m in cur] + [m for m in ms if m not in cur]
    assert o['models'][:o['cur']] == [m for m in ms if m in cur]
    o['cc'] = {m: TVS_CC[m] for m in o['models'] if TVS_CC.get(m)}
    # diagrams: each edition's figure -> its page of the book
    fi = {u: i for i, u in enumerate(dg['f'])}
    f0, used = len(dg['f']), set()
    for (v, g), (f, page) in sorted(links.items()):
        size = B[f]['size']
        assert 1 <= page <= B[f]['pages'], (f, page)
        u = 'pdf:%s/%s:0:%d:%d:0,0,1,1' % (TVS_DIR, f, size, page)
        if u not in fi:
            fi[u] = len(dg['f'])
            dg['f'].append(u)
        dg['p'].setdefault(str(v), {})[g] = fi[u]
        used.add(f)
    need = [f for f in list(TVS_BOOKS) + list(TVS_COLOUR) if f in used]
    assert need == list(TVS_BOOKS) + list(TVS_COLOUR), (set(TVS_BOOKS) | set(TVS_COLOUR)) - used
    dg['files'] = dg['files'] + [[f, B[f]['size']] for f in need]
    dg['spare'] = [[f, B[f]['size']] for f in sorted(TVS_SPARE)]
    dg['pick'] = 'All Catalogues > TVS ALL'
    for name, obj in (('tvs.json', o), ('tvs-dg.json', dg)):
        jdump(os.path.join(OUT, name), obj)
    new_pn = {r[5] for r in o['d'][len(o['d']) - len(lines):] if r[5]} - have_pn
    mrp_n = sum(1 for (nb, g, a, ref, pn, *_ ) in lines if pn and pn in o.get('mrp', {}))
    print('TVS books: %d books + %d colour lists | %d editions (%d models) | %d lines (%d book lines before merging) | %d part no. new to TVS | '
          'lines with an MRP (part no. in the MRP list) %d | %d new diagram pages | %s'
          % (len(TVS_BOOKS), len(TVS_COLOUR), len(vars_), len(set(vmod.values())), len(lines), len(rows), len(new_pn), mrp_n,
             len(dg['f']) - f0, dict(stats)))
    print('   files to upload: %d PDFs, %.1f MB | spare %d PDFs, %.1f MB' % (len(need), sum(B[f]['size'] for f in need) / 1e6, len(TVS_SPARE),
                                                                         sum(B[f]['size'] for f in TVS_SPARE) / 1e6))
    return o, dg



# ---------------------------------------------------------------- Bajaj (bajaj_parts.csv + the 90 PDF books in bajaj_spc_pdfs.zip)
# The diagrams are pages of Bajaj's own PDF books. The zip stores each PDF unchanged, so a diagram link names the PDF's bytes in the
# zip, the page and the part of the page that is the diagram: 'pdf:catalogues/bajaj/bajaj_spc_pdfs.zip:<offset>:<length>:<page>:<x0>,<y0>,<x1>,<y1>'
# (fractions of the page as shown). The app draws that page from the original PDF (vector lines stay sharp at every zoom).
# Page and crop of every plate, and the plate fixes of the 12 oldest books: bj_pages.py -> bajaj-pages.json.
BJ_ZIP = 'bajaj_spc_pdfs.zip'
BJ_NB = {'BS3': 0, 'BS4': 1, 'BS6': 2, 'OBD2': 3, 'OBD2B': 4, '': 5}
# book (path inside pdfs/) -> (models, norm, what the book covers, its date)
BJ_BOOKS = {
    "Avenger/Avenger 150/BS4/Avenger 150 Street BS IV_Rev 15 Aug'18.pdf": (['Avenger 150 Street'], 'BS4', 'BS4', 'Aug 2018'),
    "Avenger/Avenger 160/BS4/Avenger 160 Street _Rev 03 Dec'19.pdf": (['Avenger 160 Street'], 'BS4', 'BS4', 'Dec 2019'),
    'Avenger/Avenger 160/BS6/Avenger 160 Street BS VI _OBD2_Rev 11 FEB 25.pdf': (['Avenger 160 Street'], 'OBD2', 'BS6 OBD2', 'Feb 2025'),
    "Avenger/Avenger 180/BS4/Avenger 180 Street _Rev 05 Feb'19.pdf": (['Avenger 180 Street'], 'BS4', 'BS4', 'Feb 2019'),
    "Avenger/Avenger 220/BS4/Avenger 220 Cruise & Street ABS _Rev 26 Dec'19.pdf": (['Avenger 220'], 'BS4', 'Cruise & Street ABS · BS4', 'Dec 2019'),
    'Avenger/Avenger 220/BS6/Avenger 220 Cruise & Street BS VI _Rev 11 Feb 25.pdf': (['Avenger 220'], 'BS6', 'Cruise & Street · BS6', 'Feb 2025'),
    "CT/CT 100/BS4/CT100ES BS IV CBS _Rev 27 Sep'20.pdf": (['CT 100'], 'BS4', 'ES CBS · BS4', 'Sep 2020'),
    "CT/CT 100/BS4/CT100KS CBS _Rev 04 Sep'20.pdf": (['CT 100'], 'BS4', 'KS CBS · BS4', 'Sep 2020'),
    "CT/CT 100/BS6/CT 100 ES & KS BS VI UG _Rev 05 Aug'21.pdf": (['CT 100'], 'BS6', 'ES & KS · BS6', 'Aug 2021'),
    "CT/CT 110/BS4/CT 110 ES CBS _Rev 06 Sep'20.pdf": (['CT 110'], 'BS4', 'ES CBS · BS4', 'Sep 2020'),
    "CT/CT 110/BS6/CT 110 ES & KS BS VI UG _Rev 05 Aug'21.pdf": (['CT 110'], 'BS6', 'ES & KS · BS6', 'Aug 2021'),
    'CT/CT 110X/BS6/New CT 110X -FI Rev01 Nov 25.pdf': (['CT 110X'], 'BS6', 'New, FI · BS6', 'Nov 2025'),
    'CT/CT 110X/BS6/OBD2-CT 110X BS VI UG _Rev 09 Apr 2023.pdf': (['CT 110X'], 'OBD2', 'BS6 OBD2', 'Apr 2023'),
    "CT/CT 125X/BS6/CT 125X BS VI UG _OBD2_Rev02 Apr'23.pdf": (['CT 125X'], 'OBD2', 'BS6 OBD2', 'Apr 2023'),
    'Discover/Discover 100/Discover 100 _Rev 03 Jul 13 (4G 5G).pdf': (['Discover 100'], '', '4 & 5 gear', 'Jul 2013'),
    "Discover/Discover 110-125/Old/Discover 110, 125UG, 110 CBS & 125 CBS _Rev 11 Sep'20.pdf": (['Discover 110', 'Discover 125'], '', 'UG & CBS', 'Sep 2020'),
    "Discover/Discover 125-100 M/Discover 125 M 100 M Rev 04 Dec'14.pdf": (['Discover 125M', 'Discover 100M'], '', '', 'Dec 2014'),
    "Discover/Discover 125/Discover 125 (Reintroduced) _Rev 07 Oct'16.pdf": (['Discover 125'], '', 'Reintroduced', 'Oct 2016'),
    "Discover/Discover 150/Discover 150S & 150F Rev 04 May'15.pdf": (['Discover 150'], '', '150S & 150F', 'May 2015'),
    "Dominar/Dominar 250/BS6/Dominar 250 BS VI _Rev 10 Jan' 25.pdf": (['Dominar 250'], 'BS6', 'BS6', 'Jan 2025'),
    "Dominar/Dominar 400/BS4/Dominar 400 UG _Rev 24 Jan'20.pdf": (['Dominar 400'], 'BS4', 'UG · BS4', 'Jan 2020'),
    "Dominar/Dominar 400/BS6/Dominar 400_Rev 10_Jan'_25.pdf": (['Dominar 400'], 'BS6', 'BS6', 'Jan 2025'),
    'Dominar/Dominar 400/BS6/New Dominar 400_Rev 00 Mar 26.pdf': (['Dominar 400'], 'BS6', 'New · BS6', 'Mar 2026'),
    'FREEDOM/Freedom NG04_REV03_Mar 26.pdf': (['Freedom 125 CNG'], 'BS6', 'NG04 · BS6', 'Mar 2026'),
    'Obsoleted Models/Avenger 180.pdf': (['Avenger 180'], '', 'DTS-i', ''),
    "Obsoleted Models/Avenger 200 _Rev 00 Sep'08.pdf": (['Avenger 200'], '', '', 'Sep 2008'),
    "Obsoleted Models/Avenger 220 _Rev 02 Sep'15.pdf": (['Avenger 220'], '', '', 'Sep 2015'),
    "Obsoleted Models/Boxer 150 ES _Rev 00 Sep'11.pdf": (['Boxer 150'], '', 'ES', 'Sep 2011'),
    "Obsoleted Models/Discover 125ST, T 100T _Rev 05 Nov'14.pdf": (['Discover 125 ST', 'Discover 125T', 'Discover 100T'], '', '', 'Nov 2014'),
    "Obsoleted Models/Discover 135 Drum Disc _Rev 00 Dec'07.pdf": (['Discover 135'], '', 'Drum & Disc', 'Dec 2007'),
    'Obsoleted Models/Discover 135 UG Modified Final.pdf': (['Discover 135'], '', 'UG (modified)', ''),
    "Obsoleted Models/Kristal _Rev 00 Oct'07.pdf": (['Kristal'], '', '', 'Oct 2007'),
    "Obsoleted Models/Platina 125 DTS-i ES KS _Rev 00 Dec'08.pdf": (['Platina 125'], '', 'DTS-i ES & KS', 'Dec 2008'),
    "Obsoleted Models/Platina 125 _Rev 00 Mar'10.pdf": (['Platina 125'], '', '', 'Mar 2010'),
    "Obsoleted Models/Platina _Rev 02 Oct'12.pdf": (['Platina 100'], '', '', 'Oct 2012'),
    "Obsoleted Models/Pulsar 135 LS BS III _Rev 03 Jun'16.pdf": (['Pulsar 135 LS'], 'BS3', 'BS3', 'Jun 2016'),
    "Obsoleted Models/Pulsar 150 BS III UG4 _Rev 13 Nov'16.pdf": (['Pulsar 150'], 'BS3', 'UG4 · BS3', 'Nov 2016'),
    "Obsoleted Models/Pulsar 180 BS III _Rev 05 Oct'14.pdf": (['Pulsar 180'], 'BS3', 'BS3', 'Oct 2014'),
    "Obsoleted Models/Pulsar 180 UG3 _Rev 00 Nov'07.pdf": (['Pulsar 180'], '', 'UG3', 'Nov 2007'),
    "Obsoleted Models/Pulsar 200 _Rev 00 Sep'08.pdf": (['Pulsar 200'], '', 'DTS-i', 'Sep 2008'),
    "Obsoleted Models/Pulsar 220  BS III_ Rev 03 Dec'14.pdf": (['Pulsar 220F'], 'BS3', 'BS3', 'Dec 2014'),
    "Obsoleted Models/Pulsar 220 DTS-FI _Rev 00 Dec'07.pdf": (['Pulsar 220F'], '', 'DTS-Fi', 'Dec 2007'),
    "Obsoleted Models/Pulsar AS 150 & NS 150 _Rev 07 Dec'16.pdf": (['Pulsar AS150', 'Pulsar NS150'], '', '', 'Dec 2016'),
    'Obsoleted Models/Pulsar UG III 150 SPC.pdf': (['Pulsar 150'], '', 'UG3', ''),
    "Obsoleted Models/XCD 125 _Rev 00 Sep'08.pdf": (['XCD 125'], '', '', 'Sep 2008'),
    "Obsoleted Models/XCD 135 _Rev 00 Mar'09 1.pdf": (['XCD 135'], '', '', 'Mar 2009'),
    'Platina/Platina 100/BS4/Platina 100ES, KS & CBS _Rev 20 Jan 20.pdf': (['Platina 100'], 'BS4', 'ES, KS & CBS · BS4', 'Jan 2020'),
    "Platina/Platina 100/BS6/New Platina 100_ES&KS_Rev 02 Nov'25.pdf": (['Platina 100'], 'BS6', 'New, ES & KS · BS6', 'Nov 2025'),
    'Platina/Platina 100/BS6/Platina 100 BS VI _OBD2_Rev 12 Sep-23.pdf': (['Platina 100'], 'OBD2', 'BS6 OBD2', 'Sep 2023'),
    "Platina/Platina 110/BS4/Platina 110 CBS Drum & Disc _Rev 08 Sep'20.pdf": (['Platina 110'], 'BS4', 'CBS Drum & Disc · BS4', 'Sep 2020'),
    "Platina/Platina 110/BS6/New Platina 110_FI_Rev 02 Nov'25.pdf": (['Platina 110'], 'BS6', 'New, FI · BS6', 'Nov 2025'),
    "Platina/Platina 110/BS6/Platina 110 Disc BS VI UG_OBD2_Rev 11 Apr'23.pdf": (['Platina 110'], 'OBD2', 'Disc · BS6 OBD2', 'Apr 2023'),
    "Pulsar/Pulsar 125/BS4/Pulsar 125 _Rev 06 Sep'20.pdf": (['Pulsar 125'], 'BS4', 'BS4', 'Sep 2020'),
    "Pulsar/Pulsar 125/BS6/Pulsar 125 BS VI_OBD2B_LED_ headlamp_Rev 01 MAR' 26.pdf": (['Pulsar 125'], 'OBD2B', 'LED headlamp · BS6 OBD2B', 'Mar 2026'),
    'Pulsar/Pulsar 125/BS6/Pulsar 125 BS VI_OBD2_Rev 19 Apr 25.pdf': (['Pulsar 125'], 'OBD2', 'BS6 OBD2', 'Apr 2025'),
    "Pulsar/Pulsar 135/BS4/Pulsar 135 UG & BS IV_Rev 13 Aug'18.pdf": (['Pulsar 135 LS'], 'BS4', 'UG · BS4', 'Aug 2018'),
    "Pulsar/Pulsar 150/BS4/Pulsar 150 BS IV (Classic, Regular, Neon) ABS _Rev 27 Sep'20.pdf": (['Pulsar 150'], 'BS4', 'Classic, Regular, Neon, ABS · BS4', 'Sep 2020'),
    "Pulsar/Pulsar 150/BS4/Pulsar 150 Twin Disc_Rev 12 Nov'19.pdf": (['Pulsar 150 Twin Disc'], 'BS4', 'BS4', 'Nov 2019'),
    "Pulsar/Pulsar 150/BS6/Pulsar 150 BS VI_OBD2B_Rev 18 Apr' 25.pdf": (['Pulsar 150'], 'OBD2B', 'BS6 OBD2B', 'Apr 2025'),
    'Pulsar/Pulsar 150/BS6/Pulsar 150 BS_VI_(LED Headlamp)OBD2B_Rev 01 Mar 26.pdf': (['Pulsar 150'], 'OBD2B', 'LED headlamp · BS6 OBD2B', 'Mar 2026'),
    'Pulsar/Pulsar 150/BS6/Pulsar 150 Twin Disc (LED Headlamp) BS VI_OBD2B_Rev 01 Mar 26.pdf': (['Pulsar 150 Twin Disc'], 'OBD2B', 'LED headlamp · BS6 OBD2B', 'Mar 2026'),
    "Pulsar/Pulsar 150/BS6/Pulsar 150 Twin Disc BS VI_OBD2B_Rev 16 Apr' 25.pdf": (['Pulsar 150 Twin Disc'], 'OBD2B', 'BS6 OBD2B', 'Apr 2025'),
    'Pulsar/Pulsar 180/BS4/SPC_PULSAR 180 HF ABS_REV_24 JAN 20.pdf': (['Pulsar 180'], 'BS4', 'Half fairing, ABS · BS4', 'Jan 2020'),
    'Pulsar/Pulsar 180/BS6/New Pulsar 180 BS VI _Rev 00 MAR 26.pdf': (['Pulsar 180'], 'BS6', 'New · BS6', 'Mar 2026'),
    'Pulsar/Pulsar 180/BS6/Pulsar 180 BS VI _Rev 11 Jan 24.pdf': (['Pulsar 180'], 'BS6', 'BS6', 'Jan 2024'),
    "Pulsar/Pulsar 220/BS4/Pulsar 220 ABS _Rev 21 Dec'19.pdf": (['Pulsar 220F'], 'BS4', 'ABS · BS4', 'Dec 2019'),
    'Pulsar/Pulsar 220/BS6/Pulsar 220F_OBD2B_LED Indicators_Rev 16 Apr 26.pdf': (['Pulsar 220F'], 'OBD2B', 'LED indicators · BS6 OBD2B', 'Apr 2026'),
    "Pulsar/Pulsar AS-NS 150/Pulsar AS 150 NS 150 _Rev 06 Mar'16.pdf": (['Pulsar AS150', 'Pulsar NS150'], '', '', 'Mar 2016'),
    "Pulsar/Pulsar N 125/Pulsar N 125 OBD2_Rev 01 Mar' 25.pdf": (['Pulsar N125'], 'OBD2', 'BS6 OBD2', 'Mar 2025'),
    'Pulsar/Pulsar N 150/SPC_Pulsar_N150 Twin Disc_Rev01_Jan_24.pdf': (['Pulsar N150'], 'BS6', 'Twin Disc · BS6', 'Jan 2024'),
    'Pulsar/Pulsar N-F 250/BS6/Pulsar F250 & N250_Rev OBD2B_10 Mar 25.pdf': (['Pulsar N250', 'Pulsar F250'], 'OBD2B', 'BS6 OBD2B', 'Mar 2025'),
    'Pulsar/Pulsar N160/BS6/Pulsar N160 4 Valve_Rev00_Aug 26.pdf': (['Pulsar N160'], 'BS6', '4 valve · BS6', 'Aug 2026'),
    'Pulsar/Pulsar N160/BS6/Pulsar N160 USD Single Seat Fork-Rev01 Apr 26.pdf': (['Pulsar N160'], 'BS6', 'USD fork, single seat · BS6', 'Apr 2026'),
    'Pulsar/Pulsar N160/BS6/Pulsar N160_OBD2B_Rev 10 Apr 26.pdf': (['Pulsar N160'], 'OBD2B', 'BS6 OBD2B', 'Apr 2026'),
    'Pulsar/Pulsar N160/BS6/Pulsar_New N 160 UG Twin Disc-Rev 03 Feb 26.pdf': (['Pulsar N160'], 'BS6', 'New UG, Twin Disc · BS6', 'Feb 2026'),
    'Pulsar/Pulsar NS 125/BS6/Pulsar NS 125 FI BS VI_UG_Rev 12 Sep 26.pdf': (['Pulsar NS125'], 'BS6', 'FI UG · BS6', 'Sep 2026'),
    "Pulsar/Pulsar NS 160/BS4/Pulsar NS160 Single, Twin Disc & ABS _Rev 10 Mar'20.pdf": (['Pulsar NS160'], 'BS4', 'Single & Twin Disc, ABS · BS4', 'Mar 2020'),
    "Pulsar/Pulsar NS 160/BS6/New Pulsar NS 160 BS VI_UG_Rev00 Apr' 25.pdf": (['Pulsar NS160'], 'BS6', 'New UG · BS6', 'Apr 2025'),
    "Pulsar/Pulsar NS 160/BS6/Pulsar NS 160 BS VI USD Fork _OBD2_Rev 11 Feb'24.pdf": (['Pulsar NS160'], 'OBD2', 'USD fork · BS6 OBD2', 'Feb 2024'),
    "Pulsar/Pulsar NS-AS 200/BS4/Pulsar NS 200 ABS & AS 200 19 Jul'20-.pdf": (['Pulsar NS200', 'Pulsar AS200'], 'BS4', 'BS4', 'Jul 2020'),
    "Pulsar/Pulsar NS-AS 200/BS6/Pulsar NS 200 BS VI UG _Rev 00 Apr' 25.pdf": (['Pulsar NS200'], 'BS6', 'UG · BS6', 'Apr 2025'),
    "Pulsar/Pulsar NS-AS 200/BS6/Pulsar NS 200 BS VI USD Fork _Rev 13 Nov'24'.pdf": (['Pulsar NS200'], 'BS6', 'USD fork · BS6', 'Nov 2024'),
    'Pulsar/Pulsar NS400 Z/SPC_Pulsar NS 400_Rev 03 Sep 26.pdf': (['Pulsar NS400Z'], 'BS6', 'BS6', 'Sep 2026'),
    'Pulsar/Pulsar P150/BS6/OBD2-SPC_Pulsar_P150_REV01_APR_23.pdf': (['Pulsar P150'], 'OBD2', 'BS6 OBD2', 'Apr 2023'),
    'Pulsar/Pulsar RS200/BS4/SPC_Pulsar RS 200_REV_20_NOV_2019.pdf': (['Pulsar RS200'], 'BS4', 'BS4', 'Nov 2019'),
    'Pulsar/Pulsar RS200/BS6/Pulsar RS 200 BS VI_OBD2_Rev 11 Oct 24.pdf': (['Pulsar RS200'], 'OBD2', 'BS6 OBD2', 'Oct 2024'),
    "Pulsar/Pulsar RS200/BS6/Pulsar RS 200 BS VI_UG_OBD2B_Rev01 May' 25.pdf": (['Pulsar RS200'], 'OBD2B', 'UG · BS6 OBD2B', 'May 2025'),
    "V/V125 & 150/V15 & V12 BS IV UG _Rev 14 Aug'18.pdf": (['V15', 'V12'], 'BS4', 'UG · BS4', 'Aug 2018'),
}
# the same book twice in Sam's package: an identical file (same CRC) is read once; the 2nd copy of the NS160 book differs in
# 2 lines, so its lines join the first copy's edition and its diagrams are the first copy's (the larger, sharper file)
BJ_SAME = {"Discover/Discover 125/Discover 110, 125UG, 110 CBS 125 CBS _Rev 11 Sep'20.pdf": None,
           "Pulsar/Pulsar NS 160/BS4/Pulsar NS160 Single, Twin Disc & ABS _Rev 10 Mar'20 (2).pdf":
               "Pulsar/Pulsar NS 160/BS4/Pulsar NS160 Single, Twin Disc & ABS _Rev 10 Mar'20.pdf"}
BJ_ORDER = ['Pulsar 150', 'Pulsar 125', 'Pulsar 150 Twin Disc', 'Pulsar N160', 'Pulsar N150', 'Pulsar N125', 'Pulsar NS125', 'Pulsar NS160',
            'Pulsar NS200', 'Pulsar RS200', 'Pulsar N250', 'Pulsar F250', 'Pulsar NS400Z', 'Pulsar 220F', 'Pulsar 180', 'Pulsar P150',
            'Platina 100', 'Platina 110', 'CT 110X', 'CT 125X', 'CT 100', 'CT 110', 'Freedom 125 CNG', 'Avenger 160 Street', 'Avenger 220',
            'Dominar 400', 'Dominar 250',
            'Avenger 150 Street', 'Avenger 180 Street', 'Pulsar 135 LS', 'Pulsar AS150', 'Pulsar NS150', 'Pulsar AS200', 'Pulsar 200', 'V15', 'V12',
            'Discover 125', 'Discover 110', 'Discover 100', 'Discover 150', 'Discover 125M', 'Discover 100M', 'Discover 125 ST', 'Discover 125T',
            'Discover 100T', 'Discover 135', 'Platina 125', 'XCD 125', 'XCD 135', 'Boxer 150', 'Kristal', 'Avenger 180', 'Avenger 200']
BJ_CUR = 27
BJ_CC = {'Pulsar 150': 150, 'Pulsar 125': 125, 'Pulsar 150 Twin Disc': 150, 'Pulsar N160': 160, 'Pulsar N150': 150, 'Pulsar N125': 125, 'Pulsar NS125': 125,
         'Pulsar NS160': 160, 'Pulsar NS200': 200, 'Pulsar RS200': 200, 'Pulsar N250': 250, 'Pulsar F250': 250, 'Pulsar NS400Z': 400, 'Pulsar 220F': 220,
         'Pulsar 180': 180, 'Pulsar P150': 150, 'Platina 100': 100, 'Platina 110': 110, 'CT 110X': 110, 'CT 125X': 125, 'CT 100': 100, 'CT 110': 110,
         'Freedom 125 CNG': 125, 'Avenger 160 Street': 160, 'Avenger 220': 220, 'Dominar 400': 400, 'Dominar 250': 250, 'Avenger 150 Street': 150,
         'Avenger 180 Street': 180, 'Pulsar 135 LS': 135, 'Pulsar AS150': 150, 'Pulsar NS150': 150, 'Pulsar AS200': 200, 'Pulsar 200': 200, 'V15': 150,
         'V12': 125, 'Discover 125': 125, 'Discover 110': 110, 'Discover 100': 100, 'Discover 150': 150, 'Discover 125M': 125, 'Discover 100M': 100,
         'Discover 125 ST': 125, 'Discover 125T': 125, 'Discover 100T': 100, 'Discover 135': 135, 'Platina 125': 125, 'XCD 125': 125, 'XCD 135': 135,
         'Boxer 150': 150, 'Kristal': 95, 'Avenger 180': 180, 'Avenger 200': 200}
BJ_DASH = {'-', '--', '---', '–', '–-', '–--', '—', '- -', '-–', '—-'}


class PdfPages:
    """Bajaj's diagrams: a page (and the diagram's part of it) of a PDF stored in bajaj_spc_pdfs.zip -> 'pdf:' link"""
    def __init__(self, dirn, zname, plates):
        self.dir, self.z, self.plates, self.used = dirn, zname, plates, False
        self.ix, self.size = zidx(zname)
        self.at = {}

    def link(self, k):
        book, plate = k.split('|', 1)
        e, pc = self.ix.get('pdfs/' + book), self.plates.get('pdfs/' + book + '|' + plate)
        if not e or not pc:
            return None
        meth, off, csize, usize = e
        assert meth == 0 and csize == usize and off + usize <= self.size, ('stored PDF expected', book)
        self.used = True
        return 'pdf:%s/%s:%d:%d:%d:%s' % (self.dir, self.z, off, usize, pc[0], ','.join('%g' % round(v, 4) for v in pc[1:]))

    def files(self):
        return [[self.z, self.size]] if self.used else []


def bj_label(m, qual, date):
    mon, _, yr = date.partition(' ')
    return m + (' · ' + qual if qual else '') + (' (%s, %s)' % (mon, yr) if date else '')


def bj_fix(desc, qty, rem):
    """the CSV's quantity / remarks columns as the books mean them"""
    rem = rem.rstrip('\\').strip()
    ns = False
    if rem in BJ_DASH or rem in ('Y', 'Qty.', 'Qty'):
        rem = ''
    elif rem in ('N', 'Not Serviceable', 'Not serviceable'):
        ns, rem = True, ''
    if qty in ('Qty.', 'Qty', '-', '--', '–'):
        qty = ''
    if qty == '441' and desc.endswith('Petamo GHY'):        # the grease is "Petamo GHY 441"
        desc, qty = desc + ' 441', ''
    if not qty and re.fullmatch(r'\d{1,3}(\s+\d{1,3})*', rem):   # the quantity printed in the remarks column ("1   1" = 1 per variant)
        ns_ = rem.split()
        qty, rem = (ns_[0] if len(set(ns_)) == 1 else ' / '.join(ns_)), ''
    if qty and rem == qty:
        rem = ''
    if not qty and rem and (re.search(r'[&/\-(,]$', desc) or desc.count('(') > desc.count(')') or re.search(r'\b(WITH|FOR|OF|AND)$', desc, re.I)) \
            and not re.match(r'\d+(\.\d+)?\s*(ml|gm|g|ltr|l|kg)\b', rem, re.I):
        desc, rem = desc + ' ' + rem, ''                    # a description that runs on into the remarks column
    return desc, qty, rem, ns


def bajaj():
    pages = json.load(open(os.path.join(HERE, 'bajaj-pages.json')))
    raw = open(os.path.join(os.path.dirname(HERE) if False else IMP, 'Bajaj', 'bajaj_parts.csv'), 'rb').read()
    import hashlib
    assert hashlib.md5(raw).hexdigest() == pages['csv_md5'], 'bajaj-pages.json was made from another bajaj_parts.csv: run bj_pages.py again'
    drop, relabel = set(pages['drop']), pages['relabel']
    N = lambda s: re.sub(r'\s+', '', s or '').upper().lstrip('0') or '0'
    rows, seen_books, stats = [], set(), collections.Counter()
    for i, r in enumerate(rows_csv('Bajaj/bajaj_parts.csv')):
        book = r['pdf_file'][5:]
        assert r['pdf_file'].startswith('pdfs/') and (book in BJ_BOOKS or book in BJ_SAME), book
        if book in BJ_SAME and BJ_SAME[book] is None:
            stats['twin book (read once)'] += 1
            continue
        if i in drop:
            stats['dropped (no plate on that page)'] += 1
            continue
        plate, sect = N(r['Plate No.']), r['Section']
        if str(i) in relabel:
            plate, sect = relabel[str(i)]
            stats['plate fixed'] += 1
        pic = BJ_SAME.get(book) or book
        if 'pdfs/' + pic + '|' + plate not in pages['plates']:
            pic = book
        ms, norm, qual, date = BJ_BOOKS.get(BJ_SAME.get(book) or book)
        seen_books.add(BJ_SAME.get(book) or book)
        desc, qty, rem, ns = bj_fix(clean(r['Description']), clean(r['Qty']), clean(r['Remarks']))
        for m in ms:
            rows.append(dict(m=m, v=bj_label(m, qual, date), y=int(date[-4:]) if date else 0, nb=BJ_NB[norm], g=plate, asm=clean(sect).title() if clean(sect).isupper() else clean(sect),
                             ref=clean(r['Ref No.']), pn=clean(r['Part No.']), desc=desc, qty=qty, rem=rem, ns=ns, dg=pic + '|' + plate))
    assert seen_books == set(BJ_BOOKS), set(BJ_BOOKS) - seen_books
    unknown = {x['m'] for x in rows} - set(BJ_ORDER)
    assert not unknown, unknown
    print('Bajaj rows:', dict(stats))
    zp = PdfPages('catalogues/bajaj', BJ_ZIP, pages['plates'])
    bk = [[0, ['BS3'], ''], [0, ['BS4'], ''], [0, ['BS6'], ''], [0, ['BS6'], 'BS6 Phase 2 (OBD2)'], [0, ['BS6'], 'BS6 Phase 2 (OBD2B)'], [0, [], '']]
    return pack('Bajaj', 'bajaj', "Bajaj Auto spare parts catalogues (90 books), Sam's Bajaj package, Oct 2026", BJ_ORDER, set(BJ_ORDER[:BJ_CUR]), rows,
                'Plate', bk, BJ_CC, zp, pick=PKG + 'Bajaj')


# ---------------------------------------------------------------- Suzuki (Sam's folder All Catalogues > Suzuki, 11 Oct 2026)
# Suzuki's own parts catalogues (PDF: each figure's page has its diagram with the parts table beside it) and Suzuki's price list
# 03-Oct'26.pdf (MRP). sz_parse.py reads the tables and the price list -> suzuki-pages.json. A figure's diagram is its page of the
# book, drawn from the original PDF that Admin uploads unchanged: 'pdf:catalogues/suzuki/<book>:0:<size>:<page>:0,0,1,1' (the
# whole page: some diagrams' numbers sit right of the table). Price-list parts that no book has are listed too, without a diagram.
SZ_DIR = 'catalogues/suzuki'
# book -> (model, the book's model code, the bikes' years it is for: from the book's model-year code and date to the next book)
SZ_BOOKS = {
    '2025NEW ACCESS.pdf': ('Access 125', 'UZ125M5', '2025+'),
    'UZ125NDM3_P74(9900BF33GP0C026)-2022.pdf': ('Access 125', 'UZ125M3', '2022-2024'),
    'UZ125NDM2_P74-2021.pdf': ('Access 125', 'UZ125M2', '2022'),
    'UZ125M2_P74_7-2021.pdf': ('Access 125', 'UZ125M2', '2021'),            # on Sam's Mac also as UZ125NEM0_P74_14-2021.pdf (same file)
    'UZ125NEWM0_P74(9900BF33GP0C017)-2020.pdf': ('Access 125', 'UZ125NEW/NDWM0', '2020'),
    'UZ125EHCM0_P74-19.pdf': ('Access 125', 'UZ125EHCM0', '2019-2020'),      # also as "eh access.pdf" and "eh access 2.pdf"
    'UZ125EECEHCL9P74-18.pdf': ('Access 125', 'UZ125EEC/EHCL9', '2018-2019'),
    'UZ125EEL7-16.pdf': ('Access 125', 'UZ125EEL7', '2016-2018'),
    'UZ125EHL7_f-18.pdf': ('Access 125', 'UZ125EHL7', '2016-2018'),
    'UZ125EZL7_d-18.pdf': ('Access 125', 'UZ125EH/EZL7', '2016-2018'),
    'LR_PCAT_UZ125L2_2011_P74.pdf': ('Access 125', 'UZ125L2', '2011-2015'),  # also as "LR_PCAT_UZ125L2_2011_P74 2.pdf"
    'oct150cc gixxer.txt': ('Gixxer SF', 'GSX150FRLM6/FRLXM6', '2025+'),    # a PDF, named .txt on Sam's Mac
    'PARTS_CATALOGUE_GSX150FRLZM2_P74.pdf': ('Gixxer SF', 'GSX150FRLZM2', '2020-2024'),
    'PARTS_CATALOGUE_GSX150RLZM2_P74.pdf': ('Gixxer', 'GSX150RLZM2', '2020+'),
    'LR_PCAT_GSX150L6-2_P74.pdf': ('Gixxer', 'GSX150L6', '2015-2019'),
    'PARTS_CATALOGUE_GSX250FRL-RLZM2_P74.pdf': ('Gixxer SF 250', 'GSX250FRL/FRLZM2', '2020+'),
    'PARTS_CATALOGUE_GSX250RL-RLZM2_P74.pdf': ('Gixxer 250', 'GSX250RL/RLZM2', '2020+'),
    'DS250RLM3_P74_6.pdf': ('V-Strom SX', 'DS250RLM3', '2022+'),
    'GL150GLAL9_P74_14.pdf': ('Intruder 150', 'GL150GLA/RLAL9', '2019-2021'),  # 1st edition Feb 2019; price list has GL150 codes up to M1 (2021)
    'GE110DDAL8_P74.pdf': ('Hayate EP', 'GE110DDAL8', '2017-2020'),          # also as GE110DDAL8_EP-HAYATE.pdf
    'hayate2014.pdf': ('Hayate', 'GE110L5/DL5', '2014-2016'),
    'LR_PCAT_GT125K6K7_2nd_P74(9900B-20089-010).pdf': ('Zeus', 'GT125K6/K7', '2006-2009'),
    # scanned books (pictures of the pages): read by OCR (sz_ocr.py), model codes and years from their covers
    'LR_PCAT_GSX150FRFAL7_P74_2.pdf': ('Gixxer SF', 'GSX150FRFA/FRZAL7', '2017'),
    'LR_PCAT_UP125L6_P74.pdf': ('Swish 125', 'UP125L6', '2016'),
    'LR_PCAT_GS150CDL3_7-1-2013.pdf': ('GS150R', 'GS150RL3', '2013'),
    'LR_PCAT_GS150R_P74.pdf': ('GS150R', 'GS150RK9/RL0', '2009-2010'),
    'LR_PCAT_GT125CDL3_2013_P74.pdf': ('Slingshot', 'GT125(S)L3', '2013'),
}
# scanned books still to be read (none now: sz_ocr.py reads the 5 scans into suzuki-pages.json)
SZ_SCANS = []
# books whose text came from OCR: their words and titles need fixing (the first two had an OCR text layer; the rest were read by sz_ocr.py)
SZ_OCR = {'LR_PCAT_UZ125L2_2011_P74.pdf', 'LR_PCAT_GT125K6K7_2nd_P74(9900B-20089-010).pdf', 'LR_PCAT_GSX150FRFAL7_P74_2.pdf',
          'LR_PCAT_UP125L6_P74.pdf', 'LR_PCAT_GS150CDL3_7-1-2013.pdf', 'LR_PCAT_GS150R_P74.pdf', 'LR_PCAT_GT125CDL3_2013_P74.pdf'}
SZ_ORDER = ['Access 125', 'Burgman Street 125', 'Avenis 125', 'e-Access', 'Gixxer SF', 'Gixxer', 'Gixxer SF 250', 'Gixxer 250', 'V-Strom SX',
            'Intruder 150', "Let's", 'Hayate EP', 'Hayate', 'Swish 125', 'Zeus', 'Slingshot', 'GS150R',
            'Hayabusa', 'V-Strom 800DE', 'V-Strom 650', 'V-Strom 1000', 'GSX-S1000', 'GSX-S750', 'Inazuma 250', 'Intruder M1800R', 'Other models']
SZ_CUR = SZ_ORDER[:9]
SZ_CC = {'Access 125': 125, 'Burgman Street 125': 125, 'Avenis 125': 125, 'Gixxer SF': 155, 'Gixxer': 155, 'Gixxer SF 250': 249, 'Gixxer 250': 249,
         'V-Strom SX': 249, 'Intruder 150': 155, "Let's": 113, 'Hayate EP': 113, 'Hayate': 113, 'Swish 125': 125, 'Zeus': 125, 'Slingshot': 125,
         'GS150R': 150, 'Hayabusa': 1340, 'V-Strom 800DE': 776, 'V-Strom 650': 645, 'V-Strom 1000': 1037, 'GSX-S1000': 999, 'GSX-S750': 749,
         'Inazuma 250': 248, 'Intruder M1800R': 1783}
# the price list's MODEL column (Suzuki's model codes) -> the model
SZ_FAMILY = [(r'^(UZ|XB29)', 'Access 125'), (r'^UB', 'Burgman Street 125'), (r'^UN', 'Avenis 125'), (r'^ESC003', 'e-Access'),
             (r'^(UP|XD80)', 'Swish 125'), (r'^(UR110|XE34)', "Let's"), (r'^GL150', 'Intruder 150'), (r'^GSX250F', 'Gixxer SF 250'),
             (r'^GSX250', 'Gixxer 250'), (r'^GSX150F', 'Gixxer SF'), (r'^(GSX150|XF121)', 'Gixxer'), (r'^DS250', 'V-Strom SX'),
             (r'^GE110DDA', 'Hayate EP'), (r'^(GE110|XB25)', 'Hayate'), (r'^GS(150|\()', 'GS150R'), (r'^GT-?125L', 'Slingshot'), (r'^GT', 'Zeus'),
             (r'^GSX1300', 'Hayabusa'), (r'^DL800', 'V-Strom 800DE'), (r'^DL650', 'V-Strom 650'), (r'^DL1000', 'V-Strom 1000'),
             (r'^GSX-S1000', 'GSX-S1000'), (r'^GSX-S750', 'GSX-S750'), (r'^GW250', 'Inazuma 250'), (r'^VZR1800', 'Intruder M1800R')]
# price-list lines for many models at once -> every model of that kind
SZ_SCOOTERS = {'Access 125', 'Burgman Street 125', 'Avenis 125', 'e-Access', 'Swish 125', "Let's"}
SZ_BIG = {'Hayabusa', 'V-Strom 800DE', 'V-Strom 650', 'V-Strom 1000', 'GSX-S1000', 'GSX-S750', 'Inazuma 250', 'Intruder M1800R'}
SZ_MANY = {'COMMON': None, '': None, 'ALLMOTORCYCLES': 'bikes', 'ALLSCOOTERS': 'scooters', 'CBUBIKE': 'big', 'BIGBIKE': 'big'}
# price-list parts in sections by what the part is (Suzuki part numbers start with the part's system: 11-19 engine, 21-29 clutch and gears, ...)
SZ_PL = [('Engine', '11 12 13 14 15 16 17 18 19'), ('Clutch, transmission & kick starter', '21 22 23 24 25 26 27 28 29'),
         ('Electrical, lamps & meter', '31 32 33 34 35 36 37 38 39'), ('Frame, seat & body covers', '41 42 43 44 45 46 47 48 49 53 63 68 92 93 94 95'),
         ('Fork, steering & rear suspension', '51 52 61 62'), ('Handlebar, levers & cables', '56 57 58'), ('Wheels, tyres & brakes', '54 55 59 64 65 69'),
         ('Bolts, nuts, bearings & seals', '01 02 03 04 05 06 07 08 09'), ('Oils, chemicals & accessories', '99')]


class SzPages:
    """Suzuki's diagrams: a page of one of its books (a PDF uploaded as it is) -> 'pdf:' link"""
    def __init__(self, dirn, books):
        self.dir, self.books, self.used = dirn, books, set()

    def link(self, k):
        book, page = k.split('|', 1)
        b = self.books.get(book)
        if not b:
            return None
        self.used.add(book)
        return 'pdf:%s/%s:0:%d:%d:0,0,1,1' % (self.dir, book, b['size'], int(page))

    def files(self):
        return [[f, self.books[f]['size']] for f in sorted(self.used)]


# words the OCR of the two old scanned books got wrong: fixed against the words of Suzuki's clean books and price list
SZ_CHARS = {'π': 'TT', '廿': 'TT', '戊': '/X', '氏': '/X', '民': '/X', '・': '-', '·': '-', '•': '', '˜': '~', '{': '(', '}': ')', '‘': "'"}
SZ_DIG = str.maketrans({'0': 'O', '1': 'I', '5': 'S', '8': 'B', '6': 'G', '2': 'Z', '4': 'A', '|': 'I'})
SZ_OCRW = {'ORING': 'O RING', 'SHAFTNALVE': 'SHAFT/VALVE'}
SZ_KEEP = {'VENDOR', 'MINDA', 'ERING', 'INDER'}                   # right as read, or too unsure to change
# words the scanned books print in titles and remarks that Suzuki's clean books use too rarely to count (COWLING INSTALLATION PARTS,
# NOT ILLUSTRATED, DRAWING NOT SHOWN, vendors NING HAI COUNTY MOTORCYCLE, ROOP POLYMERS, VIR RUBBER)
SZ_WORDS = 'INSTALLATION ILLUSTRATED DRAWING SHOWN NOT NEW COUNTY MOTORCYCLE NING HAI ROOP POLYMERS VIR'
SZ_LOOK = {'O': 'DQC', 'D': 'O', 'I': 'LT', 'L': 'I', 'E': 'F', 'F': 'E', 'C': 'GO', 'G': 'C', 'U': 'V', 'V': 'U', 'N': 'H', 'H': 'N', 'B': 'E'}


def sz_text(s, ocr=False):
    import unicodedata
    s = unicodedata.normalize('NFKC', s or '')
    for a, b in SZ_CHARS.items():
        s = s.replace(a, b)
    if ocr:
        s = re.sub(r'(?<=[A-Z])l|l(?=[A-Z])', 'I', s)          # CYLlNDER, CALlPER
    s = re.sub(r'^[\s.•]+', '', s)                             # Suzuki's sub-part dots (".O RING", "..SEAL SET")
    s = re.sub(r'\s*,\s*', ', ', s)
    return re.sub(r'\s+', ' ', s).strip(' ,')


class SzWords:
    def __init__(self, texts):
        self.f = collections.Counter(w for t in texts for w in re.findall(r'[A-Z]{2,}', t.upper()))

    def ok(self, w):
        return self.f.get(w, 0) >= 2

    def word(self, w):
        """one OCR word -> the word Suzuki prints (or the word as it is)"""
        if self.ok(w) or len(w) < 3 or w in SZ_KEEP:
            return w
        if w in SZ_OCRW:
            return SZ_OCRW[w]
        if re.search(r'\d', w):                                # CAM5HAFT, HOU51NG, 5WITCH; never 2ND, 5AH, 12V, P74, 6X16, GT125XK6
            m = re.fullmatch(r'(\d+(?:ST|ND|RD|TH))([A-Z]{3,})', w)        # 2NDAIR
            if m:
                return m.group(1) + ' ' + self.word(m.group(2))
            if re.fullmatch(r'\d+(?:ST|ND|RD|TH|[A-Z]{1,2})|[A-Z]{1,2}\d+[A-Z]?', w) or len(re.findall(r'\d', w)) > 2 or not re.search(r'[A-Z]{2}', w):
                return w
            t = w.translate(SZ_DIG)
            if self.ok(t):
                return t
            u = self.word(t) if len(t) >= 6 and not re.search(r'\d', t) else t      # COWLINGB0DY = COWLING BODY
            return u if u != t else w
        if len(w) >= 6:                                        # words run together (OILPUMP, CYLINDERHEADCOVER): the fewest words Suzuki uses
            best = [None] * (len(w) + 1)
            best[0] = (0, 0, [])
            for i in range(len(w)):
                if best[i] is None:
                    continue
                for j in range(i + 2, len(w) + 1):
                    a = w[i:j]
                    if self.ok(a):
                        c = (best[i][0] + 1, -min(-best[i][1] if best[i][2] else 1e9, self.f[a]), best[i][2] + [a])
                        if best[j] is None or c[:2] < best[j][:2]:
                            best[j] = c
            if best[-1] and len(best[-1][2]) > 1:
                return ' '.join(best[-1][2])
        if len(w) == 4:                                        # a letter OCR mistakes for one that looks like it (HEAO)
            cand = [(self.f[c], c) for c in {w[:i] + x + w[i + 1:] for i in range(4) for x in SZ_LOOK.get(w[i], '')} if self.f.get(c, 0) >= 3]
            return max(cand)[1] if cand else w
        if len(w) >= 5:                                        # one letter wrong, missing, extra or swapped (VAVLE, COTIER)
            L = 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'
            cand = {w[:i] + c + w[i + 1:] for i in range(len(w)) for c in L} | {w[:i] + w[i + 1:] for i in range(len(w))} \
                | {w[:i] + c + w[i:] for i in range(len(w) + 1) for c in L} | {w[:i] + w[i + 1] + w[i] + w[i + 2:] for i in range(len(w) - 1)}
            cand = [(self.f[c], c) for c in cand if self.f.get(c, 0) >= 3]
            if cand:
                return max(cand)[1]
        return w

    def fix(self, s):
        s = re.sub(r'\b(\d+(?:\.\d+)?)X([\dOIl.X]+)\b', lambda m: m.group(1) + 'X' + m.group(2).replace('O', '0').replace('I', '1').replace('l', '1'), s)
        s = re.sub(r'\bN0\.', 'NO.', s)
        s = re.sub(r'^-?0 RING\b', 'O RING', s)
        s = re.sub(r'[A-Za-z0-9|]{3,}', lambda m: self.word(m.group(0).upper()) if m.group(0).isupper() or re.search(r'\d', m.group(0)) else m.group(0), s)
        # a word the OCR split in two (COWL ING, CL IP)
        s = re.sub(r'\b([A-Z]{2,}) ([A-Z]{1,4})\b', lambda m: m.group(1) + m.group(2) if not self.ok(m.group(2)) and self.ok(m.group(1) + m.group(2)) else m.group(0), s)
        # Zeus model codes in titles and remarks: GT125XCOK71X0UK7IXEUK7 = GT125XCDK7/XDUK7/XEUK7
        s = re.sub(r'(K[67])\s*[I1lJ|/]?\s*(?=[XC][A-Z0-9]{0,3}K[67])', r'\1/', s)
        return re.sub(r'(?<=[/5])(X[A-Z]?)[O0]([A-Z]?K[67])', r'\1D\2', s)

    def clean(self, s):
        """True when every word of s is one Suzuki uses"""
        return not re.search(r'[^\x20-\x7e]|\?', s) and all(self.ok(w) for w in re.findall(r'(?<![A-Za-z0-9])[A-Z]{3,}(?![A-Za-z0-9])', s.upper()))


def sz_rem(r, colours):
    """an OCR book's remark with its words run together: FORYVBJBANEWPN47011-30H00-YVB = FOR YVB, JBA NEW PN 47011-30H00-YVB"""
    r = re.sub(r'NEW ?PN ?(?=[\d(])', ' NEW PN ', re.sub(r'\bFOR(?=MODEL)', 'FOR ', r))
    r = re.sub(r'\bFOR((?:[A-Z0-9]{3})+)\b', lambda m: 'FOR ' + ', '.join(re.findall('...', m.group(1)))
               if all(c in colours for c in re.findall('...', m.group(1))) else m.group(0), r)
    return re.sub(r'\s+', ' ', r).strip()


class SzNear:
    """a description the OCR lost or added a letter or two in, matched to the one name Suzuki prints that is that close
    (the same letters but up to 2, the same digits)"""
    def __init__(self, names):
        self.by = collections.defaultdict(dict)
        for n in names:
            k = re.sub(r'[^A-Z0-9]', '', n.upper())
            if len(k) >= 8:
                self.by[len(k)].setdefault(k, n)

    def get(self, d):
        s = re.sub(r'[^A-Z0-9]', '', d.upper())
        if len(s) < 8:
            return None
        k = min(2, len(s) // 8)
        hits = {}
        for n in range(len(s) - k, len(s) + k + 1):
            for key, name in self.by.get(n, {}).items():
                if re.sub(r'\D', '', key) == re.sub(r'\D', '', s):
                    e = lev(s, key, k)
                    if e <= k:
                        hits.setdefault(e, set()).add(name)
        if not hits:
            return None
        best = hits[min(hits)]
        return best.pop() if len(best) == 1 else None


def lev(a, b, k):
    """edit distance of a and b when it is at most k, else k + 1"""
    if abs(len(a) - len(b)) > k:
        return k + 1
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i] + [0] * len(b)
        for j, cb in enumerate(b, 1):
            cur[j] = min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb))
        if min(cur) > k:
            return k + 1
        prev = cur
    return prev[-1]


def sz_model(t):
    """the model years an OCR book prints run together or with O for 0: MODELLO / MODEL LO = MODEL L0, MODELK9 = MODEL K9"""
    return re.sub(r'\bMODEL ?([KL])([O0-9])\b', lambda m: 'MODEL ' + m.group(1) + m.group(2).replace('O', '0'), t)


def sz_title(t, ocr):
    """a figure's title as the book prints it (its page address and OCR noise left out)"""
    t = sz_text(t, ocr)
    if ocr:
        t = re.sub(r'^(?:FIG|F1G|FlG|\S?G)\.\s*\S{0,3}\s+', '', t)     # "FIG.S ", "FIG. a- ", OCR'd "円G.46 "
        t = re.sub(r'^.*?(?=\b(?:[A-Z]{3,}|\d(?:ST|ND|RD|TH))\b)', '', t, count=1)  # "(8- , 2) ", "1-3 ( ) ", "( -11) "
        t = re.sub(r'(?<=[A-Z0-9])\(', ' (', sz_model(t))
    return t.strip(' -.')


def suzuki():
    SZ = json.load(open(os.path.join(HERE, 'suzuki-pages.json')))
    books, price = SZ['books'], SZ['price']
    assert set(books) == set(SZ_BOOKS) | set(SZ_SCANS), set(books) ^ (set(SZ_BOOKS) | set(SZ_SCANS))
    assert all(not books[f]['lines'] for f in SZ_SCANS)
    W = SzWords([sz_text(l[3] + ' ' + l[6]) for f, b in books.items() if f not in SZ_OCR for l in b['lines']] + [n for n, _, _ in price.values()]
                + [SZ_WORDS] * 2)
    colours = {pn[-3:] for pn in itertools.chain(price, (l[5] for b in books.values() for l in b['lines'])) if re.search(r'[A-Z]', pn[-3:])}  # YPA, AAC: a part no.'s colour
    colours |= {c for b in books.values() for l in b['lines'] for m in re.finditer(r'\bFOR ((?:[A-Z0-9]{3}\b,? ?)+)', sz_text(l[8]))
                for c in re.findall(r'[A-Z0-9]{3}', m.group(1))}                      # FOR YVB, AAC, APT
    near = SzNear([sz_text(l[6]) for f, b in books.items() if f not in SZ_OCR for l in b['lines']] + [sz_text(n) for n, _, _ in price.values()])
    pl_name = lambda pn: sz_text(price[pn][0]) if pn in price else ''
    seen_desc = {}                                             # a part's description in Suzuki's clean books
    for f, b in books.items():
        if f not in SZ_OCR:
            for l in b['lines']:
                if sz_text(l[6]):
                    seen_desc.setdefault(l[5], sz_text(l[6]))
    squash = lambda t: re.sub(r'[^A-Z0-9]', '', t.upper())
    title_of = {}                                              # a title as the clean books print it, by its letters (CYLINDERHEADCOVER)
    for f, b in books.items():
        if f not in SZ_OCR:
            for l in b['lines']:
                t = sz_title(l[3], False)
                if t:
                    title_of.setdefault(squash(t), t)
    rows, stats, in_books = [], collections.Counter(), set()
    for f, (m, code, yrs) in SZ_BOOKS.items():
        ocr = f in SZ_OCR
        v = '%s · %s (%s)' % (m, code, yrs)
        y = max(int(x) for x in re.findall(r'\d{4}', yrs))
        first, fig_at = {}, {}
        if ocr and not books[f].get('ocr'):   # the old OCR text layers: figure numbers read wrong (238 = 23B, 440 = 44D) or lost -> per page
            pages = collections.OrderedDict()
            for l in books[f]['lines']:
                pages.setdefault(l[0], (l[1], W.fix(sz_title(l[3], True))))
            fix = lambda g: g[:2] + {'8': 'B', '0': 'D'}[g[2]] if re.fullmatch(r'\d\d[08]', g) and int(g[:2]) <= 60 else g
            used, last = {fix(g) for g, _ in pages.values()}, None
            for page, (g, t) in pages.items():
                g = new = fix(g)
                if last and g == last[0]:
                    new = last[2]                                              # the same figure going on over this page
                    if t != last[1]:
                        nxt = str(int(re.match(r'\d+', last[2]).group(0)) + 1)  # a new figure whose number the OCR lost
                        if nxt not in used:
                            new = nxt
                            used.add(nxt)
                            stats['figure number put back'] += 1
                fig_at[page], last = new, (g, t, new)
        for page, fig, addr, ttl, ref, pn, desc, qty, rem in books[f]['lines']:
            asm = sz_title(ttl, ocr)
            if ocr:
                asm = title_of.get(squash(asm)) or asm
                fig, asm, ref = fig_at.get(page, fig), re.sub(r'[^\x20-\x7e]+', '?', W.fix(asm)), sz_text(ref, True)
            first.setdefault(fig, page)
            d = sz_text(desc, ocr)
            if ocr:
                d = W.fix(d.lstrip('-~_ ').replace('_', ''))
                if (not W.clean(d) or not re.search(r'[A-Z]{3}', d)) and (pn in seen_desc or pn in price):   # a word the OCR could not read: the same part as another book or the price list names it
                    d = seen_desc.get(pn) or pl_name(pn)
                    stats['OCR description from another book or the price list'] += 1
                elif not W.clean(d) and near.get(d):                      # CLUTCHSETTARTER = CLUTCH SET, STARTER (a letter lost)
                    d = near.get(d)
                    stats['OCR description put right from a name Suzuki prints'] += 1
                d = re.sub(r'[^\x20-\x7e]+', '?', d)
            if not d and pn in price:
                d, stats['description from the price list'] = pl_name(pn), stats['description from the price list'] + 1
            r = sz_text(rem, ocr)
            if ocr:
                r = re.sub(r'[^\x20-\x7e]+', '?', W.fix(sz_model(sz_rem(r, colours))))
            ns = bool(re.search(r'\bNOT? ?SERVICEABLE\b|\bNON[ -]?SERVICEABLE\b', r))
            if ns:
                r = re.sub(r'\s*\bNOT? ?SERVICEABLE\b|\s*\bNON[ -]?SERVICEABLE\b', '', r).strip()
            in_books.add((m, pn))
            rows.append(dict(m=m, v=v, y=y, nb=0, g=fig, asm=asm.upper() if asm else 'FIG. ' + fig, ref=clean(ref), pn=pn, desc=d.upper(),
                             qty=clean(qty), rem=r, ns=ns, dg='%s|%d' % (f, first[fig])))
    # price-list parts that no book of THEIR model has: one more edition per model, in sections by what the part is. The price list names one
    # model per part, so a part the Gixxer book also lists but the price list files under GL150 still goes to Intruder 150 (Sam 2026-10-15:
    # Intruder 150 bearing missing). A model with no book gets years from the price list's model codes (L8 = 2018, M1 = 2021: Suzuki's
    # model-year letters K 2000s, L 2010s, M 2020s) so Find Part's Year list is not empty.
    sec_of = {p: i for i, (_, ps) in enumerate(SZ_PL) for p in ps.split()}
    fam_of = lambda s: next((f for rx, f in SZ_FAMILY if re.match(rx, s)), None)
    many, fams = [], set()
    for pn, (name, model, mrp) in price.items():
        s = re.sub(r'\s+', '', model.upper())
        if (fam_of(s) or 'Other models', pn) in in_books:
            continue
        if s in SZ_MANY:
            many.append((pn, name, model, SZ_MANY[s]))
            continue
        fam = fam_of(s) or 'Other models'
        fams.add(fam)
        rows.append(dict(m=fam, v=fam + ' · more parts (price list)', y=0, nb=1, g='PL%d' % (sec_of.get(pn[:2], len(SZ_PL)) + 1), asm='', ref='', pn=pn,
                         desc=sz_text(name).upper(), qty='', rem=clean(model), ns=False, dg=None))
        stats['price-list part (no book)'] += 1
    for pn, name, model, kind in many:
        for fam in sorted(fams):
            if kind is None or kind == 'scooters' and fam in SZ_SCOOTERS or kind == 'big' and fam in SZ_BIG \
                    or kind == 'bikes' and fam not in SZ_SCOOTERS and fam != 'Other models':
                rows.append(dict(m=fam, v=fam + ' · more parts (price list)', y=0, nb=1, g='PL%d' % (sec_of.get(pn[:2], len(SZ_PL)) + 1), asm='', ref='',
                                 pn=pn, desc=sz_text(name).upper(), qty='', rem=clean(model), ns=False, dg=None))
        stats['price-list part for many models'] += 1
    names = [n for n, _ in SZ_PL] + ['Other parts']
    booked = {x['m'] for x in rows if x['nb'] == 0}
    yrs = collections.defaultdict(set)
    for pn, (name, model, mrp) in price.items():
        s = re.sub(r'\s+', '', model.upper())
        f = fam_of(s)
        for c, d in re.findall(r'(?<=[A-Z0-9])([KLM])(\d)(?!\d)', s):
            if f and f not in booked:
                yrs[f].add({'K': 2000, 'L': 2010, 'M': 2020}[c] + int(d))
    for f, ys in yrs.items():
        lo, hi = min(ys), max(ys)
        lab = ' (%d)' % lo if lo == hi else ' (%d-%d)' % (lo, hi)
        stats['price-list years ' + f + lab] += 1
        for x in rows:
            if x['m'] == f and x['nb'] == 1:
                x['v'] = f + ' · price list parts' + lab
    for x in rows:
        if x['nb'] == 1:
            x['asm'] = names[int(x['g'][2:]) - 1].upper() + ' (PRICE LIST, NO DIAGRAM)'
    unknown = {x['m'] for x in rows} - set(SZ_ORDER)
    assert not unknown, unknown
    print('Suzuki rows:', dict(stats), '| scanned books left for OCR:', len(SZ_SCANS))
    mrp = {pn: (int(p) if p == int(p) else p) for pn, (_, _, p) in price.items() if p and p > 0}
    zp = SzPages(SZ_DIR, books)
    bk = [[0, [], ''], [1, [], 'Parts from Suzuki’s price list that no catalogue book has: no diagram']]
    o, dg = pack('Suzuki', 'suzuki', "Suzuki Motorcycle India parts catalogues, Sam's Suzuki folder, Oct 2026", SZ_ORDER, set(SZ_CUR), rows, 'Fig.', bk, SZ_CC,
                 zp, mrp=mrp, pick='All Catalogues > Suzuki', by_group=True)
    o['src'].append("Suzuki price list (MRP) 03-Oct-2026, Sam's Suzuki folder")
    jdump(os.path.join(OUT, 'suzuki.json'), o)
    return o, dg


# ---------------------------------------------------------------- Royal Enfield (update 9h): RE's diagram books, Sam's folder
# Downloads > Royal-Enfield-Catalogue > Models (16 Oct 2026, from spares.royalenfield.com): one book per edition (colour / version), every
# variant's diagrams, vector, with each item's part no. and MRP. re-pages.json = re_parse.py. Replaces the RE package (EPC data) in Find Part;
# the package's part photos stay (same part numbers, its zips are already on Sam's server).
RE_DIR = 'catalogues/royal-enfield'
RE_NEW = [('FF.C6', 'Flying Flea C6')]
RE_PL, RE_PL_V = 'Other parts (price list)', 'Royal Enfield MRP price list (Oct 2026)'
RE_NORM = {'BSIII': 'BS3', 'BSIV': 'BS4', 'BS IV': 'BS4', 'BSVI': 'BS6'}


class RePages(SzPages):
    """a page of one of RE's diagram books -> 'pdf:' link; the package's photo zips stay in the file list"""
    def __init__(self, dirn, books, keep):
        super().__init__(dirn, books)
        self.keep = keep

    def files(self):
        return super().files() + self.keep


def royal_enfield():
    B = json.load(open(os.path.join(HERE, 're-pages.json')))
    old = json.load(open(os.path.join(HERE, '..', '..', 'public', 'cat', 'royal-enfield.json')))
    key = lambda t: re.sub(r'[^a-z0-9]', '', t.lower())
    old_year = {}
    for v in old['vars']:
        y = re.search(r' \((\d{4})\)$', v)
        if y:
            old_year[key(v[:y.start()])] = int(y.group(1))
    order = HR['RE_ORDER'][:HR['RE_CUR']] + [m for _, m in RE_NEW] + HR['RE_ORDER'][HR['RE_CUR']:] + [RE_PL]
    base = lambda e: next((m for k, m in RE_NEW if e.startswith(k)), None) or HR['re_base'](e)
    rows, mrp, st = [], [], collections.Counter()
    for f, b in sorted(B.items(), key=lambda kv: (kv[1]['edition'], kv[1]['part'])):
        e = b['edition']
        m = base(e)
        n = re.search(r'\b(BSIII|BSIV|BS IV|BSVI)\b', e)
        norms = {x for d in b['dgs'] for x in d['norms']}
        nb = HR['NB'][RE_NORM[n.group(1)] if n else 'BS6' if 'BS6' in norms or not m.startswith('Bullet Trials') else 'BS4']
        lab0 = HR['re_label'](e, 0, nb)
        y = old_year.get(key(lab0), 0)
        v = HR['re_label'](e, y, nb)
        st['editions with the old year'] += bool(y)
        for d in b['dgs']:
            asm = d['asm'] if d['dgm'] in ('', d['asm']) else '%s – %s' % (d['asm'], d['dgm'])
            if d['nv']:
                asm += ' (%d of %d variants%s)' % (d['nv'], b['nvar'], ': ' + ', '.join(d['vars'][:3]) + (' +%d' % (d['nv'] - 3) if d['nv'] > 3 else '')
                                                 if d['vars'] else '')
            for item, pn, desc, p, uom, moq, note in d['rows']:
                mrp.append((pn, p or 0))
                rows.append(dict(m=m, v=v, y=y, nb=nb, g=str(d['n']), asm=asm, ref=item, pn=pn, desc=clean(desc), qty='', rem=note, ns=False,
                                 dg='%s|%d' % (f, d['page'])))
    # MRP: the books' (16 Oct 2026); a part with none there from Sam's RE MRP price list (Oct 2026)
    mrp, two = merge_mrp(mrp)
    import openpyxl
    pl = os.path.join(os.path.dirname(IMP), '..', 'Downloads', 'Royal-Enfield-Catalogue', '_compare', 'RE_MRP_PriceList_2026-10.xlsx')
    pl_mrp, _ = merge_mrp([(clean(str(r[0])), r[2]) for r in itertools.islice(openpyxl.load_workbook(pl, read_only=True).active.iter_rows(values_only=True), 1, None)
                           if r[0] is not None])
    pns = {x['pn'] for x in rows}
    add = {pn: p for pn, p in pl_mrp.items() if pn in pns and pn not in mrp}
    # price-list parts that no book has (special tools, older numbers): model 'Other parts (price list)', a section per first letter, no diagram
    names = {clean(str(r[0])): clean(str(r[1] or '')).upper() for r in itertools.islice(openpyxl.load_workbook(pl, read_only=True).active.iter_rows(values_only=True), 1, None)
             if r[0] is not None}
    flat = {re.sub(r'[^A-Z0-9]', '', p) for p in pns}
    for pn, p in pl_mrp.items():
        if pn not in pns and re.sub(r'[^A-Z0-9]', '', pn) not in flat:
            k = (names.get(pn) or '#')[0]
            k = k if k.isalpha() else '#'
            rows.append(dict(m=RE_PL, v=RE_PL_V, y=0, nb=len(HR['NORMS']), g='PL' + k, asm='PRICE LIST PARTS: ' + k + ' (NO DIAGRAM)', ref='', pn=pn,
                             desc=names.get(pn, ''), qty='', rem='', ns=False, dg=None))
            add[pn] = p
            st['price-list part (no book)'] += 1
    pns = {x['pn'] for x in rows}
    print('Royal Enfield: %s | MRP from the books %d (%d with two prices left out), from the price list %d' % (dict(st), len(mrp), len(two), len(add)))
    mrp.update(add)
    cc = {m: int(re.search(r'\d{3}', m).group(0)) for m in order if re.search(r'\d{3}', m)}
    keep = [[z, s] for z, s in json.load(open(os.path.join(HERE, '..', '..', 'public', 'cat', 'royal-enfield-dg.json')))['files']
            if RE_DIR + '/' + z in old['phz']]                              # only the zip the photos are in (the diagrams in part 1 are not used now)
    zp = RePages(RE_DIR, B, keep)
    bk = [[0, n[1], n[2]] for n in HR['NORMS']] + [[1, [], 'Parts from Royal Enfield’s MRP price list that no diagram book has: no diagram']]
    o, dg = pack('Royal Enfield', 'royal-enfield', "Royal Enfield diagram books (spares.royalenfield.com, MRP as listed 16 Oct 2026), Sam's folder Royal-Enfield-Catalogue",
                 order, set(order[:HR['RE_CUR'] + len(RE_NEW)]), rows, 'Diagram', bk, cc, zp, mrp=mrp, pick='Downloads > Royal-Enfield-Catalogue > Models', by_group=True)
    o['src'].append("Royal Enfield MRP price list (Oct 2026), Sam's folder Royal-Enfield-Catalogue")
    have = {r[5] for r in o['d']}
    o['phf'] = {pn: x for pn, x in old['phf'].items() if pn in have}
    o['phz'] = old['phz']
    jdump(os.path.join(OUT, 'royal-enfield.json'), o)
    print('   photos from the RE package: %d of %d part no.' % (len(o['phf']), len(have)))
    return o, dg


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    for w in WHICH:
        {'hero': hero, 're': royal_enfield, 'yamaha': yamaha, 'honda': honda, 'tvs': tvs, 'bajaj': bajaj, 'suzuki': suzuki}[w]()
